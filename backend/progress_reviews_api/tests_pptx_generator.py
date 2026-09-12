import io

from django.test import SimpleTestCase
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .pptx_generator import generate_progress_review_pptx
from .review_pack import NOT_AVAILABLE

# Names/employer/manager that belong to the two source decks the template was
# built from — see progress_reviews_api/README.md. None of these may ever
# appear in a deck generated for a different learner; every occurrence found
# during development turned out to be a shape the generator had missed.
SOURCE_DECK_NAMES = ["Bethanie", "Grenfell", "Helen McCaughran", "Curtis Cooper", "Andrew"]


def _sample_pack(**overrides) -> dict:
    pack = {
        "learner": {
            "learner_id": "42", "full_name": "Jordan Example", "email": "jordan@example.com",
            "programme": "Marketing Executive Level 4", "apprenticeship_standard": "Marketing Executive Level 4",
            "employer": "Acme Ltd", "role": NOT_AVAILABLE, "manager_name": "Sam Manager",
            "cohort": "Cohort A", "group": "Group 1", "coach": "Coach Casey",
            "programme_start_date": "2025-01-01", "planned_end_date": "2026-06-30", "active_status": True,
        },
        "review": {
            "review_id": "run-1", "review_number": 2, "review_date": "2026-06-15",
            "review_period_start": "2026-03-24", "review_period_end": "2026-06-15",
            "action_period_start": "2026-06-16", "action_period_end": "2026-09-07",
            "generated_by": "coach@kbc.example", "generated_at": "2026-06-15T10:00:00+00:00",
        },
        "attendance": {
            "attendance_percentage": 92, "monthly_summary": [{"month": "May 2026", "sessions": 4, "present": 4, "absent": 0, "late": 0}],
            "missed_sessions": 1, "catch_ups_completed": 1, "catch_ups_needed": 0, "engagement_notes": "Attendance is strong.",
        },
        "progress": {
            "current_programme_progress_percentage": 55, "target_progress_percentage": 60,
            "progress_variance": -5, "next_module": "Module 3: Digital Campaigns",
            "overdue_lms_activities": NOT_AVAILABLE, "action_notes": "Agree the next module with the coach.",
        },
        "otj": {
            "completed_otj_hours": 220.0, "required_otj_hours_to_date": 240.0, "variance": -20.0,
            "forecast_hours": 410.0, "minimum_required_hours": 400.0, "planned_hours": 450.0,
            "risk_status": "Need attention", "duplicate_or_weak_otj_warning": NOT_AVAILABLE,
        },
        "lms_modules": [
            {"module": "Module 1", "completion_percentage": 100, "components_done": 5, "components_total": 5},
            {"module": "Module 2", "completion_percentage": 60, "components_done": 3, "components_total": 5},
        ],
        "evidence": [
            {
                "evidence_title": "Campaign brief.pdf", "evidence_date": "2026-05-01", "evidence_type": "application/pdf",
                "evidence_status": "approved", "evidence_summary": "Wrote a campaign brief for a client launch.",
                "evidence_file_link": NOT_AVAILABLE, "image_or_screenshot_link": NOT_AVAILABLE,
                "manager_verification_status": "accepted", "ai_classification_summary": NOT_AVAILABLE,
                "ksb_mappings": ["K1", "S2"], "evidence_strength": "strong",
            },
        ],
        "assignments": [],
        "workplace_activities": [],
        "ksbs": {
            "knowledge_evidenced": [{"code": "K1", "description": "Understands marketing principles.", "coverage_percentage": 100}],
            "skills_evidenced": [],
            "behaviours_evidenced": [{"code": "B1", "description": "Acts with integrity.", "coverage_percentage": 100}],
            "priority_next": [
                {"code": "S2", "description": "Applies data analysis.", "coverage_percentage": 40,
                 "accepted_count": 1, "how_to_evidence": "Capture a live workplace task.", "evidence_to_retain": "Brief and outcome."},
            ],
        },
        "epa": {
            "current_readiness": "62%", "multiple_choice_test_confidence": "70% average across 3 quiz attempt(s) to date.",
            "project_showcase_confidence": NOT_AVAILABLE, "professional_discussion_confidence": NOT_AVAILABLE,
            "portfolio_risks": ["No outstanding portfolio issues identified."],
            "evidence_admin_risks": ["1 KSB(s) below full coverage — see priority KSBs for detail."],
        },
        "actions": [
            {"title": "Close evidence admin", "detail": "Review pending uploads.", "owner": "Coach and learner", "due_by": "2026-06-16"},
        ],
        "manager_questions": [
            "What improvements have you seen in confidence, independence, or professional judgement?",
        ],
        "source_warnings": [],
    }
    pack.update(overrides)
    return pack


def _iter_all_text(shapes):
    """Recurse into groups — several of this template's dynamic fields (bottom
    action bars, pill-shaped card headers) live one level inside a group, and
    a text search that only checks top-level shapes misses them entirely."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_all_text(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text


def _slide_texts(slide) -> str:
    return "\n".join(_iter_all_text(slide.shapes))


def _deck_texts(prs) -> str:
    return "\n".join(_slide_texts(slide) for slide in prs.slides)


class GenerateProgressReviewPptxTests(SimpleTestCase):
    def test_builds_exactly_nineteen_slides(self):
        data = generate_progress_review_pptx(_sample_pack(), fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        self.assertEqual(len(prs.slides), 19)

    def test_title_slide_carries_the_learner_employer_and_manager(self):
        data = generate_progress_review_pptx(_sample_pack(), fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        text = _slide_texts(prs.slides[0])
        self.assertIn("Jordan Example", text)
        self.assertIn("Marketing Executive Level 4", text)
        self.assertIn("Acme Ltd", text)
        self.assertIn("Sam Manager", text)

    def test_no_source_deck_names_leak_into_a_different_learners_deck(self):
        """Regression guard: every one of these names was found leaking out of
        an unmapped shape (a missed subheading, a nested pill header, a bottom
        action bar) during development — see the module docstring in
        pptx_generator.py for the shape-index provenance of each fix."""
        data = generate_progress_review_pptx(_sample_pack(), fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        full_text = _deck_texts(prs)
        for name in SOURCE_DECK_NAMES:
            self.assertNotIn(name, full_text, f"Found leaked source-deck name {name!r} in the generated deck.")

    def test_missing_evidence_image_uses_a_placeholder_not_a_crash(self):
        pack = _sample_pack()
        pack["evidence"][0]["image_or_screenshot_link"] = "https://example.invalid/evidence.png"
        data = generate_progress_review_pptx(pack, fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        # Evidence detail slides are 8, 9, 10 (0-indexed).
        text = _slide_texts(prs.slides[8])
        self.assertIn("Campaign brief.pdf", text)

    def test_empty_pack_sections_render_without_crashing(self):
        pack = _sample_pack(
            evidence=[], assignments=[], workplace_activities=[],
            ksbs={"knowledge_evidenced": [], "skills_evidenced": [], "behaviours_evidenced": [], "priority_next": []},
            actions=[], manager_questions=[],
        )
        data = generate_progress_review_pptx(pack, fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        self.assertEqual(len(prs.slides), 19)
        text = _slide_texts(prs.slides[8])
        self.assertIn(NOT_AVAILABLE, text)

    def test_not_available_values_are_shown_literally_not_blanked(self):
        data = generate_progress_review_pptx(_sample_pack(), fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        text = _slide_texts(prs.slides[5])  # EPA slide references project showcase confidence
        self.assertIn(NOT_AVAILABLE, text)

    def test_progress_bar_widths_are_proportional_to_the_real_percentage(self):
        """Regression guard: the template's example bars happen to both be
        ~75% wide, which silently hid that bar width was never actually
        resized — see slide_cloner.set_proportional_fill_width."""
        pack = _sample_pack()
        pack["progress"]["current_programme_progress_percentage"] = 20
        pack["progress"]["target_progress_percentage"] = 80
        data = generate_progress_review_pptx(pack, fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        shapes = list(prs.slides[4].shapes)
        current_fill, current_track = shapes[9], shapes[8]
        target_fill, target_track = shapes[13], shapes[12]
        self.assertAlmostEqual(current_fill.width / current_track.width, 0.20, delta=0.02)
        self.assertAlmostEqual(target_fill.width / target_track.width, 0.80, delta=0.02)

    def test_ksb_tables_are_populated_with_real_rows(self):
        data = generate_progress_review_pptx(_sample_pack(), fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        knowledge_table = next(s for s in prs.slides[12].shapes if s.has_table).table
        self.assertEqual(knowledge_table.cell(1, 0).text, "K1")
        self.assertIn("marketing principles", knowledge_table.cell(1, 1).text)

    def test_priority_ksb_codes_are_shown_not_just_descriptions(self):
        data = generate_progress_review_pptx(_sample_pack(), fetch_image=lambda url: None)
        prs = Presentation(io.BytesIO(data))
        text = _slide_texts(prs.slides[15])
        self.assertIn("S2", text)

    def test_template_has_the_expected_slide_count(self):
        from .pptx_generator import TEMPLATE_PATH
        prs = Presentation(str(TEMPLATE_PATH))
        self.assertEqual(len(prs.slides), 19)
