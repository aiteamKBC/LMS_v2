import io
from datetime import date

from django.test import SimpleTestCase
from pptx import Presentation

from . import mcm
from .tests_pptx_generator import _sample_pack


def _all_text(prs) -> str:
    def walk(shapes):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                yield from walk(shape.shapes)
            elif shape.has_text_frame:
                yield shape.text_frame.text
    return "\n".join(t for slide in prs.slides for t in walk(slide.shapes))


class McmPeriodTests(SimpleTestCase):
    def test_default_window_is_the_four_weeks_up_to_the_meeting(self):
        period = mcm.build_mcm_period(date(2026, 9, 29))
        self.assertIsNone(period.review_number)
        self.assertEqual(period.review_period_start, date(2026, 9, 2))
        self.assertEqual(period.review_period_end, date(2026, 9, 29))
        self.assertEqual(period.action_period_start, date(2026, 9, 30))
        self.assertEqual(period.action_period_end, date(2026, 10, 27))


class McmDeckTests(SimpleTestCase):
    def setUp(self):
        self.prs = Presentation(io.BytesIO(mcm.generate_mcm_pptx(_sample_pack(), fetch_image=lambda _url: None)))

    def test_keeps_only_the_monthly_slides(self):
        self.assertEqual(len(self.prs.slides), len(mcm.MCM_SLIDES))

    def test_is_labelled_as_a_monthly_coaching_meeting(self):
        text = _all_text(self.prs)
        self.assertIn("Monthly Coaching Meeting", text)
        self.assertNotIn("Progress Review", text)
        self.assertIn("Wellbeing, Safeguarding & British Values", text)

    def test_deck_opens_after_slide_removal(self):
        buffer = io.BytesIO()
        self.prs.save(buffer)
        self.assertEqual(len(Presentation(io.BytesIO(buffer.getvalue())).slides), len(mcm.MCM_SLIDES))
