import io
from copy import deepcopy
from unittest.mock import patch

from django.test import SimpleTestCase
from pptx import Presentation

from learner_api.standard_mcm_presentation import FORMAT
from . import edits, mcm
from .tests_pptx_generator import _sample_pack


class StandardMcmPipelineTests(SimpleTestCase):
    def setUp(self):
        for target in ('socket.socket.connect', 'socket.create_connection'):
            self.enterContext(patch(target, side_effect=AssertionError('No network permitted')))

    def pack(self):
        pack = _sample_pack()
        pack['review']['presentation_format'] = FORMAT
        pack['manual_slides'] = []
        return pack

    def test_standard_deck_uses_manager_template_real_metrics_and_full_notes_without_mutation(self):
        pack = self.pack(); before = deepcopy(pack)
        deck = Presentation(io.BytesIO(mcm.generate_mcm_pptx(pack, fetch_image=lambda ref: None)))
        self.assertTrue(8 <= len(deck.slides) <= 10)
        text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('Monthly progress and engagement', text)
        self.assertIn('220', text)
        self.assertIn('240', text)
        self.assertIn('92', text)
        notes = '\n'.join(s.notes_slide.notes_text_frame.text for s in deck.slides)
        self.assertIn('Wrote a campaign brief for a client launch.', notes)
        self.assertNotIn(pack['learner']['email'], notes)
        self.assertEqual(pack, before)

    def test_large_evidence_volume_stops_at_twelve_generated_slides(self):
        pack = self.pack()
        pack['evidence'] = [dict(pack['evidence'][0], evidence_title=f'Example {i}', evidence_summary=f'Full original account {i}. ' * 100) for i in range(50)]
        deck = Presentation(io.BytesIO(mcm.generate_mcm_pptx(pack, fetch_image=lambda ref: None)))
        self.assertEqual(len(deck.slides), 12)
        notes = '\n'.join(s.notes_slide.notes_text_frame.text for s in deck.slides)
        for i in range(50): self.assertIn(f'Full original account {i}.', notes)

    def test_learner_can_add_manual_slides_and_partial_edits_preserve_them(self):
        pack = self.pack()
        self.assertEqual(edits.editable_view(pack, review_kind='mcm')['manual_slides'], [])
        manual = [{'title': f'My own slide {i}', 'body': 'My own workplace example.', 'notes': 'My full account.'} for i in range(6)]
        changed = edits.apply_edits(pack, {'manual_slides': manual}, review_kind='mcm', learner_id=42)
        preserved = edits.apply_edits(changed, {'attendance': {'engagement_notes': 'Updated note.'}}, review_kind='mcm', learner_id=42)
        self.assertEqual(preserved['manual_slides'], manual)
        deck = Presentation(io.BytesIO(mcm.generate_mcm_pptx(preserved, fetch_image=lambda ref: None)))
        self.assertGreater(len(deck.slides), 12)
        self.assertIn('My full account.', deck.slides[-1].notes_slide.notes_text_frame.text)
        self.assertEqual(pack['manual_slides'], [])

    def test_old_decks_keep_existing_edit_contract_and_cannot_gain_manual_slides_implicitly(self):
        for kind in ('mcm', 'progress_review'):
            self.assertNotIn('manual_slides', edits.editable_view(_sample_pack(), review_kind=kind))
            with self.assertRaises(edits.EditError):
                edits.apply_edits(_sample_pack(), {'manual_slides': [{'title': 'Title', 'body': 'Body'}]}, review_kind=kind, learner_id=42)

    def test_empty_or_oversized_manual_slide_is_rejected(self):
        for manual in ([{'title': '', 'body': 'Body'}], [{'title': 'Title', 'body': 'x' * 601}], [{}] * 189):
            with self.assertRaises(edits.EditError):
                edits.apply_edits(self.pack(), {'manual_slides': manual}, review_kind='mcm', learner_id=42)

    def test_existing_four_evidence_photo_slots_still_embed_owned_images(self):
        from learner_api.compact_presentation import ASSETS
        raw = (ASSETS / 'kbc-original-logo.png').read_bytes()
        pack = self.pack()
        pack['evidence'] = [dict(pack['evidence'][0], evidence_summary=f'Distinct caption {i}', image_or_screenshot_link=f'blob:test/photo-{i}') for i in range(4)]
        requested = []
        def fetch(ref):
            requested.append(ref)
            return raw
        deck = Presentation(io.BytesIO(mcm.generate_mcm_pptx(pack, fetch_image=fetch)))
        self.assertEqual(requested, [f'blob:test/photo-{i}' for i in range(4)])
        text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        for i in range(4):
            self.assertIn(f'Distinct caption {i}', text)
        self.assertTrue(8 <= len(deck.slides) <= 10)
        self.assertEqual(sum(s.shape_type == 13 and s.image.blob == raw for slide in deck.slides for s in slide.shapes), 4)

    def test_unavailable_or_invalid_image_is_visible_without_losing_evidence_text(self):
        for raw in (None, b'invalid image'):
            pack = self.pack()
            pack['evidence'][0]['image_or_screenshot_link'] = 'blob:test/photo'
            deck = Presentation(io.BytesIO(mcm.generate_mcm_pptx(pack, fetch_image=lambda ref: raw)))
            text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
            self.assertIn('Image unavailable.', text)
            self.assertIn('Wrote a campaign brief', '\n'.join(s.notes_slide.notes_text_frame.text for s in deck.slides))

    def test_fresh_generate_selects_standard_but_uploaded_decks_keep_their_existing_contract(self):
        from datetime import date
        from types import SimpleNamespace
        from .views import _generate_for_learner
        for uploaded in (None, b'own presentation'):
            pack = _sample_pack()
            with patch('progress_reviews_api.views._programme_window', return_value=(date(2025, 1, 1), None, 'apprenticeship', SimpleNamespace())), \
                 patch('progress_reviews_api.views._curriculum_programme_id', return_value=None), \
                 patch('progress_reviews_api.views.build_review_pack', return_value=pack), \
                 patch('progress_reviews_api.views._store_deck', return_value={'reviewId': 'new'}) as store:
                _generate_for_learner(42, review_date=date(2026, 10, 15), review_kind='mcm', uploaded_pptx=uploaded)
            saved = store.call_args.kwargs['pack']
            if uploaded is None:
                self.assertEqual(saved['review']['presentation_format'], FORMAT)
                self.assertEqual(saved['manual_slides'], [])
            else:
                self.assertNotIn('presentation_format', saved['review'])

    def test_missing_percentages_are_reported_without_inventing_values(self):
        pack = self.pack()
        pack.update(progress={}, attendance={}, lms_modules=[{'module': 'New module'}])
        deck = Presentation(io.BytesIO(mcm.generate_mcm_pptx(pack, fetch_image=lambda ref: None)))
        text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('Current: Not recorded.', text)
        self.assertNotIn('Not recorded.%', text)
