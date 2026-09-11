"""Low-level slide manipulation for template-cloning PPTX generation.

python-pptx has no native "duplicate a slide" operation. `duplicate_slide`
below is the standard community recipe: add a slide on the same layout, strip
the placeholder shapes it comes with, deep-copy every shape element from the
source slide's shape tree, and rewrite every relationship id the copied XML
references (image fills, hyperlinks) so they point at newly-added
relationships on the destination slide's part rather than the source's.

Only needed for CONTINUATION slides (a KSB table or an evidence list that
overflows what one slide can hold) — every other slide in a generated deck is
one of the 19 template slides, mutated in place, never duplicated.

The rest of this module is text/table/picture-fill mutation that preserves
run-level formatting: setting a run's text is always cheaper and safer than
deleting and recreating paragraphs, which silently drops the exact
Aptos/Aptos Bold/colour/size formatting extracted into pptx_theme.py.
"""
from __future__ import annotations

import copy
from typing import Optional

from pptx.oxml.ns import qn
from pptx.slide import Slide

R_EMBED = qn("r:embed")
R_LINK = qn("r:link")
R_ID = qn("r:id")
REL_ATTRS = (R_EMBED, R_LINK, R_ID)


def duplicate_slide(prs, index: int) -> Slide:
    """Append a copy of `prs.slides[index]` at the end of the deck and return it."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    # add_slide() populates placeholder shapes from the layout — remove them,
    # the source's own shapes (including its own placeholder copies) replace them.
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    rid_map = {}
    for rid, rel in source.part.rels.items():
        if rel.is_external:
            new_rid = dest.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            new_rid = dest.part.relate_to(rel.target_part, rel.reltype)
        rid_map[rid] = new_rid

    spTree = source.shapes._spTree
    skip_tags = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
    for element in spTree:
        if element.tag in skip_tags:
            continue
        new_element = copy.deepcopy(element)
        for node in new_element.iter():
            for attr in REL_ATTRS:
                old_rid = node.get(attr)
                if old_rid and old_rid in rid_map:
                    node.set(attr, rid_map[old_rid])
        dest.shapes._spTree.append(new_element)

    return dest


def move_slide(prs, from_index: int, to_index: int) -> None:
    """Reposition a slide (e.g. move a just-appended continuation slide to sit
    directly after the slide it continues, rather than at the end of the deck)."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[from_index])
    xml_slides.insert(to_index, slides[from_index])


def iter_text_shapes(slide: Slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape


def find_table_shape(slide: Slide, occurrence: int = 0):
    """The `occurrence`-th (0-based) table shape on the slide, in shape order."""
    found = 0
    for shape in slide.shapes:
        if shape.has_table:
            if found == occurrence:
                return shape
            found += 1
    return None


def set_run_text(shape, text: str, *, paragraph_index: int = 0, run_index: int = 0) -> None:
    """Overwrite one run's text, preserving that run's own formatting.

    Extra runs in the same paragraph (a paragraph split into several runs for
    partial bold/colour, e.g. "Current: **75%**") are cleared to '' rather than
    removed, so paragraph-level spacing/alignment is untouched.
    """
    paragraphs = shape.text_frame.paragraphs
    if paragraph_index >= len(paragraphs):
        return
    runs = paragraphs[paragraph_index].runs
    if not runs:
        return
    for i, run in enumerate(runs):
        run.text = text if i == run_index else ""


def set_paragraph_text(shape, paragraph_index: int, text: str) -> None:
    """Overwrite an entire paragraph's visible text in its first run, using
    that run's formatting for the whole replacement (for a paragraph whose
    original content was a single run — the common case for body copy)."""
    set_run_text(shape, text, paragraph_index=paragraph_index, run_index=0)


def set_all_text(shape, text: str) -> None:
    """Replace a whole text frame with one paragraph, but reuse the FIRST
    existing run's formatting rather than python-pptx's own default run
    formatting (which would silently drop the extracted Aptos styling)."""
    frame = shape.text_frame
    first_paragraph = frame.paragraphs[0]
    if not first_paragraph.runs:
        first_paragraph.text = text
        return
    template_run = first_paragraph.runs[0]
    for extra in list(first_paragraph.runs[1:]):
        extra._r.getparent().remove(extra._r)
    for extra_paragraph in list(frame.paragraphs[1:]):
        extra_paragraph._p.getparent().remove(extra_paragraph._p)
    template_run.text = text


def set_bullet_paragraphs(shape, items: list[str], *, header_count: int = 0) -> None:
    """Populate a bullet-list shape's paragraphs AFTER its first `header_count`
    paragraphs (a card title like "Coach summary" left untouched), cloning the
    first bullet paragraph's formatting for each item and dropping/adding
    paragraphs to match `items`' length.
    """
    frame = shape.text_frame
    paragraphs = frame.paragraphs
    if len(paragraphs) <= header_count or not paragraphs[header_count].runs:
        return
    template_p = paragraphs[header_count]._p
    parent = template_p.getparent()

    while len(frame.paragraphs) > header_count + 1:
        frame.paragraphs[-1]._p.getparent().remove(frame.paragraphs[-1]._p)

    for _ in range(max(0, len(items) - 1)):
        parent.append(copy.deepcopy(template_p))

    for paragraph, text in zip(frame.paragraphs[header_count:], items):
        set_paragraph_text_on(paragraph, text)


def set_paragraph_text_on(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for extra in list(runs[1:]):
        extra._r.getparent().remove(extra._r)


def set_table_cell(shape, row: int, col: int, text: str) -> None:
    table = shape.table
    if row >= len(table.rows) or col >= len(table.columns):
        return
    cell = table.cell(row, col)
    set_all_text_in_cell(cell, text)


def set_all_text_in_cell(cell, text: str) -> None:
    frame = cell.text_frame
    first_paragraph = frame.paragraphs[0]
    if not first_paragraph.runs:
        first_paragraph.text = text
        return
    template_run = first_paragraph.runs[0]
    for extra in list(first_paragraph.runs[1:]):
        extra._r.getparent().remove(extra._r)
    for extra_paragraph in list(frame.paragraphs[1:]):
        extra_paragraph._p.getparent().remove(extra_paragraph._p)
    template_run.text = text


def delete_table_row(shape, row_index: int) -> None:
    table = shape.table
    tbl = table._tbl
    rows = tbl.findall(qn("a:tr"))
    if row_index < len(rows):
        tbl.remove(rows[row_index])


def duplicate_table_row(shape, row_index: int) -> None:
    table = shape.table
    tbl = table._tbl
    rows = tbl.findall(qn("a:tr"))
    if row_index >= len(rows):
        return
    new_row = copy.deepcopy(rows[row_index])
    rows[row_index].addnext(new_row)


def _blip_of(shape):
    """The <a:blip> element behind this shape's picture fill, or None.

    This template's chrome (logo, background art, evidence photos) is built as
    FREEFORM shapes with a picture fill (see pptx_theme.py's module
    docstring), each wrapped in its own single-child GROUP — never a bare
    top-level shape, so callers must walk into groups to find them.
    """
    spPr = shape._element.find(qn("p:spPr"))
    if spPr is None:
        return None
    return spPr.find(qn("a:blipFill") + "/" + qn("a:blip"))


def _picture_fill_shapes_in(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    results = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            results.extend(_picture_fill_shapes_in(shape.shapes))
            continue
        blip = _blip_of(shape)
        if blip is not None:
            results.append((shape, blip))
    return results


def picture_fill_shapes(slide: Slide):
    """Every shape whose fill is a picture, searched recursively through
    groups. Returns (shape, blip_element) pairs in document order."""
    return _picture_fill_shapes_in(slide.shapes)


def replace_picture_fill(shape, image_bytes: bytes, slide_part) -> None:
    """Swap the image behind a picture-filled shape (an evidence photo slot)
    for `image_bytes`, keeping the shape's exact position/size/crop/rounding."""
    import io

    image_part, rId = slide_part.get_or_add_image_part(io.BytesIO(image_bytes))
    blip = _blip_of(shape)
    if blip is not None:
        blip.set(R_EMBED, rId)


def shape_bounds_in(shape) -> Optional[tuple]:
    if shape.left is None or shape.top is None:
        return None
    return (
        shape.left / 914400, shape.top / 914400,
        (shape.left + shape.width) / 914400, (shape.top + shape.height) / 914400,
    )
