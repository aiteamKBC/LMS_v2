import json
import os
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase

from . import edits
from .review_pack import NOT_AVAILABLE
from .tests_pptx_generator import _sample_pack

UPLOAD = "blob:progress-review-decks/edits/42/abc.jpg"


class EditableViewTests(SimpleTestCase):
    def test_not_available_becomes_null_and_evidence_is_addressable(self):
        pack = _sample_pack()
        pack["attendance"]["attendance_percentage"] = NOT_AVAILABLE
        view = edits.editable_view(pack, review_kind="progress_review")
        self.assertIsNone(view["attendance"]["attendance_percentage"])
        self.assertTrue(all(item["ref"] for item in view["evidence"]))
        self.assertIn("epa", view)

    def test_mcm_view_has_no_review_only_fields(self):
        view = edits.editable_view(_sample_pack(), review_kind="mcm")
        self.assertNotIn("epa", view)
        self.assertNotIn("manager_questions", view)


class ApplyEditsTests(SimpleTestCase):
    def apply(self, changes, pack=None, kind="progress_review"):
        return edits.apply_edits(pack or _sample_pack(), changes, review_kind=kind, learner_id=42)

    def test_corrected_figures_recompute_their_variance(self):
        pack = self.apply({
            "progress": {"current_programme_progress_percentage": 30, "target_progress_percentage": 40},
            "otj": {"completed_otj_hours": 50.5, "required_otj_hours_to_date": 60, "risk_status": "Need attention"},
        })
        self.assertEqual(pack["progress"]["progress_variance"], -10)
        self.assertEqual(pack["otj"]["variance"], -9.5)
        self.assertEqual(pack["otj"]["risk_status"], "Need attention")

    def test_cleared_field_becomes_not_available(self):
        pack = self.apply({"attendance": {"attendance_percentage": None}})
        self.assertEqual(pack["attendance"]["attendance_percentage"], NOT_AVAILABLE)

    def test_invalid_values_are_rejected(self):
        for changes in (
            {"attendance": {"attendance_percentage": 140}},
            {"otj": {"risk_status": "Great"}},
            {"learner": {"full_name": ""}},
            {"actions": [{"title": "x", "due_by": "next week"}]},
            {"priority_ksbs": [{"code": "not a code"}]},
        ):
            with self.subTest(changes=changes), self.assertRaises(edits.EditError):
                self.apply(changes)

    def test_the_original_pack_is_not_modified(self):
        pack = _sample_pack()
        self.apply({"learner": {"employer": "New Co"}}, pack=pack)
        self.assertNotEqual(pack["learner"]["employer"], "New Co")

    def test_evidence_photo_can_be_replaced_with_this_learners_upload(self):
        view = edits.editable_view(_sample_pack(), review_kind="progress_review")
        view["evidence"][0]["image_ref"] = UPLOAD
        pack = self.apply({"evidence": view["evidence"]})
        self.assertEqual(pack["slide_evidence"][0]["image_or_screenshot_link"], UPLOAD)

    def test_another_learners_or_foreign_photo_is_refused(self):
        view = edits.editable_view(_sample_pack(), review_kind="progress_review")
        for ref in ("blob:progress-review-decks/edits/7/abc.jpg", "https://example.com/x.jpg", "blob:evidence-approved/x.jpg"):
            view["evidence"][0]["image_ref"] = ref
            with self.subTest(ref=ref), self.assertRaises(edits.EditError):
                self.apply({"evidence": view["evidence"]})

    def test_evidence_items_can_be_added_and_removed(self):
        view = edits.editable_view(_sample_pack(), review_kind="progress_review")
        kept = view["evidence"][:1]
        added = {"ref": None, "evidence_title": "Campaign launch", "evidence_summary": "Ran the launch.",
                 "evidence_date": "2026-09-10", "ksb_mappings": ["s2"], "image_ref": UPLOAD}
        pack = self.apply({"evidence": [*kept, added]})
        self.assertEqual([i["evidence_title"] for i in pack["slide_evidence"]][-1], "Campaign launch")
        self.assertEqual(edits.editable_view(pack, review_kind="progress_review")["evidence"][-1]["ksb_mappings"], ["S2"])

    def test_the_owners_order_is_the_order_on_the_slides(self):
        from .pptx_generator import _evidence_blocks

        source = _sample_pack()
        source["evidence"].append({**source["evidence"][0], "evidence_title": "Second.pdf"})
        view = edits.editable_view(source, review_kind="progress_review")
        reordered = list(reversed(view["evidence"]))
        self.assertEqual(reordered[0]["evidence_title"], "Second.pdf")
        pack = self.apply({"evidence": reordered}, pack=source)
        self.assertEqual(
            [b["evidence_title"] for b in _evidence_blocks(pack, 99)],
            [i["evidence_title"] for i in reordered],
        )
        # A second edit addresses the saved order.
        again = edits.editable_view(pack, review_kind="progress_review")["evidence"]
        self.assertTrue(all(i["ref"].startswith("slide_evidence:") for i in again))

    def test_an_unknown_evidence_ref_is_refused(self):
        with self.assertRaises(edits.EditError):
            self.apply({"evidence": [{"ref": "assignments:99", "evidence_title": "x"}]})

    def test_mcm_edit_ignores_review_only_fields(self):
        pack = self.apply({"epa": {"current_readiness": 99}}, kind="mcm")
        self.assertNotEqual(pack["epa"].get("current_readiness"), "99%")


def _account(role, subject_id=0, email="user@example.invalid"):
    return SimpleNamespace(role=role, subject_id=subject_id, email=email, is_staff=role != "learner")


RUN = {
    "id": "run-1", "learner_id": 42, "learner_kind": "apprenticeship", "generation_status": "completed",
    "review_number": None, "review_date": date(2026, 9, 29),
    "review_period_start": date(2026, 9, 2), "review_period_end": date(2026, 9, 29),
    "action_period_start": date(2026, 9, 30), "action_period_end": date(2026, 10, 27),
}


class EditOwnershipTests(SimpleTestCase):
    """The learner owns the MCM deck; the coach owns the Progress Review deck."""

    def status(self, kind, account):
        with patch("progress_reviews_api.views.runs") as mock_runs, \
                patch("progress_reviews_api.views.authenticate_request", return_value=account), \
                patch.dict(os.environ, {"LEARNER_API_REQUIRE_AUTH": "1"}):
            mock_runs.get_run.return_value = {**RUN, "review_kind": kind}
            mock_runs.get_snapshot.return_value = _sample_pack()
            return self.client.get("/progress_reviews_api/run-1/edit/").status_code

    def test_mcm_deck_is_editable_by_its_learner_only(self):
        self.assertEqual(self.status("mcm", _account("learner", 42)), 200)
        self.assertEqual(self.status("mcm", _account("learner", 7)), 403)
        self.assertEqual(self.status("mcm", _account("coach")), 403)

    def test_progress_review_deck_is_editable_by_staff_only(self):
        self.assertEqual(self.status("progress_review", _account("coach")), 200)
        self.assertEqual(self.status("progress_review", _account("learner", 42)), 403)

    @patch("progress_reviews_api.views._store_deck", return_value={"reviewId": "run-2"})
    def test_saving_an_edit_creates_a_new_version_of_the_deck(self, mock_store):
        with patch("progress_reviews_api.views.runs") as mock_runs, \
                patch("progress_reviews_api.views.authenticate_request", return_value=_account("coach")):
            mock_runs.get_run.return_value = {**RUN, "review_kind": "progress_review"}
            mock_runs.get_snapshot.return_value = _sample_pack()
            mock_runs.new_run_id.return_value = "run-2"
            response = self.client.post(
                "/progress_reviews_api/run-1/edit/",
                data=json.dumps({"edits": {"learner": {"employer": "New Co"}}}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 201)
        kwargs = mock_store.call_args.kwargs
        self.assertEqual(kwargs["parent_run_id"], "run-1")
        self.assertEqual(kwargs["revision_source"], "edited")
        self.assertEqual(kwargs["pack"]["learner"]["employer"], "New Co")

    def test_an_invalid_edit_is_a_400_with_the_reason(self):
        with patch("progress_reviews_api.views.runs") as mock_runs, \
                patch("progress_reviews_api.views.authenticate_request", return_value=_account("coach")):
            mock_runs.get_run.return_value = {**RUN, "review_kind": "progress_review"}
            mock_runs.get_snapshot.return_value = _sample_pack()
            response = self.client.post(
                "/progress_reviews_api/run-1/edit/",
                data=json.dumps({"edits": {"attendance": {"attendance_percentage": 500}}}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Attendance", json.loads(response.content)["error"])

    def test_a_non_pptx_upload_is_refused(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        with patch("progress_reviews_api.views.runs") as mock_runs, \
                patch("progress_reviews_api.views.authenticate_request", return_value=_account("coach")):
            mock_runs.get_run.return_value = {**RUN, "review_kind": "progress_review"}
            response = self.client.post(
                "/progress_reviews_api/run-1/edit/upload/",
                data={"file": SimpleUploadedFile("deck.pptx", b"not a deck")},
            )
        self.assertEqual(response.status_code, 400)


class RoundTripTests(SimpleTestCase):
    def test_saving_the_generated_fields_unchanged_always_succeeds(self):
        """Real generated text (full KSB standard wording, long file names)
        must never fail validation just by being saved back as it was."""
        pack = _sample_pack()
        pack["ksbs"]["priority_next"][0]["description"] = "x" * 900
        pack["evidence"][0]["evidence_title"] = "y" * 250
        for kind in ("mcm", "progress_review"):
            view = edits.editable_view(pack, review_kind=kind)
            fields = {k: v for k, v in view.items() if k not in ("kind", "evidence_photo_slots")}
            edits.apply_edits(pack, fields, review_kind=kind, learner_id=42)


class UploadOwnDeckTests(SimpleTestCase):
    """The owner can use their own presentation instead of the generated one."""

    def post(self, data):
        from django.core.files.uploadedfile import SimpleUploadedFile

        with patch("progress_reviews_api.views.authenticate_request", return_value=_account("coach")), \
                patch.dict(os.environ, {"LEARNER_API_REQUIRE_AUTH": "0"}):
            return self.client.post(
                "/progress_reviews_api/42/upload/",
                data={"file": SimpleUploadedFile("mine.pptx", data), "review_date": "2026-09-16"},
            )

    def test_a_non_pptx_file_is_refused(self):
        self.assertEqual(self.post(b"not a deck").status_code, 400)

    @patch("progress_reviews_api.views._generate_for_learner", return_value={"reviewId": "run-9"})
    def test_the_uploaded_file_is_stored_as_is_with_the_real_data_pack(self, mock_generate):
        import io as _io

        from pptx import Presentation

        buffer = _io.BytesIO()
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0])
        deck.save(buffer)

        response = self.post(buffer.getvalue())

        self.assertEqual(response.status_code, 201)
        kwargs = mock_generate.call_args.kwargs
        self.assertEqual(kwargs["review_kind"], "progress_review")
        self.assertEqual(kwargs["review_date"], date(2026, 9, 16))
        self.assertEqual(kwargs["uploaded_pptx"], buffer.getvalue())
