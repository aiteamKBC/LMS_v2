"""(Re)builds templates/kbc_mcm_template.pptx — the MCM-only visual template.

    python manage.py build_mcm_template

This is a BUILD-TIME tool, not part of any request path: it writes the
template file to disk; mcm_shapes.py only ever *reads* the resulting .pptx at
generation time, exactly the way pptx_generator.py only ever reads
kbc_progress_review_template.pptx (see that module's own docstring). Re-run
this after editing the layout constants below, then re-render sample decks
(tests_mcm_template.py + a manual render) before shipping a design change.

Why a separate template from the Progress Review one
------------------------------------------------------
kbc_progress_review_template.pptx is a real, human-produced deck (see
progress_reviews_api/README.md) shared by both the Progress Review and the
(now former) sliced-down MCM renderer. Redesigning its visuals would have
silently redesigned the standard Progress Review deck too — the audit's stop
condition. This template is authored FROM SCRATCH with python-pptx instead of
being cloned from a human deck, carries only the ten MCM slides, and is never
read by pptx_generator.py's Progress Review path. Building it this way (as
opposed to hand-authoring a .pptx in PowerPoint) means the visual language can
be regenerated deterministically from these constants rather than requiring a
design tool round-trip for every tweak.

Visual language, extracted from the supplied reference deck ("Learner Monthly
Coaching Presentation"), not invented:
  - background FBFAFF (content slides), a KBC-branded wave/logo picture-fill
    background on the cover slide only (ppt/media/image1.png in the reference)
  - a repeating "card" grammar: white rounded-rectangle body, a slightly
    inset colour header strip carrying a bold white label, muted-navy body
    copy underneath (colours 4E36C8 / 6B5AE0 / D8A23D header rotation, 2A2356
    headings, 403A6B body — read directly from the reference's runs)
  - a full-width, deep-purple "band" at the foot of every content slide
    carrying a short bold gold-coloured label + a white sentence (the
    reference's "COACH TIP" band; generalised here to whichever short
    practical note each slide already carries)
  - Aptos Display for slide titles, Aptos for everything else (matches the
    reference's own font choice, and pptx_theme.py's existing FONT_BODY/
    FONT_BOLD, so no new font dependency is introduced)

Every shape that the MCM populator (mcm_shapes.py) writes dynamic data into is
given an explicit, semantic `shape.name` at build time here (e.g.
"field:learner_name", "bar:progress:fill", "photo:evidence:1") — see
mcm_shapes.SHAPE prefixes. Nothing in this file is looked up by numeric index
at generation time; numeric shape indices are exactly the fragility the audit
flagged, and this template exists specifically so MCM no longer depends on
them.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

from django.core.management.base import BaseCommand

from ... import pptx_theme

ASSETS_DIR = Path(__file__).resolve().parent.parent.parent / "templates" / "assets"
TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "templates" / "kbc_mcm_template.pptx"

EMU_PER_IN = 914400
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5

# Header rotation across a card row, in the exact three colours the
# reference's own 3-card slides alternate (see progress_reviews_api/README.md
# provenance note and the reference dump).
HEADER_ROTATION = (pptx_theme.THEME.mcm_purple_deep, pptx_theme.THEME.mcm_gold, pptx_theme.THEME.mcm_purple_light)


def _c(rgb: RGBColor) -> RGBColor:
    return rgb


def _in(value: float) -> Emu:
    return Inches(value)


def _set_run(shape, text, *, size, color, bold=False, font=pptx_theme.FONT_BODY, align=PP_ALIGN.LEFT):
    frame = shape.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    for extra in list(p.runs):
        extra._r.getparent().remove(extra._r)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color


def _textbox(slide, name, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
    box.name = name
    box.text_frame.vertical_anchor = anchor
    box.text_frame.margin_left = box.text_frame.margin_right = Pt(2)
    box.text_frame.margin_top = box.text_frame.margin_bottom = Pt(1)
    return box


def _rect(slide, name, x, y, w, h, *, fill, rounded=True, line=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, _in(x), _in(y), _in(w), _in(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = pptx_theme.THEME.mcm_card_border
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    if rounded:
        try:
            shape.adjustments[0] = 0.06
        except (IndexError, AttributeError):
            pass
    return shape


def _background(slide, colour):
    bg = _rect(slide, "background", 0, 0, SLIDE_W_IN, SLIDE_H_IN, fill=colour, rounded=False)
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return bg


def _logo_mark(slide, *, x=None, y=None, w=0.62):
    """The small KBC shield watermark every content slide carries. A single,
    explicitly-named picture this module places itself — there is no
    "foreign" logo inherited from a source deck to mistake it for (see
    pptx_generator._remove_source_employer_logos, which exists only because
    the shared Progress Review template is cloned from two human-made decks;
    this template has no such inherited chrome)."""
    h = w * (582 / 662)
    x = SLIDE_W_IN - w - 0.28 if x is None else x
    y = SLIDE_H_IN - h - 0.55 if y is None else y
    pic = slide.shapes.add_picture(str(ASSETS_DIR / "kbc_logo_mark.png"), _in(x), _in(y), width=_in(w), height=_in(h))
    pic.name = "chrome:kbc_logo_mark"
    return pic


def _header_block(slide, *, title_key, subtitle_key, breadcrumb_key):
    """Breadcrumb / title / subheading, in that order, on every content
    slide — the same three fields pptx_generator._set_breadcrumb and
    _populate_* write today, just under semantic names instead of a shape
    index."""
    breadcrumb = _textbox(slide, breadcrumb_key, 0.7, 0.2, 9.5, 0.3)
    _set_run(breadcrumb, "", size=10, color=pptx_theme.THEME.mcm_purple_light, bold=True)
    title = _textbox(slide, title_key, 0.7, 0.48, 10.4, 0.62)
    _set_run(title, "", size=25, color=pptx_theme.THEME.mcm_heading, bold=True, font=pptx_theme.FONT_DISPLAY)
    subtitle = _textbox(slide, subtitle_key, 0.7, 1.08, 10.4, 0.32)
    _set_run(subtitle, "", size=13, color=pptx_theme.THEME.mcm_purple_light, bold=True)
    return breadcrumb, title, subtitle


def _band(slide, name_prefix, *, label, y=None, h=0.46):
    """The deep-purple footer band every content slide carries — the
    reference's "COACH TIP" treatment, generalised: a short bold gold label
    plus a body sentence, both dynamic."""
    y = SLIDE_H_IN - h if y is None else y
    _rect(slide, f"{name_prefix}:bg", 0, y, SLIDE_W_IN, h, fill=pptx_theme.THEME.mcm_purple_deep, rounded=False)
    label_box = _textbox(slide, f"{name_prefix}:label", 0.55, y, 2.6, h, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(label_box, label, size=11.5, color=pptx_theme.THEME.mcm_gold, bold=True)
    text_box = _textbox(slide, f"{name_prefix}:text", 3.0, y, SLIDE_W_IN - 3.4, h, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(text_box, "", size=11.5, color=RGBColor.from_string("FFFFFF"), bold=False)
    return label_box, text_box


def _card(slide, name, x, y, w, h, *, header_color, header_label, header_h=0.5, body_size=12.5):
    """One reference-style card: white rounded body, an inset colour header
    strip with a bold white label, and a body text box beneath it for
    bulleted or free-text content."""
    _rect(slide, f"{name}:card", x, y, w, h, fill=RGBColor.from_string("FFFFFF"), line=True)
    inset = 0.02
    _rect(slide, f"{name}:header", x + inset, y + inset, w - 2 * inset, header_h, fill=header_color)
    label_box = _textbox(slide, f"{name}:header:label", x + 0.16, y + inset, w - 0.32, header_h, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(label_box, header_label, size=13.5, color=RGBColor.from_string("FFFFFF"), bold=True)
    body_box = _textbox(slide, f"{name}:body", x + 0.16, y + header_h + 0.14, w - 0.32, h - header_h - 0.28)
    _set_run(body_box, "", size=body_size, color=pptx_theme.THEME.mcm_body)
    return label_box, body_box


def _kpi_card(slide, name, x, y, w, h, *, header_color, header_label):
    """A card whose body is one large value plus a small caption, for a
    single headline metric (attendance %, OTJ hours, …) rather than a bullet
    list."""
    _rect(slide, f"{name}:card", x, y, w, h, fill=RGBColor.from_string("FFFFFF"), line=True)
    inset = 0.02
    _rect(slide, f"{name}:header", x + inset, y + inset, w - 2 * inset, 0.4, fill=header_color)
    label_box = _textbox(slide, f"{name}:header:label", x + 0.14, y + inset, w - 0.28, 0.4, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(label_box, header_label, size=11.5, color=RGBColor.from_string("FFFFFF"), bold=True)
    value_box = _textbox(slide, f"{name}:value", x + 0.14, y + 0.5, w - 0.28, h - 0.95, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(value_box, "", size=26, color=pptx_theme.THEME.mcm_heading, bold=True, align=PP_ALIGN.CENTER)
    caption_box = _textbox(slide, f"{name}:caption", x + 0.1, y + h - 0.4, w - 0.2, 0.36, anchor=MSO_ANCHOR.TOP)
    _set_run(caption_box, "", size=9.5, color=pptx_theme.THEME.mcm_body, align=PP_ALIGN.CENTER)
    return value_box, caption_box


def _progress_bar(slide, name, x, y, w, h=0.22, *, track_color=None):
    track_color = track_color or pptx_theme.THEME.mcm_track
    track = _rect(slide, f"bar:{name}:track", x, y, w, h, fill=track_color)
    fill = _rect(slide, f"bar:{name}:fill", x, y, 0.02, h, fill=pptx_theme.THEME.mcm_purple_deep)
    return track, fill


def _chip(slide, name, x, y, w, h=0.34):
    """A status pill (RAG). Fill and text colour are both set at generation
    time from the same business RAG value — see mcm_shapes._apply_rag_chip."""
    chip = _rect(slide, f"chip:{name}:bg", x, y, w, h, fill=pptx_theme.THEME.rag_amber_bg)
    text = _textbox(slide, f"chip:{name}:text", x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(text, "", size=11, color=pptx_theme.THEME.rag_amber, bold=True, align=PP_ALIGN.CENTER)
    return chip, text


def _photo_frame(slide, name, x, y, w, h):
    frame = _rect(slide, f"photo:{name}:frame", x, y, w, h, fill=RGBColor.from_string("FFFFFF"), line=True)
    picture = slide.shapes.add_picture(str(ASSETS_DIR / "kbc_logo_mark.png"), _in(x + 0.03), _in(y + 0.03), width=_in(w - 0.06), height=_in(h - 0.06))
    picture.name = f"photo:{name}:picture"
    caption = _textbox(slide, f"photo:{name}:caption", x, y + h + 0.03, w, 0.3)
    _set_run(caption, "", size=10, color=pptx_theme.THEME.mcm_body, bold=True, align=PP_ALIGN.CENTER)
    return picture, caption


# --------------------------------------------------------------------------- #
# slide builders
# --------------------------------------------------------------------------- #

def _new_slide(prs, layout_index=6):
    return prs.slides.add_slide(prs.slide_layouts[layout_index])


def _build_cover(prs):
    slide = _new_slide(prs)
    pic = slide.shapes.add_picture(
        str(ASSETS_DIR / "kbc_cover_background.png"), 0, 0, width=_in(SLIDE_W_IN), height=_in(SLIDE_H_IN),
    )
    pic.name = "chrome:cover_background"
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)

    title = _textbox(slide, "field:cover_title", 3.9, 0.85, 9.0, 1.0)
    _set_run(title, "", size=30, color=pptx_theme.THEME.mcm_heading, bold=True, font=pptx_theme.FONT_DISPLAY)
    _rect(slide, "chrome:cover_accent_bar", 3.98, 1.86, 6.8, 0.06, fill=pptx_theme.THEME.mcm_purple_deep, rounded=False)
    subtitle = _textbox(slide, "field:cover_subtitle", 3.98, 2.05, 8.6, 0.4)
    _set_run(subtitle, "", size=13, color=pptx_theme.THEME.mcm_body, bold=True)

    for i, key in enumerate(("field:cover_learner", "field:cover_programme")):
        box = _textbox(slide, key, 0.62, 3.35 + i * 0.42, 6.3, 0.38)
        _set_run(box, "", size=13.5, color=pptx_theme.THEME.mcm_heading, bold=(i == 0))

    meta_labels = ("Employer", "Coach", "Month reviewed")
    for i, (key, label) in enumerate(zip(
        ("field:cover_employer", "field:cover_coach", "field:cover_period"), meta_labels,
    )):
        cx = 0.62 + i * 2.35
        pill = _rect(slide, f"chrome:cover_meta_pill:{i}", cx, 4.55, 2.15, 0.34, fill=(
            pptx_theme.THEME.mcm_purple_deep if i != 1 else pptx_theme.THEME.mcm_gold
        ))
        label_box = _textbox(slide, f"chrome:cover_meta_label:{i}", cx, 4.55, 2.15, 0.34, anchor=MSO_ANCHOR.MIDDLE)
        _set_run(label_box, label, size=10, color=RGBColor.from_string("FFFFFF"), bold=True, align=PP_ALIGN.CENTER)
        value_box = _textbox(slide, key, cx, 4.95, 2.15, 0.36)
        _set_run(value_box, "", size=11.5, color=pptx_theme.THEME.mcm_body, align=PP_ALIGN.CENTER)

    _band(slide, "band:cover", label="THIS MONTH'S FOCUS")
    return slide


def _build_snapshot(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    kpi_y, kpi_h, kpi_w, gap = 1.55, 1.55, 1.68, 0.22
    kpis = [
        ("kpi:attendance", "Attendance", HEADER_ROTATION[0]),
        ("kpi:progress", "Programme progress", HEADER_ROTATION[1]),
        ("kpi:target", "Target", HEADER_ROTATION[2]),
        ("kpi:otj_hours", "OTJ hours", HEADER_ROTATION[0]),
        ("kpi:module", "Key LMS module", HEADER_ROTATION[1]),
    ]
    for i, (name, label, color) in enumerate(kpis):
        _kpi_card(slide, name, 0.6 + i * (kpi_w + gap), kpi_y, kpi_w, kpi_h, header_color=color, header_label=label)

    chip_y = kpi_y + kpi_h + 0.2
    _chip(slide, "otj_status", 0.6, chip_y, 2.0)
    label = _textbox(slide, "field:otj_status_label", 2.75, chip_y, 3.0, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(label, "OTJ status this period", size=10.5, color=pptx_theme.THEME.mcm_body)

    row_y = chip_y + 0.55
    _card(slide, "card:coach_summary", 0.6, row_y, 5.55, 1.7, header_color=HEADER_ROTATION[0], header_label="Coach summary")
    _card(slide, "card:focus", 6.3, row_y, 5.55, 1.7, header_color=HEADER_ROTATION[2], header_label="Immediate focus")

    _band(slide, "band:judgement", label="OVERALL")
    return slide


def _build_attendance(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    row_y, row_h, card_w, gap = 1.55, 2.2, 3.66, 0.2
    for i in range(3):
        name = f"card:month:{i + 1}"
        _card(slide, name, 0.6 + i * (card_w + gap), row_y, card_w, row_h, header_color=HEADER_ROTATION[i], header_label="", body_size=11.5)

    note_y = row_y + row_h + 0.18
    _card(slide, "card:engagement", 0.6, note_y, 5.55, 1.05, header_color=HEADER_ROTATION[1], header_label="Working well")
    _card(slide, "card:future_protocol", 6.3, note_y, 5.55, 1.05, header_color=HEADER_ROTATION[2], header_label="Attendance protocol")

    _band(slide, "band:practical_action", label="PRACTICAL ACTION")
    return slide


def _build_progress_otj_lms(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    bar_y = 1.7
    for i, key in enumerate(("progress", "target")):
        y = bar_y + i * 0.62
        label = _textbox(slide, f"field:{key}_label", 0.6, y, 3.2, 0.3)
        _set_run(label, "", size=11.5, color=pptx_theme.THEME.mcm_heading, bold=True)
        _progress_bar(slide, key, 3.95, y + 0.03, 5.2, 0.24)
        value = _textbox(slide, f"field:{key}_value", 9.3, y, 1.4, 0.3)
        _set_run(value, "", size=11.5, color=pptx_theme.THEME.mcm_heading, bold=True, align=PP_ALIGN.RIGHT)

    _chip(slide, "otj_status", 0.6, 3.1, 2.0)
    otj_value = _textbox(slide, "field:otj_hours_value", 2.75, 3.1, 3.4, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(otj_value, "", size=12, color=pptx_theme.THEME.mcm_body, bold=True)
    otj_variance = _textbox(slide, "field:otj_variance", 6.3, 3.1, 2.9, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    _set_run(otj_variance, "", size=11, color=pptx_theme.THEME.mcm_body)
    otj_warning = _textbox(slide, "field:otj_warning", 0.6, 3.5, 8.6, 0.32)
    _set_run(otj_warning, "", size=10, color=pptx_theme.THEME.mcm_body)

    module_y = 3.95
    module_w = 3.66
    for i in range(3):
        x = 0.6 + i * (module_w + 0.2)
        label = _textbox(slide, f"field:module_label:{i + 1}", x, module_y, module_w, 0.28)
        _set_run(label, "", size=10.5, color=pptx_theme.THEME.mcm_heading, bold=True)
        _progress_bar(slide, f"module:{i + 1}", x, module_y + 0.34, module_w, 0.2)

    _band(slide, "band:action_note", label="ACTION")
    return slide


def _build_assignments(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    row_y, row_h, card_w, gap = 1.55, 3.15, 3.66, 0.2
    for i in range(3):
        name = f"card:evidence:{i + 1}"
        card_label, body_box = _card(slide, name, 0.6 + i * (card_w + gap), row_y, card_w, row_h, header_color=HEADER_ROTATION[i], header_label="", body_size=11.5)
        value_box = _textbox(slide, f"{name}:value", 0.6 + i * (card_w + gap) + 0.16, row_y + row_h - 0.42, card_w - 0.32, 0.32)
        _set_run(value_box, "", size=10, color=HEADER_ROTATION[i], bold=True)

    _band(slide, "band:next_step", label="NEXT STEP")
    return slide


def _build_evidence_photos(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    frame_y, frame_h, frame_w, gap = 1.55, 2.35, 2.7, 0.22
    for i in range(4):
        x = 0.55 + i * (frame_w + gap)
        _photo_frame(slide, f"evidence:{i + 1}", x, frame_y, frame_w, frame_h)

    note = _textbox(slide, "field:evidence_note", 0.55, frame_y + frame_h + 0.42, 11.2, 0.7)
    _set_run(note, "", size=12.5, color=pptx_theme.THEME.mcm_body)

    _band(slide, "band:evidence_value", label="EVIDENCE VALUE")
    return slide


def _build_workplace_impact(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    _card(slide, "card:impact", 0.6, 1.6, 11.25, 3.5, header_color=HEADER_ROTATION[0], header_label="Impact this month", body_size=14)
    _band(slide, "band:impact_note", label="EMPLOYER VALUE")
    return slide


def _build_priority_ksbs(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    card_w, gap, y, h = 2.72, 0.18, 1.55, 3.35
    for i in range(4):
        x = 0.55 + i * (card_w + gap)
        name = f"card:ksb:{i + 1}"
        _rect(slide, f"{name}:card", x, y, card_w, h, fill=RGBColor.from_string("FFFFFF"), line=True)
        code_chip = _rect(slide, f"{name}:code:bg", x + 0.18, y + 0.18, 1.15, 0.42, fill=HEADER_ROTATION[i % 3])
        code_text = _textbox(slide, f"{name}:code:text", x + 0.18, y + 0.18, 1.15, 0.42, anchor=MSO_ANCHOR.MIDDLE)
        _set_run(code_text, "", size=13, color=RGBColor.from_string("FFFFFF"), bold=True, align=PP_ALIGN.CENTER)
        desc = _textbox(slide, f"{name}:description", x + 0.18, y + 0.74, card_w - 0.36, 1.15)
        _set_run(desc, "", size=11, color=pptx_theme.THEME.mcm_heading, bold=True)
        idea = _textbox(slide, f"{name}:idea", x + 0.18, y + 1.95, card_w - 0.36, h - 2.1)
        _set_run(idea, "", size=10.5, color=pptx_theme.THEME.mcm_body)

    _band(slide, "band:decision", label="DECISION TODAY")
    return slide


def _build_smart_targets(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    _card(slide, "card:targets", 0.6, 1.55, 6.85, 3.35, header_color=HEADER_ROTATION[0], header_label="This month's targets", body_size=12)
    _card(slide, "card:owners", 7.65, 1.55, 4.2, 1.55, header_color=HEADER_ROTATION[1], header_label="Owners & dates", body_size=11)
    _card(slide, "card:evidence_upload", 7.65, 3.3, 4.2, 1.6, header_color=HEADER_ROTATION[2], header_label="Evidence to upload", body_size=11)

    _band(slide, "band:decision", label="DECISION TODAY")
    return slide


def _build_wellbeing(prs):
    slide = _new_slide(prs)
    _background(slide, pptx_theme.THEME.mcm_page_bg)
    _logo_mark(slide)
    _header_block(slide, title_key="field:title", subtitle_key="field:subtitle", breadcrumb_key="field:breadcrumb")

    _card(slide, "card:safeguarding", 0.6, 1.55, 5.55, 3.35, header_color=HEADER_ROTATION[0], header_label="Wellbeing & safeguarding", body_size=12.5)
    _card(slide, "card:british_values", 6.3, 1.55, 5.55, 1.55, header_color=HEADER_ROTATION[1], header_label="British Values", body_size=11.5)
    _card(slide, "card:epa_brief", 6.3, 3.3, 5.55, 1.6, header_color=HEADER_ROTATION[2], header_label="EPA brief", body_size=11)

    _band(slide, "band:tip", label="COACH TIP")
    return slide


SLIDE_BUILDERS = (
    _build_cover,
    _build_snapshot,
    _build_attendance,
    _build_progress_otj_lms,
    _build_assignments,
    _build_evidence_photos,
    _build_workplace_impact,
    _build_priority_ksbs,
    _build_smart_targets,
    _build_wellbeing,
)


def build_template() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    for builder in SLIDE_BUILDERS:
        builder(prs)
    return prs


class Command(BaseCommand):
    help = "Rebuild templates/kbc_mcm_template.pptx from the layout constants in this module."

    def handle(self, *args, **options):
        prs = build_template()
        TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(TEMPLATE_PATH))
        self.stdout.write(self.style.SUCCESS(
            f"Wrote {TEMPLATE_PATH} ({len(prs.slides)} slides)."
        ))
