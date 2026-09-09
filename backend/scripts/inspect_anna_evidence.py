"""Read-only Anna evidence inventory and in-memory PDF/DOCX text extraction.

Run from backend, optionally passing evidence ids to inspect their original
files. Does not invoke Django startup hooks or write to the database/storage.
"""
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from inspect_student_activity_pilot import connect


def main():
    with connect("audit") as connection, connection.cursor() as cursor:
        cursor.execute('''
            SELECT evidence_id, evidence_name, component_name, evidence_kind,
                   file_blob, completed_date, submission_date
            FROM fetching_evidence.evidence_items WHERE learner_id = %s
            ORDER BY evidence_id
        ''', [4176])
        rows = cursor.fetchall()
    selected = {int(value) for value in sys.argv[1:]}
    if not selected:
        print(json.dumps([{key: value for key, value in row.items() if key != "file_blob"}
                          for row in rows], default=str, indent=2))
        return
    from learner_api.evidence_storage import _service_client
    service = _service_client(retry_total=1)
    for row in rows:
        if row["evidence_id"] not in selected or not row["file_blob"]:
            continue
        data = service.get_blob_client("fetch-aptem-evidences", row["file_blob"]).download_blob(
            connection_timeout=15, read_timeout=30).readall()
        print(f"\nEVIDENCE {row['evidence_id']}: {row['evidence_name']}")
        if data.startswith(b"%PDF"):
            import fitz
            with fitz.open(stream=data, filetype="pdf") as pdf:
                for index, page in enumerate(pdf, 1):
                    print(f"PAGE {index}\n{page.get_text()}")
        elif data.startswith(b"PK"):
            from docx import Document
            document = Document(io.BytesIO(data))
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    print(paragraph.text)
            for table in document.tables:
                for table_row in table.rows:
                    print(" | ".join(dict.fromkeys(cell.text for cell in table_row.cells)))
        else:
            print("Unsupported file format; needs separate review.")


if __name__ == "__main__":
    main()
