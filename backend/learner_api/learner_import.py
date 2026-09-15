"""Staff-only Excel enrolment: validate the whole workbook before any writes."""
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import DatabaseError, connections, transaction
from django.db.models.functions import Lower, Trim
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from login.models import LoginAccount
from login.permissions import staff_only
from .mappers import ValidationError, write_fields
from .models import Employer, EnrolmentUser, LearnerProfile, StaffUser

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_ROWS = 500
MAX_SHEET_ROWS = 2000
COLUMNS = (
    ("First name", "firstName"), ("Surname", "surname"), ("Email", "email"),
    ("Mobile", "phone"), ("Date of birth", "dob"), ("Programme", "programme"),
    ("Cohort", "cohort"), ("Group", "group"), ("Employer ID", "employerId"),
    ("Case owner", "caseOwner"), ("Preferred name", "preferredName"),
    ("Title", "title"), ("Gender", "gender"), ("National Insurance number", "niNumber"),
    ("Postcode", "postcode"), ("Address 1", "addressLine1"), ("Address 2", "addressLine2"),
    ("Town/City", "townCity"), ("County", "county"), ("Country", "country"),
    ("Learning provider", "learningProvider"), ("Manager", "lineManager"),
    ("Mentor", "mentor"), ("Reference number", "referenceNumber"),
    ("Referrer", "referrer"), ("Referrer address", "referrerAddress"),
    ("Referrer contact", "referrerContact"), ("Employer address", "employerAddress"),
    ("Extended break", "extendedBreak"),
)
REQUIRED = {"firstName", "surname", "email"}
LABELS = {key: label for label, key in COLUMNS}


def _text(value):
    return "" if value is None else str(value).strip()


def _key(value):
    return _text(value).casefold()


def _header(value):
    return re.sub(r"[\s_*/-]+", "", _key(value))


def _issue(row, message, field=None):
    result = {"row": row, "message": message}
    if field:
        result["field"] = LABELS.get(field, field)
    return result


def _body(count=0, preview=None, errors=None, results=None):
    return {"count": count, "preview": preview or [], "errors": errors or [],
            "results": results or [], "imported": len(results or [])}


def _error(message, *, errors=None, count=0, preview=None, status=400):
    return JsonResponse({**_body(count, preview, errors or [_issue(0, message)]),
                         "error": message}, status=status)


def read_students(upload):
    """Return (row-number, payload) pairs and cell/header errors; never hit DB."""
    if Path(upload.name).suffix.lower() != ".xlsx":
        raise ValidationError("Upload an Excel .xlsx file downloaded from the template.")
    if upload.size > MAX_FILE_BYTES:
        raise ValidationError("The file must be 5 MB or smaller.")
    data = upload.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValidationError("The file must be 5 MB or smaller.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValidationError("The workbook expands beyond the supported size. Use a fresh template.")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ValidationError:
        raise
    except (BadZipFile, KeyError, ValueError, OSError, SyntaxError) as exc:
        raise ValidationError("This file is not a readable Excel workbook. Download a fresh template.") from exc

    try:
        if "Students" not in workbook.sheetnames:
            raise ValidationError("The workbook must contain a sheet named Students.")
        sheet = workbook["Students"]
        if (sheet.max_row or 0) > MAX_SHEET_ROWS or (sheet.max_column or 0) > len(COLUMNS):
            raise ValidationError("The Students sheet is too large. Use the template and at most 500 students.")
        # Some spreadsheet writers under-report dimensions. Read the actual
        # rows so a misleading dimension cannot silently omit students.
        sheet.reset_dimensions()
        aliases = {_header(label): key for label, key in COLUMNS}
        aliases.update({_header(key): key for _, key in COLUMNS})
        columns, errors = [], []
        for cell in next(sheet.iter_rows(min_row=1, max_row=1), []):
            label = _text(cell.value)
            key = aliases.get(_header(label)) if label else None
            if label and key is None:
                errors.append(_issue(1, f"Unknown column: {label}. Keep the template headings."))
            if key and key in columns:
                errors.append(_issue(1, f"Duplicate column: {label}."))
            columns.append(key)
        for missing in sorted(REQUIRED - set(columns)):
            errors.append(_issue(1, f"Missing required column: {LABELS[missing]}.", missing))
        if errors:
            return [], errors
        students = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if row_number > MAX_SHEET_ROWS or len(cells) > len(COLUMNS):
                raise ValidationError("The Students sheet is too large. Use the template and at most 500 students.")
            if not any(cell.value is not None and _text(cell.value) for cell in cells):
                continue
            if len(students) == MAX_ROWS:
                raise ValidationError("Import at most 500 students at a time.")
            payload = {}
            for index, cell in enumerate(cells):
                key = columns[index] if index < len(columns) else None
                value = cell.value
                if value is None or not _text(value):
                    continue
                if not key:
                    errors.append(_issue(row_number, "A value has no column heading."))
                    continue
                if cell.data_type in {"f", "e"}:
                    errors.append(_issue(row_number, "Use a plain value, not an Excel formula or error.", key))
                    continue
                if isinstance(value, (datetime, date)):
                    if key != "dob":
                        errors.append(_issue(row_number, "Use text in this column.", key))
                        continue
                    value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                elif key == "employerId" and isinstance(value, (float, int)) and not isinstance(value, bool):
                    value = str(int(value)) if int(value) == value else str(value)
                elif key in {"phone", "niNumber", "postcode", "referenceNumber"} and not isinstance(value, str):
                    errors.append(_issue(row_number, "Enter this value as text to preserve leading zeros.", key))
                    continue
                value = _text(value)
                if len(value) > 2000:
                    errors.append(_issue(row_number, "The value is too long (maximum 2000 characters).", key))
                else:
                    payload[key] = value
            students.append((row_number, payload))
        if not students:
            errors.append(_issue(2, "The Students sheet is empty. Add at least one student below the headings."))
        return students, errors
    except (BadZipFile, KeyError, ValueError, TypeError, OSError, SyntaxError, OverflowError) as exc:
        raise ValidationError("This file contains unreadable Excel cells. Download a fresh template.") from exc
    finally:
        workbook.close()


def load_references():
    """Read authored choices without invoking curriculum's setup/write paths."""
    from curriculum_api.views import curriculum_row_effectively_deleted, is_archived_program_config

    connection = connections["enrolment"]
    rows = {}
    for table in ("programmes", "cohorts", "groups"):
        table_name = f'curriculum."{table}"' if connection.vendor == "postgresql" else f'"{table}"'
        with connection.cursor() as cursor:
            cursor.execute(f"SELECT * FROM {table_name}")
            names = [column[0] for column in cursor.description]
            rows[table] = [dict(zip(names, row)) for row in cursor.fetchall()]
        rows[table] = [row for row in rows[table] if not curriculum_row_effectively_deleted(row)
                       and (table != "programmes" or not is_archived_program_config(row))]
    rows["employers"] = list(Employer.objects.values(
        "id", "first_name", "surname", "employer_group_names"))
    rows["owners"] = list(StaffUser.objects.filter(position__in=["Caseowner", "Admin"])
                          .values("username", "email"))
    for programme in rows["programmes"]:
        programme["name"] = _text(programme.get("name")) or _text(programme.get("programme_id"))
    return rows


def _matches_parent(row, parent, kind):
    parent_id = _text(parent.get(f"{kind}_id") or parent.get("id"))
    row_id = _text(row.get(f"{kind}_id"))
    if row_id and parent_id:
        return row_id == parent_id
    return _key(row.get(f"{kind}_name")) == _key(parent.get("name") or parent.get(f"{kind}_name"))


def _resolve_placement(payload, references, row_number):
    errors, programme, cohort = [], None, None
    for kind, table, parent in (("programme", "programmes", None),
                                 ("cohort", "cohorts", "programme"),
                                 ("group", "groups", "cohort")):
        value = payload.get(kind)
        if not value:
            continue
        if parent and not payload.get(parent):
            errors.append(_issue(row_number, f"Select a {parent} before a {kind}.", kind))
            continue
        matches = [row for row in references[table]
                   if _key(row.get("name") if kind == "programme" else row.get(f"{kind}_name")) == _key(value)]
        if parent:
            parent_row = programme if parent == "programme" else cohort
            matches = [row for row in matches if parent_row and _matches_parent(row, parent_row, parent)]
        if kind == "group":
            matches = [row for row in matches if programme and _matches_parent(row, programme, "programme")]
        if len(matches) != 1:
            message = f"Unknown {kind} or incorrect parent selection. Use the Placements sheet."
            if len(matches) > 1:
                message = f"This {kind} name is ambiguous. Ask an administrator to make its name unique."
            errors.append(_issue(row_number, message, kind))
            continue
        selected = matches[0]
        payload[kind] = _text(selected.get("name") if kind == "programme" else selected.get(f"{kind}_name"))
        if kind == "programme":
            programme = selected
        elif kind == "cohort":
            cohort = selected
    return errors


def existing_emails(emails):
    found = set()
    # Login spans every subject type: reusing a staff/employer login email would
    # make the email-only login form ambiguous even if no learner exists yet.
    for manager in (EnrolmentUser.all_learners, LearnerProfile.objects, LoginAccount.objects):
        found.update(manager.annotate(import_email=Lower(Trim("email")))
                     .filter(import_email__in=emails).values_list("import_email", flat=True))
    return found


def validate_students(students, references):
    errors, prepared, preview, seen = [], [], [], {}
    emails = {_key(payload.get("email")) for _, payload in students if payload.get("email")}
    existing = existing_emails(emails) if emails else set()
    employers = {str(row["id"]): row for row in references["employers"]}
    for row_number, original in students:
        payload = dict(original)
        for required in sorted(REQUIRED):
            if not payload.get(required):
                errors.append(_issue(row_number, f"{LABELS[required]} is required.", required))
        email = _key(payload.get("email"))
        if email:
            try:
                validate_email(email)
            except DjangoValidationError:
                errors.append(_issue(row_number, "Enter a valid email address.", "email"))
            if email in seen:
                errors.append(_issue(row_number, f"This email is repeated on row {seen[email]}.", "email"))
            elif email in existing:
                errors.append(_issue(row_number, "This email already belongs to a learner or platform account.", "email"))
            seen[email] = row_number
        if payload.get("dob"):
            try:
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", payload["dob"]):
                    raise ValueError
                birthday = date.fromisoformat(payload["dob"])
                if birthday > timezone.localdate():
                    raise ValueError
            except ValueError:
                errors.append(_issue(row_number, "Use a valid date of birth in YYYY-MM-DD format, no later than today.", "dob"))
        errors.extend(_resolve_placement(payload, references, row_number))
        if payload.get("employerId"):
            employer = employers.get(payload["employerId"])
            if not employer:
                errors.append(_issue(row_number, "Unknown employer ID. Use the Employers sheet.", "employerId"))
            else:
                payload["employer"] = " ".join(filter(None, [_text(employer["first_name"]), _text(employer["surname"])]))
                groups = employer.get("employer_group_names")
                payload["organization"] = ", ".join(map(str, groups)) if isinstance(groups, list) else ""
        if payload.get("caseOwner"):
            owners = [row for row in references["owners"] if _key(row["username"]) == _key(payload["caseOwner"])]
            if len(owners) != 1:
                errors.append(_issue(row_number, "Unknown or ambiguous case owner. Use the Case owners sheet.", "caseOwner"))
            else:
                payload["caseOwner"] = owners[0]["username"]
        payload.update(username=" ".join(filter(None, [payload.get("firstName"), payload.get("surname")])),
                       email=email, learnerType="commercial", type="User", status="FullUser")
        payload.setdefault("country", "United Kingdom")
        payload.setdefault("learningProvider", "Kent Business College")
        preview.append({"row": row_number, "name": payload["username"], "email": email,
                        **{field: payload.get(field, "") for field in ("programme", "cohort", "group")}})
        try:
            prepared.append((row_number, write_fields(payload, require_create=True)))
        except ValidationError as exc:
            # Required cells already have precise column errors above.
            if REQUIRED <= {key for key, value in payload.items() if value}:
                errors.append(_issue(row_number, str(exc)))
    return prepared, preview, errors


def _sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for values in rows:
        sheet.append([_text(value) for value in values])
        # Treat reference data as literal text, including values starting '='.
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="16384F")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, heading in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = max(20, min(42, len(heading) + 8))
    return sheet


def template_workbook(references):
    workbook = Workbook()
    workbook.remove(workbook.active)
    students = _sheet(workbook, "Students", [label for label, _ in COLUMNS], [])
    for index, (_, key) in enumerate(COLUMNS, 1):
        students.cell(1, index).comment = Comment(
            "Required. Enter one student's value per row." if key in REQUIRED else "Optional. See Instructions.", "LMS")
        for row_number in range(2, MAX_ROWS + 2):
            students.cell(row_number, index).number_format = "@"
    _sheet(workbook, "Instructions", ["Step", "Instructions"], [
        ("1", "Fill the Students sheet, beginning on row 2. Leave the headings unchanged. One student per row."),
        ("2", "First name, Surname and Email are required. All other columns are optional. Delete unused blank rows if desired."),
        ("3", "Maximum 500 students and 5 MB per .xlsx file. Do not enter formulas. Phone, postcode, NI and reference values must be text."),
        ("4", "Dates of birth must use YYYY-MM-DD (for example 2000-04-25). Keep mobile numbers as text to preserve leading zeros."),
        ("5", "Copy Programme, Cohort and Group names from Placements. Cohort requires Programme; Group requires both."),
        ("6", "Optional Employer ID and Case owner must match the Employers and Case owners sheets. Organisation follows the employer."),
        ("7", "Upload for preview, correct any listed errors, then import. If any row is invalid, no students are saved."),
        ("8", "Existing learner or platform-account emails and repeated emails are rejected. Existing records are never updated."),
        ("9", "New students are commercial learners (Type User, subscription FullUser). Country defaults to United Kingdom and provider to Kent Business College."),
        ("10", "Cohort dates and the case owner's coach details follow the normal Add user flow. Accounts are created without sending emails; send invitations from Accounts when ready."),
    ])
    workbook["Instructions"].column_dimensions["B"].width = 135
    placements = []
    for programme in references["programmes"]:
        name = _text(programme.get("name"))
        placements.append((name, "", ""))
        for cohort in references["cohorts"]:
            if not _matches_parent(cohort, programme, "programme"):
                continue
            placements.append((name, _text(cohort.get("cohort_name")), ""))
            for group in references["groups"]:
                if _matches_parent(group, cohort, "cohort") and _matches_parent(group, programme, "programme"):
                    placements.append((name, _text(cohort.get("cohort_name")), _text(group.get("group_name"))))
    _sheet(workbook, "Placements", ["Programme", "Cohort", "Group"], sorted(set(placements)))
    _sheet(workbook, "Employers", ["Employer ID", "Employer", "Organisation"], [
        (row["id"], f'{_text(row["first_name"])} {_text(row["surname"])}'.strip(),
         ", ".join(map(str, row["employer_group_names"])) if isinstance(row["employer_group_names"], list) else "")
        for row in references["employers"]])
    _sheet(workbook, "Case owners", ["Case owner", "Email"], [
        (row["username"], row["email"]) for row in references["owners"]])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@require_GET
@staff_only()
def import_template(request):
    try:
        content = template_workbook(load_references())
    except DatabaseError:
        return _error("The template choices could not be loaded. Try again shortly.", status=503)
    response = HttpResponse(content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="student-import-template.xlsx"'
    response["Cache-Control"] = "private, no-store"
    return response


@csrf_exempt
@require_POST
@staff_only()
def import_students(request):
    if request.headers.get("X-Requested-With") != "XMLHttpRequest":
        return _error("Missing X-Requested-With header.", status=403)
    files = request.FILES.getlist("file")
    if len(files) != 1:
        return _error("Upload exactly one Excel file using the file field.")
    dry_run = request.POST.get("dryRun", "true").lower()
    if dry_run not in {"true", "false"}:
        return _error("dryRun must be true or false.")
    try:
        students, errors = read_students(files[0])
        if errors:
            return _error("Correct the workbook errors and upload it again.", errors=errors, count=len(students))
        prepared, preview, errors = validate_students(students, load_references())
        if errors:
            return _error("Correct the workbook errors and upload it again.", errors=errors, count=len(students), preview=preview)
        if dry_run == "true":
            return JsonResponse(_body(len(students), preview))

        from .views import _create_enrolment_user

        results, current_row = [], 0
        try:
            with transaction.atomic(using="enrolment"):
                # Every commit revalidates the upload. A login uniqueness race
                # also fails account provisioning and rolls the entire batch back.
                for current_row, fields in prepared:
                    results.append(_create_enrolment_user(request, fields, require_account=True))
        except ValidationError as exc:
            return _error(str(exc), errors=[_issue(current_row, str(exc))], count=len(students), preview=preview)
        return JsonResponse(_body(len(students), preview, results=results), status=201)
    except ValidationError as exc:
        return _error(str(exc))
    except DatabaseError:
        return _error("The import could not be completed. No students were imported. Try again shortly.", status=503)
