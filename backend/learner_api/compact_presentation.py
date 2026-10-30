"""Opt-in MCM format: one editable PowerPoint slide per saved section.

Legacy decks and uploaded files keep their existing renderers. No database,
network, learner-data mutation or automatic pagination occurs here.
"""
import io
import textwrap
from copy import deepcopy
from pathlib import Path

FORMAT = 'kbc-mcm-v1'
ASSETS = Path(__file__).parent / 'assets' / 'mcm-v1'
PURPLE, PLUM, GOLD, LAVENDER, INK = '51258B', '24103E', '8A5D10', 'F0EAFB', '30263E'
FONT = 'Calibri'


def _validate(slides):
    if not isinstance(slides, list) or not 1 <= len(slides) <= 200:
        raise ValueError('Add between 1 and 200 slides before exporting.')
    for number, section in enumerate(slides, 1):
        if not isinstance(section, dict) or any(not isinstance(section.get(key), str) or not section[key].strip() for key in ('title', 'body')):
            raise ValueError(f'Slide {number}: add a title and content before exporting.')
        if len(section['title']) > 80:
            raise ValueError(f'Slide {number}: keep the title within 80 characters.')
        if len(section['body']) > 600 or len(section['body'].split('\n')) > 6:
            raise ValueError(f'Slide {number}: keep the content to 6 short lines and 600 characters. Move detail to speaker notes or add a slide.')
        if not isinstance(section.get('notes', ''), str) or len(section.get('notes', '')) > 2000000:
            raise ValueError(f'Slide {number}: speaker notes must be text of no more than 2,000,000 characters.')
    if sum(len(s.get('notes', '')) for s in slides) > 2000000:
        raise ValueError('The speaker notes are too large. Split this material into separate presentations.')


def _write(frame, text, size, width, height, number, color=None, bold=False, font=FONT, preserve_style=False):
    """Conservative wrapping; never shrink, hide or discard overflowing text."""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    width_pt, height_pt = width / 12700, height / 12700
    columns = max(1, int((width_pt - 6) / (size * .60)))
    paragraphs = [textwrap.wrap(p, columns, break_long_words=True, break_on_hyphens=False) or [''] for p in text.split('\n')]
    line_height = size * 1.16
    required = sum(len(p) for p in paragraphs) * line_height + max(0, len(paragraphs) - 1) * 9
    if required > height_pt - 6:
        raise ValueError(f'Slide {number}: this text does not fit the selected layout at a readable size. Shorten the highlights, move detail to speaker notes, or add a slide.')
    first = frame.paragraphs[0]
    alignment = first.alignment if preserve_style else None
    style = deepcopy(first.runs[0]._r.rPr) if preserve_style and first.runs and first.runs[0]._r.rPr is not None else None
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = Pt(0)
    for index, lines in enumerate(paragraphs):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = alignment
        p.text = '\v'.join(lines)
        for run in p.runs:
            if style is not None:
                run._r.insert(0, deepcopy(style))
            run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        p.font.name, p.font.size, p.font.bold = font, Pt(size), bold
        if color:
            p.font.color.rgb = RGBColor.from_string(color)
        p.space_before, p.space_after, p.line_spacing = Pt(0), Pt(9), Pt(line_height)


def _text(slide, text, x, y, width, height, size, number, color, bold=False):
    from pptx.util import Inches
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    _write(shape.text_frame, text, size, shape.width, shape.height, number, color, bold)
    return shape


def _logo(slide, x, y, width):
    """Native picture crop removes transparent padding; original PNG stays intact."""
    from pptx.util import Inches
    picture = slide.shapes.add_picture(str(ASSETS / 'kbc-original-logo.png'), Inches(x), Inches(y),
                                       width=Inches(width), height=Inches(width * 800 / 779))
    picture.crop_left, picture.crop_top = 394 / 1563, 437 / 1563
    picture.crop_right, picture.crop_bottom = 390 / 1563, 326 / 1563
    return picture


def _notes(slide, section):
    # The visible highlights and the full account both remain editable and readable.
    if slide.notes_slide.notes_text_frame is None:
        from pptx.oxml.shapes.autoshape import CT_Shape
        from pptx.enum.shapes import PP_PLACEHOLDER
        notes = slide.notes_slide
        shape_id = max((s.shape_id for s in notes.shapes), default=1) + 1
        notes.shapes._spTree.insert_element_before(
            CT_Shape.new_placeholder_sp(shape_id, 'Speaker notes', PP_PLACEHOLDER.BODY, 'horz', 'full', 1), 'p:extLst')
    slide.notes_slide.notes_text_frame.text = f"{section['title']}\n\nSlide highlights\n{section['body']}\n\nFull detail\n{section.get('notes') or section['body']}"


def build_compact_deck(slides, reference=None, design=None):
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    _validate(slides)
    if reference is not None:
        return _from_template(slides, reference, design or {})
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = 'MCM Meeting'
    deck.core_properties.subject = 'Monthly learning, reflection and workplace impact'
    deck.core_properties.author = 'Kent Business College'
    for index, section in enumerate(slides):
        number = index + 1
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        cover = index == 0
        closing = section['title'] == 'EPA preparation'
        dark = not cover and (closing or section['title'] == 'Next-month action plan')
        background = PLUM if dark else LAVENDER if cover or index % 4 == 0 else 'FFFFFF'
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        if cover:
            _logo(slide, .72, .45, 1.02)
            slide.shapes.add_picture(str(ASSETS / 'kbc-ibis-wreath.png'), Inches(7.4), Inches(.95), height=Inches(5.85))
            _text(slide, section['title'], .75, 2.25, 6.4, 1.5, 46, number, PURPLE, True)
            _text(slide, section['body'], .8, 4.0, 6.4, 2.4, 22, number, INK)
        else:
            if dark:
                _text(slide, 'Kent Business College', .75, .42, 9.5, .4, 17, number, 'DAC3FA')
            else:
                _logo(slide, 11.65, .4, .82)
            _text(slide, section['title'], .75, 1.05, 10.6, 1.25, 34, number, 'FFFFFF' if dark else PURPLE, True)
            if closing:
                slide.shapes.add_picture(str(ASSETS / 'kbc-horse-growth.png'), Inches(9.1), Inches(2.65), width=Inches(3.6))
            _text(slide, section['body'], .8, 2.75, 8.05 if closing else 11.5, 3.9, 22 if closing else 24, number, 'F0EAFB' if dark else INK)
        footer_color = 'DAC3FA' if dark else GOLD
        _text(slide, 'KBC · Monthly learning', .8, 6.98, 9.5, .35, 11, number, footer_color)
        _text(slide, f'{number:02d} / {len(slides):02d}', 11.5, 6.98, 1, .35, 11, number, footer_color)
        _notes(slide, section)
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


def _from_template(slides, content, design):
    """Preserve an explicitly selected reference, including its slide dimensions."""
    from pptx import Presentation
    from pptx.util import Inches
    from .presentation_template import _clone_slide, _all_shapes
    deck = Presentation(io.BytesIO(content))
    originals = list(deck.slides)
    try:
        cover = int(design.get('coverSlide', 1)) - 1
        body = int(design.get('contentSlide', min(2, len(originals)))) - 1
        if not 0 <= cover < len(originals) or not 0 <= body < len(originals):
            raise ValueError
    except (ValueError, TypeError):
        raise ValueError('Choose valid cover and content slide numbers for your reference.') from None
    width, height = deck.slide_width, deck.slide_height
    for index, section in enumerate(slides):
        slide = _clone_slide(deck, originals[cover if index == 0 else body])
        shapes = [s for s in _all_shapes(slide.shapes) if s.has_text_frame and (s.text.strip() or s.is_placeholder)]
        title = slide.shapes.title
        if title is None:
            title = min(shapes, key=lambda s: s.top) if shapes else None
        candidates = [s for s in shapes if s.shape_id != getattr(title, 'shape_id', None) and s.width >= width * .35 and s.height >= height * .15]
        body_shape = max(candidates, key=lambda s: s.width * s.height) if candidates else None
        for shape in shapes:
            if shape.shape_id not in (getattr(title, 'shape_id', None), getattr(body_shape, 'shape_id', None)):
                shape.text_frame.clear()
        if title is None:
            title = slide.shapes.add_textbox(int(width * .08), int(height * .08), int(width * .78), int(height * .18))
        if body_shape is None:
            body_shape = slide.shapes.add_textbox(int(width * .08), int(height * .32), int(width * .82), int(height * .54))
            body_shape.fill.solid()
            from pptx.dml.color import RGBColor
            body_shape.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
        font = str(design.get('font') or FONT)[:80]
        _write(title.text_frame, section['title'], 32, title.width, title.height, index + 1, bold=True, font=font, preserve_style=True)
        _write(body_shape.text_frame, section['body'], 20, body_shape.width, body_shape.height, index + 1, font=font, preserve_style=True)
        _logo(slide, (width - Inches(.95)) / 914400, (height - Inches(.7)) / 914400, .55)
        _notes(slide, section)
    # Private reference slide text/notes must not be published as learner work.
    for _ in originals:
        slide_id = deck.slides._sldIdLst[0]
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    deck.core_properties.title = 'MCM Meeting'
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()
