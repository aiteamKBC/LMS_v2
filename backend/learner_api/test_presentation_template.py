import io
from pathlib import Path
from unittest import TestCase
from unittest.mock import MagicMock, patch

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from .presentation_template import build_from_template, load_owned_template


class PresentationTemplateTests(TestCase):
    def reference(self):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13), Inches(7)
        for i in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = f'Original title {i}'
            slide.placeholders[1].text = f'Original body {i}'
            triangle = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, 0, Inches(5), Inches(2), Inches(2))
            triangle.name = f'Reference artwork {i}'
            triangle.fill.solid()
            triangle.fill.fore_color.rgb = RGBColor(210, 0, 100)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(245, 240, 230)
            slide.notes_slide.notes_text_frame.text = 'PRIVATE ORIGINAL NOTES'
            slide.shapes.add_picture(str(Path(__file__).parent / 'assets' / 'kbc-logo.png'), Inches(11), 0, width=Inches(1))
        stream = io.BytesIO()
        deck.save(stream)
        return stream.getvalue()

    def test_retains_package_layout_artwork_images_and_replaces_original_content(self):
        reference = self.reference()
        result = Presentation(io.BytesIO(build_from_template([
            {'title': 'New cover', 'body': 'The learner answer'},
            {'title': 'New content', 'body': 'What the learner learned'}], reference, {'coverSlide': 1, 'contentSlide': 2})))
        source = Presentation(io.BytesIO(reference))
        self.assertEqual(result.slide_width, source.slide_width)
        self.assertEqual(result.slide_height, source.slide_height)
        self.assertEqual(len(result.slide_masters), len(source.slide_masters))
        self.assertEqual(len(result.slides), 3)
        self.assertEqual(result.slides[0].shapes.title.text, "MCM Meeting")
        self.assertEqual(result.core_properties.title, "MCM Meeting")
        for i, slide in enumerate(result.slides):
            self.assertTrue(any(s.name == f'Reference artwork {0 if i == 0 else 1}' for s in slide.shapes))
            self.assertEqual(slide.background.fill.fore_color.rgb, RGBColor(245, 240, 230))
            self.assertEqual(sum(s.shape_type == 13 for s in slide.shapes), 2)
            self.assertFalse(slide.has_notes_slide)
            text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            self.assertNotIn('Original', text)
            if i > 0:
                self.assertIn('learner', text)

    def test_paginates_without_keeping_original_slides(self):
        body = ' '.join(f'word{i}' for i in range(900))
        result = Presentation(io.BytesIO(build_from_template([{'title': 'New', 'body': body}], self.reference(), {})))
        self.assertGreater(len(result.slides), 2)
        all_text = '\n'.join(s.text for slide in result.slides for s in slide.shapes if s.has_text_frame)
        for i in range(900):
            self.assertIn(f'word{i}', all_text)
        self.assertNotIn('Original body', all_text)

    def test_requires_approved_file_owned_by_learner(self):
        cursor = MagicMock()
        cursor.fetchall.return_value = []
        cursor.fetchone.return_value = None
        conn = MagicMock()
        conn.cursor.return_value.__enter__.return_value = cursor
        with patch('learner_api.presentation_template.connections', {'enrolment': conn}), patch('learner_api.evidence_storage.download_blob_bytes') as download:
            with self.assertRaises(ValueError):
                load_owned_template('commercial', '123', 'other-file')
        self.assertEqual(cursor.execute.call_args.args[1], ['other-file', 'commercial', '123', 'approved'])
        download.assert_not_called()
