"""Excel validation and all-or-nothing import orchestration; no database access."""
from datetime import date
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError
from django.test import RequestFactory, SimpleTestCase
from openpyxl import Workbook, load_workbook

from . import learner_import as imports, views
from .mappers import ValidationError
from .models import EnrolmentUser


def references():
    return {
        "programmes": [{"programme_id": "p1", "name": "Business"},
                       {"programme_id": "p2", "name": "Computing"}],
        "cohorts": [{"cohort_id": "c1", "programme_id": "p1", "programme_name": "Business", "cohort_name": "September"},
                    {"cohort_id": "c2", "programme_id": "p2", "programme_name": "Computing", "cohort_name": "October"}],
        "groups": [{"group_id": "g1", "programme_id": "p1", "programme_name": "Business", "cohort_id": "c1", "cohort_name": "September", "group_name": "A"}],
        "employers": [{"id": 42, "first_name": "Test", "surname": "Employer", "employer_group_names": ["Company"]}],
        "owners": [{"username": "Test Coach", "email": "coach@example.test"}],
    }


def student(**overrides):
    return {"firstName": "Test", "surname": "Learner", "email": "student@example.test", **overrides}


def upload(rows=None, headers=None, sheet_name="Students"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers or ["First name", "Surname", "Email"])
    for row in rows if rows is not None else [["Test", "Learner", "student@example.test"]]:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return SimpleUploadedFile("students.xlsx", output.getvalue())


class WorkbookParsingTests(SimpleTestCase):
    def test_corrupt_numeric_cell_is_a_validation_error(self):
        original = upload([["Test", "Learner", "student@example.test", 123]],
                          ["First name", "Surname", "Email", "Mobile"])
        output = BytesIO()
        with ZipFile(BytesIO(original.read())) as source, ZipFile(output, "w") as target:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace(b"<v>123</v>", b"<v>broken</v>")
                target.writestr(member, content)
        with self.assertRaisesMessage(ValidationError, "unreadable Excel cells"):
            imports.read_students(SimpleUploadedFile("students.xlsx", output.getvalue()))

    def test_underreported_sheet_dimensions_do_not_drop_students(self):
        original = upload([["A", "B", "a@example.test"], ["C", "D", "b@example.test"]])
        output = BytesIO()
        with ZipFile(BytesIO(original.read())) as source, ZipFile(output, "w") as target:
            for member in source.infolist():
                content = source.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace(b'dimension ref="A1:C3"', b'dimension ref="A1:C2"')
                target.writestr(member, content)
        students, errors = imports.read_students(SimpleUploadedFile("students.xlsx", output.getvalue()))
        self.assertFalse(errors)
        self.assertEqual([row for row, _ in students], [2, 3])

    def test_template_has_blank_text_cells_and_round_trips_real_values(self):
        content = imports.template_workbook(references())
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["Students", "Instructions", "Placements", "Employers", "Case owners"])
        sheet = workbook["Students"]
        self.assertEqual(sheet["A2"].value, None)
        self.assertEqual(sheet["D2"].number_format, "@")
        sheet["A2"], sheet["B2"], sheet["C2"] = "Test", "Learner", "student@example.test"
        sheet["D2"] = "00123456789"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        rows, errors = imports.read_students(SimpleUploadedFile("students.xlsx", output.getvalue()))
        self.assertFalse(errors)
        self.assertEqual(rows[0][1]["phone"], "00123456789")
        self.assertEqual(len(rows), 1)

    def test_header_reordering_and_blank_rows_preserve_excel_row_numbers(self):
        rows, errors = imports.read_students(upload(
            [[None, None, None], [" Learner ", " Test ", "student@example.test"]],
            ["Surname", "firstName", " EMAIL "]))
        self.assertFalse(errors)
        self.assertEqual(rows, [(3, student())])

    def test_rejects_wrong_headers_missing_headers_duplicate_headers_and_empty_sheet(self):
        for headers, values, expected in (
            (["First name", "Surname", "Mystery"], [["A", "B", "C"]], "Unknown column"),
            (["First name", "Surname"], [["A", "B"]], "Missing required column"),
            (["First name", "Surname", "Email", "email"], [["A", "B", "a@b.test", "a@b.test"]], "Duplicate column"),
            (["First name", "Surname", "Email"], [], "empty"),
        ):
            with self.subTest(expected=expected):
                _, errors = imports.read_students(upload(values, headers))
                self.assertTrue(any(expected in item["message"] for item in errors))

    def test_rejects_formula_and_numeric_phone_without_silently_changing_them(self):
        rows, errors = imports.read_students(upload(
            [["=1+1", "Learner", "student@example.test", 12345]],
            ["First name", "Surname", "Email", "Mobile"]))
        self.assertEqual(len(rows), 1)
        self.assertEqual({error["field"] for error in errors}, {"First name", "Mobile"})

    def test_accepts_excel_date_cells_and_keeps_text_leading_zeros(self):
        rows, errors = imports.read_students(upload(
            [["Test", "Learner", "student@example.test", date(2000, 4, 25), "001234"]],
            ["First name", "Surname", "Email", "Date of birth", "Mobile"]))
        self.assertFalse(errors)
        self.assertEqual(rows[0][1]["dob"], "2000-04-25")
        self.assertEqual(rows[0][1]["phone"], "001234")

    def test_rejects_bad_file_and_missing_sheet(self):
        for file in (SimpleUploadedFile("wrong.csv", b"a,b"),
                     SimpleUploadedFile("wrong.xlsx", b"not a workbook"),
                     upload(sheet_name="Sheet1")):
            with self.subTest(name=file.name), self.assertRaises(ValidationError):
                imports.read_students(file)

    def test_limits_compressed_and_expanded_file_sizes(self):
        oversized = upload()
        oversized.size = imports.MAX_FILE_BYTES + 1
        with self.assertRaisesMessage(ValidationError, "5 MB"):
            imports.read_students(oversized)
        with patch.object(imports, "MAX_UNCOMPRESSED_BYTES", 1), \
             self.assertRaisesMessage(ValidationError, "expands"):
            imports.read_students(upload())

    def test_row_limit_counts_students(self):
        with patch.object(imports, "MAX_ROWS", 1), self.assertRaisesMessage(ValidationError, "500"):
            imports.read_students(upload([["A", "B", "a@example.test"], [None, None, None], ["C", "D", "b@example.test"]]))

    def test_reference_values_cannot_become_excel_formulas(self):
        refs = references()
        refs["owners"] = [{"username": "=HYPERLINK(\"https://example.test\")", "email": "coach@example.test"}]
        workbook = load_workbook(BytesIO(imports.template_workbook(refs)))
        self.assertEqual(workbook["Case owners"]["A2"].data_type, "s")
        workbook.close()


class ReferenceLoadingTests(SimpleTestCase):
    def test_references_use_enrolment_database_and_current_case_owner_roles(self):
        database = MagicMock(vendor="postgresql")
        cursor = database.cursor.return_value.__enter__.return_value
        cursor.description = [("programme_id",), ("name",)]
        cursor.fetchall.side_effect = [[("p1", "")], [], []]
        with patch.object(imports, "connections", {"enrolment": database}), \
             patch.object(imports.Employer.objects, "values", return_value=[]), \
             patch.object(imports.StaffUser.objects, "filter") as staff:
            staff.return_value.values.return_value = [{"username": "Coach", "email": "coach@example.test"}]
            result = imports.load_references()
        self.assertEqual(result["programmes"][0]["name"], "p1")
        self.assertEqual(result["owners"][0]["username"], "Coach")
        staff.assert_called_once_with(position__in=["Caseowner", "Admin"])
        self.assertEqual(cursor.execute.call_count, 3)
        self.assertTrue(all('curriculum.' in call.args[0] for call in cursor.execute.call_args_list))


class StudentValidationTests(SimpleTestCase):
    def validate(self, rows, *, existing=None, refs=None):
        with patch.object(imports, "existing_emails", return_value=existing or set()):
            return imports.validate_students(rows, refs or references())

    def test_optional_placement_and_same_defaults_as_add_user(self):
        prepared, preview, errors = self.validate([(2, student(email="  STUDENT@EXAMPLE.TEST  "))])
        self.assertFalse(errors)
        fields = prepared[0][1]
        self.assertEqual(fields["email"], "student@example.test")
        self.assertEqual(fields["username"], "Test Learner")
        self.assertEqual((fields["learner_type"], fields["type"], fields["status"]), ("commercial", "User", "FullUser"))
        self.assertEqual((fields["country"], fields["learning_provider"]), ("United Kingdom", "Kent Business College"))
        self.assertEqual(preview[0]["row"], 2)

    def test_existing_and_file_duplicates_are_case_insensitive_and_row_specific(self):
        _, _, errors = self.validate([(2, student()), (4, student(email="STUDENT@EXAMPLE.TEST"))],
                                     existing={"student@example.test"})
        self.assertEqual([error["row"] for error in errors], [2, 4])
        self.assertIn("already belongs", errors[0]["message"])
        self.assertIn("row 2", errors[1]["message"])

    def test_existing_email_query_checks_all_enrolments_profiles_and_account_subject_types(self):
        with patch.object(imports.EnrolmentUser.all_learners, "annotate") as learners, \
             patch.object(imports.LearnerProfile.objects, "annotate") as profiles, \
             patch.object(imports.LoginAccount.objects, "annotate") as accounts:
            learners.return_value.filter.return_value.values_list.return_value = ["a@example.test"]
            profiles.return_value.filter.return_value.values_list.return_value = ["b@example.test"]
            accounts.return_value.filter.return_value.values_list.return_value = ["c@example.test"]
            emails = {"a@example.test", "b@example.test", "c@example.test"}
            self.assertEqual(imports.existing_emails(emails), emails)
            for manager in (learners, profiles, accounts):
                manager.return_value.filter.assert_called_once_with(import_email__in=emails)

    def test_required_values_email_and_birth_date_errors(self):
        _, _, errors = self.validate([(2, student(firstName="", email="invalid", dob="2026-02-31"))])
        self.assertEqual({error["field"] for error in errors}, {"First name", "Email", "Date of birth"})

    def test_canonical_names_employer_organisation_and_owner_are_resolved(self):
        prepared, _, errors = self.validate([(2, student(programme="business", cohort="SEPTEMBER", group="a",
                                                       employerId="42", caseOwner="test coach"))])
        self.assertFalse(errors)
        fields = prepared[0][1]
        self.assertEqual((fields["programme"], fields["cohort"], fields["group"]), ("Business", "September", "A"))
        self.assertEqual(fields["employer_id"], 42)
        self.assertEqual(fields["employer"], "Test Employer")
        self.assertEqual(fields["organization"], "Company")
        self.assertEqual(fields["case_owner"], "Test Coach")

    def test_placement_rejects_unknown_wrong_parent_and_missing_parent(self):
        for payload in (student(programme="Unknown"), student(cohort="September"),
                        student(group="A"), student(programme="Computing", cohort="September"),
                        student(programme="Computing", cohort="October", group="A")):
            with self.subTest(payload=payload):
                _, _, errors = self.validate([(2, payload)])
                self.assertTrue(errors)

    def test_ambiguous_name_is_rejected(self):
        refs = references()
        refs["programmes"].append({"programme_id": "p3", "name": "Business"})
        _, _, errors = self.validate([(2, student(programme="Business"))], refs=refs)
        self.assertIn("ambiguous", errors[0]["message"])

    def test_unknown_employer_and_case_owner_are_rejected(self):
        _, _, errors = self.validate([(2, student(employerId="99", caseOwner="Missing"))])
        self.assertEqual({error["field"] for error in errors}, {"Employer ID", "Case owner"})


class ImportEndpointTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory(HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.enterContext(patch("login.permissions.authenticate_request", return_value=SimpleNamespace(role="staff")))
        self.enterContext(patch.object(imports, "load_references", side_effect=references))
        self.enterContext(patch.object(imports, "existing_emails", return_value=set()))

    def request(self, **data):
        return self.factory.post("/learner_api/enrolment-users/import/", {"file": upload(), **data})

    def test_missing_custom_header_cannot_write(self):
        with patch.object(views, "_create_enrolment_user") as create:
            response = imports.import_students(RequestFactory().post("/", {"file": upload(), "dryRun": "false"}))
        self.assertEqual(response.status_code, 403)
        create.assert_not_called()

    def test_preview_defaults_to_read_only_and_never_creates_anything(self):
        with patch.object(views, "_create_enrolment_user") as create, patch.object(imports.transaction, "atomic") as atomic:
            response = imports.import_students(self.request())
        body = json.loads(response.content)
        self.assertEqual(response.status_code, 200)
        self.assertEqual((body["count"], body["imported"], body["results"]), (1, 0, []))
        self.assertEqual(body["preview"][0]["name"], "Test Learner")
        create.assert_not_called()
        atomic.assert_not_called()

    def test_invalid_rows_prevent_all_writes(self):
        request = self.factory.post("/", {"dryRun": "false", "file": upload([["A", "B", "invalid"]])})
        with patch.object(views, "_create_enrolment_user") as create:
            response = imports.import_students(request)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(json.loads(response.content)["errors"][0]["row"], 2)
        create.assert_not_called()

    def test_commit_revalidates_existing_emails(self):
        with patch.object(imports, "existing_emails", return_value={"student@example.test"}), \
             patch.object(views, "_create_enrolment_user") as create:
            response = imports.import_students(self.request(dryRun="false"))
        self.assertEqual(response.status_code, 400)
        create.assert_not_called()

    def test_success_requires_account_and_uses_enrolment_transaction(self):
        row = {"id": "12", "name": "Test Learner", "invitation": {"awaitingInvitation": True}}
        with patch.object(views, "_create_enrolment_user", return_value=row) as create, \
             patch.object(imports.transaction, "atomic") as atomic:
            response = imports.import_students(self.request(dryRun="false"))
        self.assertEqual(response.status_code, 201)
        body = json.loads(response.content)
        self.assertEqual((body["count"], body["imported"], body["results"]), (1, 1, [row]))
        atomic.assert_called_once_with(using="enrolment")
        self.assertTrue(create.call_args.kwargs["require_account"])

    def test_account_failure_exits_transaction_with_error_and_reports_no_partial_success(self):
        request = self.factory.post("/", {"dryRun": "false", "file": upload(
            [["A", "B", "a@example.test"], ["C", "D", "b@example.test"]])})
        with patch.object(views, "_create_enrolment_user", side_effect=[{"id": "1"}, ValidationError("Account failed")]), \
             patch.object(imports.transaction, "atomic") as atomic:
            response = imports.import_students(request)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(atomic.return_value.__exit__.call_args.args[0], ValidationError)
        body = json.loads(response.content)
        self.assertEqual((body["results"], body["imported"], body["errors"][0]["row"]), ([], 0, 3))

    def test_database_failure_is_not_a_success_and_does_not_expose_database_details(self):
        with patch.object(imports, "load_references", side_effect=DatabaseError("secret database host")):
            response = imports.import_students(self.request())
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("secret", response.content.decode())

    def test_invalid_dry_run_or_missing_file_cannot_write(self):
        for request in (self.factory.post("/", {}), self.request(dryRun="maybe")):
            with self.subTest(request=request), patch.object(views, "_create_enrolment_user") as create:
                self.assertEqual(imports.import_students(request).status_code, 400)
                create.assert_not_called()

    def test_template_download_is_excel_private_and_contains_instructions(self):
        response = imports.import_template(self.factory.get("/"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertEqual(response["Cache-Control"], "private, no-store")
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("Instructions", workbook.sheetnames)
        workbook.close()

    def test_both_endpoints_deny_learner_and_anonymous_access(self):
        for role, status in (("learner", 403), (None, 401)):
            with patch("login.permissions.authenticate_request", return_value=SimpleNamespace(role=role) if role else None):
                self.assertEqual(imports.import_template(self.factory.get("/")).status_code, status)
                self.assertEqual(imports.import_students(self.request()).status_code, status)


class SharedCreationTests(SimpleTestCase):
    def test_shared_create_preserves_cohort_dates_coach_progression_and_account_provisioning(self):
        learner = EnrolmentUser(id=1, username="Test Learner", email="a@example.test", learner_type="commercial")
        request = RequestFactory().post("/")
        with patch.object(views, "cohort_delivery_window", return_value=(date(2026, 9, 1), date(2027, 9, 1), date(2027, 12, 1))), \
             patch.object(views, "case_owner_coach", return_value={"coach_name": "Test Coach", "coach_email": "coach@example.test"}), \
             patch.object(EnrolmentUser.all_learners, "create", return_value=learner) as create, \
             patch.object(views, "advance_learner") as advance, \
             patch.object(views, "_send_platform_invitation", return_value={"awaitingInvitation": True}) as invite:
            result = views._create_enrolment_user(request, {"programme": "Business", "cohort": "September", "case_owner": "Test Coach"}, require_account=True)
        fields = create.call_args.kwargs
        self.assertEqual(fields["start_date"], date(2026, 9, 1))
        self.assertEqual(fields["practical_period_end_date"], "2027-09-01")
        self.assertEqual(fields["apprenticeship_end_date"], "2027-12-01")
        self.assertEqual(fields["coach_email"], "coach@example.test")
        self.assertTrue(fields["invite_to_platform"])
        advance.assert_called_once_with(learner)
        invite.assert_called_once_with(request, "learner", 1, subject=learner)
        self.assertEqual(result["invitation"], {"awaitingInvitation": True})

    def test_bulk_account_error_raises_so_enclosing_transaction_rolls_back(self):
        learner = EnrolmentUser(id=1, username="Test", email="a@example.test")
        with patch.object(views, "cohort_delivery_window", return_value=(None, None, None)), \
             patch.object(EnrolmentUser.all_learners, "create", return_value=learner), \
             patch.object(views, "advance_learner"), \
             patch.object(views, "_send_platform_invitation", return_value={"error": "Unavailable"}), \
             self.assertRaisesMessage(ValidationError, "No students were imported"):
            views._create_enrolment_user(RequestFactory().post("/"), {}, require_account=True)
