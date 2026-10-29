"""Real PPTX regression checks, synthetic sources, no database or external services."""
import io
import json
from copy import deepcopy
from inspect import unwrap
from unittest.mock import patch

from django.test import RequestFactory, SimpleTestCase
from pptx import Presentation
from pptx.util import Inches

from .compact_presentation import build_compact_deck, FORMAT, ASSETS


def sections():
    return [{'title': 'MCM Meeting' if i == 0 else 'EPA preparation' if i == 16 else f'Monthly section {i}',
             'body': 'I used the learning to improve my planning.\nI will review the results with my coach.',
             'notes': f'Source {i}: ' + 'Full original learner detail. ' * 400} for i in range(17)]


class CompactPresentationTests(SimpleTestCase):
    def setUp(self):
        for target in ('socket.socket.connect', 'socket.create_connection'):
            self.enterContext(patch(target, side_effect=AssertionError('No network permitted')))

    def test_exact_slide_count_original_logo_and_full_notes_without_source_mutation(self):
        source = sections()
        before = deepcopy(source)
        deck = Presentation(io.BytesIO(build_compact_deck(source)))
        self.assertEqual(len(deck.slides), 17)
        original_logo = (ASSETS / 'kbc-original-logo.png').read_bytes()
        for index, slide in enumerate(deck.slides):
            self.assertIn(source[index]['notes'], slide.notes_slide.notes_text_frame.text)
            text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            self.assertIn(source[index]['title'], text)
            for sentence in source[index]['body'].splitlines():
                self.assertIn(sentence, text.replace('\v', ' '))
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)
        self.assertTrue(any(s.shape_type == 13 and s.image.blob == original_logo for s in deck.slides[0].shapes))
        self.assertEqual(source, before)

    def test_manual_added_slides_are_retained_without_silent_cap_or_extra_cover(self):
        source = sections() + [{'title': 'My additional slide', 'body': 'My own additional example.', 'notes': 'Full additional example'}]
        result = Presentation(io.BytesIO(build_compact_deck(source)))
        self.assertEqual(len(result.slides), 18)
        self.assertIn('My additional slide', result.slides[-1].notes_slide.notes_text_frame.text)

    def test_readability_overflow_and_invalid_notes_are_rejected_without_silent_pagination(self):
        for changes in ({'body': 'x' * 601}, {'body': '\n'.join(['a'] * 7)}, {'title': 't' * 81}, {'notes': ['not text']}, {'notes': 'x' * 2000001}, {'body': 'Wide words ' * 54}):
            with self.subTest(changes=list(changes)):
                source = sections()
                source[1].update(changes)
                with self.assertRaisesRegex(ValueError, 'Slide 2'):
                    build_compact_deck(source)

    def reference(self):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13), Inches(7.5)
        for index in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = 'Private reference title'
            slide.placeholders[1].text = 'Private reference content'
            slide.shapes.add_picture(str(ASSETS / 'kbc-ibis-wreath.png'), Inches(11.7), 0, width=Inches(.6))
            slide.notes_slide.notes_text_frame.text = 'Private reference notes'
        output = io.BytesIO(); deck.save(output)
        return output.getvalue()

    def test_selected_reference_retains_dimensions_and_images_without_extra_cover_or_private_notes(self):
        source = sections()
        deck = Presentation(io.BytesIO(build_compact_deck(source, self.reference(), {'coverSlide': 1, 'contentSlide': 2})))
        self.assertEqual(len(deck.slides), 17)
        self.assertEqual(deck.slide_width, Inches(13))
        for index, slide in enumerate(deck.slides):
            self.assertNotIn('Private reference', slide.notes_slide.notes_text_frame.text)
            self.assertIn(source[index]['notes'], slide.notes_slide.notes_text_frame.text)
            text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            self.assertNotIn('Private reference', text)
            self.assertEqual(sum(s.shape_type == 13 for s in slide.shapes), 2)

    def test_small_reference_reports_fit_problem_instead_of_cutting_text_or_shrinking_fonts(self):
        deck = Presentation(io.BytesIO(self.reference()))
        for slide in deck.slides:
            slide.shapes.title.height = Inches(.1)
        output = io.BytesIO(); deck.save(output)
        with self.assertRaisesRegex(ValueError, 'does not fit'):
            build_compact_deck(sections(), output.getvalue())

    def payload(self, kind='commercial'):
        return {'learnerKind': kind, 'learnerId': '1', 'activityId': 'C1', 'actualTimeHours': '2',
                'monthlyAssignment': {'month': '2026-10', 'presentationFormat': FORMAT, 'slides': sections()}}

    def test_official_both_learner_kinds_export_compact_and_notes_affect_review_token(self):
        from .monthly_assignment import export_presentation, presentation_fingerprint
        for kind in ('commercial', 'apprenticeship'):
            payload = self.payload(kind)
            before = deepcopy(payload)
            request = RequestFactory().post('/presentation', data=json.dumps(payload), content_type='application/json')
            response = unwrap(export_presentation)(request)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(Presentation(io.BytesIO(response.content)).slides), 17)
            self.assertTrue(response['X-Presentation-Token'])
            self.assertEqual(payload, before)
            previous = presentation_fingerprint(payload)
            payload['monthlyAssignment']['slides'][0]['notes'] += ' Changed detail'
            self.assertNotEqual(previous, presentation_fingerprint(payload))

    def test_personal_course_export_uses_same_compact_renderer_without_store_or_network_writes(self):
        from .personal_learning_activities import presentation
        payload = self.payload()
        response = presentation({'id': 'personal'}, payload)
        result = Presentation(io.BytesIO(response.content))
        self.assertEqual(len(result.slides), 17)
        self.assertIn(payload['monthlyAssignment']['slides'][7]['notes'], result.slides[7].notes_slide.notes_text_frame.text)

    def test_coach_export_and_upload_priority_keep_the_saved_selection(self):
        from coach_api.presentations import presentation_file, PPTX
        payload = self.payload()
        with patch('coach_api.presentations._file_bytes', return_value=(b'original upload', 'own.pptx', PPTX)) as stored:
            content, _, _, source = presentation_file(payload, lambda *args, **kwargs: {})
            self.assertEqual(source, 'generated')
            self.assertEqual(len(Presentation(io.BytesIO(content)).slides), 17)
            stored.assert_not_called()
            payload['monthlyAssignment']['uploadedPresentation'] = {'id': 'own', 'name': 'own.pptx'}
            content, filename, _, source = presentation_file(payload, lambda *args, **kwargs: {})
            self.assertEqual((content, filename, source), (b'original upload', 'own.pptx', 'uploaded'))

    def test_legacy_absence_of_format_uses_existing_renderer(self):
        from coach_api.presentations import presentation_file
        payload = self.payload()
        del payload['monthlyAssignment']['presentationFormat']
        with patch('learner_api.presentation_design.build_deck', return_value=b'legacy') as old, patch('learner_api.compact_presentation.build_compact_deck') as new:
            self.assertEqual(presentation_file(payload, lambda *args: {})[0], b'legacy')
            old.assert_called_once_with(payload['monthlyAssignment']['slides'])
            new.assert_not_called()
