"""Generate inside the reference package, retaining its masters and artwork."""
import io
import textwrap
from copy import deepcopy
from pathlib import Path

from django.db import connections


def load_owned_template(kind, learner_id, evidence_id):
    from .evidence_storage import download_blob_bytes
    with connections['enrolment'].cursor() as cur:
        cur.execute('SELECT original_filename, container, blob_name FROM "Learner"."evidence_files" '
                    'WHERE id::text = %s AND learner_kind = %s AND learner_id = %s AND status = %s',
                    [str(evidence_id), kind, str(learner_id), 'approved'])
        row = cur.fetchone()
    if not row or not row[0].lower().endswith('.pptx'):
        raise ValueError('The PowerPoint reference is unavailable. Upload it again or choose the default KBC design.')
    return row[0], download_blob_bytes(row[1], row[2], max_bytes=15 * 1024 * 1024)


def _all_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from _all_shapes(shape.shapes)
        else:
            yield shape


def _clone_slide(deck, source):
    from pptx.oxml.ns import qn
    slide = deck.slides.add_slide(source.slide_layout)
    for element in list(slide.shapes._spTree):
        if element.tag not in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            slide.shapes._spTree.remove(element)
    rel_map = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith(('/notesSlide', '/comments', '/slide', '/slideLayout')):
            continue
        rel_map[rel.rId] = slide.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,
                                              rel.reltype, is_external=rel.is_external)
    for shape in source.shapes:
        # Tables, charts and embedded media represent original content, not a backdrop.
        if shape._element.tag == qn('p:graphicFrame') or shape._element.xpath('.//a:videoFile|.//a:audioFile'):
            continue
        element = deepcopy(shape._element)
        for original_content in element.xpath('.//p:graphicFrame'):
            original_content.getparent().remove(original_content)
        for node in element.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    if value in rel_map:
                        node.set(attr, rel_map[value])
                    else:
                        node.attrib.pop(attr)
        # Links in the old template text must not survive replacement.
        for link in element.xpath('.//a:hlinkClick|.//a:hlinkMouseOver'):
            link.getparent().remove(link)
        slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    background = source._element.cSld.find(qn('p:bg'))
    if background is not None:
        copied = deepcopy(background)
        for node in copied.iter():
            for attr, value in list(node.attrib.items()):
                if value in rel_map and attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    node.set(attr, rel_map[value])
        slide._element.cSld.insert(0, copied)
    for key in ('showMasterSp', 'showMasterPhAnim'):
        if source._element.get(key) is not None:
            slide._element.set(key, source._element.get(key))
    return slide


def _replace(frame, text, size, preserve_alignment=False):
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN
    first = frame.paragraphs[0]
    alignment = first.alignment
    style = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    frame.clear()
    frame.word_wrap = True
    for i, line in enumerate(text.splitlines() or ['']):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = alignment if preserve_alignment else PP_ALIGN.LEFT
        p._p.get_or_add_pPr().set('rtl', '0')
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.15
        run = p.add_run()
        if style is not None:
            run._r.insert(0, deepcopy(style))
        run.text = line
        run.font.size = Pt(size)


def build_from_template(slides, content, design):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation(io.BytesIO(content))
    deck.core_properties.title = "MCM Meeting"
    originals = list(deck.slides)
    if not originals:
        raise ValueError('The reference has no slides.')
    cover_index = int(design.get('coverSlide', 1)) - 1
    body_index = int(design.get('contentSlide', min(2, len(originals)))) - 1
    if not (0 <= cover_index < len(originals) and 0 <= body_index < len(originals)):
        raise ValueError('Choose valid cover and content slide numbers for your reference.')
    width, height = deck.slide_width, deck.slide_height
    logo = Path(__file__).parent / 'assets' / 'kbc-logo.png'
    # Keep the reference cover composition intact; put the narrative on content slides.
    cover = _clone_slide(deck, originals[cover_index])
    cover_text = [s for s in _all_shapes(cover.shapes) if s.has_text_frame and s.text.strip()]
    cover_title = cover.shapes.title
    if cover_title is None:
        cover_title = max(cover_text, key=lambda s: s.width * s.height) if cover_text else None
    for shape in cover_text:
        if shape.shape_id != getattr(cover_title, 'shape_id', None):
            shape.text_frame.clear()
    if cover_title is None:
        cover_title = cover.shapes.add_textbox(int(width * .15), int(height * .35), int(width * .7), int(height * .25))
    _replace(cover_title.text_frame, 'MCM Meeting', 28, preserve_alignment=True)
    cover.shapes.add_picture(str(logo), int(width - Inches(1.3)), int(height - Inches(.75)), width=Inches(1.05))
    page_count = 1
    for section_index, section in enumerate(slides):
        remaining = section['body']
        continued = False
        while remaining:
            page_count += 1
            if page_count > 500:
                raise ValueError('This presentation exceeds 500 pages. Reduce the content before exporting.')
            source = originals[body_index]
            slide = _clone_slide(deck, source)
            text_shapes = [s for s in _all_shapes(slide.shapes) if s.has_text_frame and (s.text.strip() or s.is_placeholder)]
            title_shape = slide.shapes.title
            if title_shape is None:
                title_shape = min(text_shapes, key=lambda s: s.top) if text_shapes else None
            candidates = [s for s in text_shapes if s is not title_shape and s.shape_id != getattr(title_shape, 'shape_id', None)
                          and s.width >= width * .35 and s.height >= height * .15]
            body_shape = max(candidates, key=lambda s: s.width * s.height) if candidates else None
            for shape in text_shapes:
                if shape.shape_id not in (getattr(title_shape, 'shape_id', None), getattr(body_shape, 'shape_id', None)):
                    shape.text_frame.clear()
            if title_shape is None:
                title_shape = slide.shapes.add_textbox(int(width * .08), int(height * .08), int(width * .78), int(height * .14))
            if body_shape is None:
                # References without content placeholders get a readable area over their artwork.
                title_shape.left, title_shape.top = int(width * .12), int(height * .08)
                title_shape.width, title_shape.height = int(width * .75), int(height * .14)
                body_shape = slide.shapes.add_textbox(int(width * .12), int(height * .28), int(width * .75), int(height * .58))
                body_shape.fill.solid()
                from pptx.dml.color import RGBColor
                body_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
            body_size = 18
            cols = max(20, int(body_shape.width / 12700 / (body_size * .58)))
            lines_per_page = max(1, int((body_shape.height - body_shape.text_frame.margin_top - body_shape.text_frame.margin_bottom) / 12700 / (body_size * 1.25)))
            lines = []
            for paragraph in remaining.splitlines():
                lines.extend(textwrap.wrap(paragraph, width=cols, break_long_words=True) or [''])
            page_lines, rest = lines[:lines_per_page], lines[lines_per_page:]
            _replace(title_shape.text_frame, section['title'] + (' (continued)' if continued else ''), 24, preserve_alignment=True)
            _replace(body_shape.text_frame, '\n'.join(page_lines), body_size)
            slide.shapes.add_picture(str(logo), int(width - Inches(1.3)), int(height - Inches(.75)), width=Inches(1.05))
            remaining = '\n'.join(rest)
            continued = True
    # Remove the original reference slides and their notes from the output deck.
    for _ in originals:
        slide_id = deck.slides._sldIdLst[0]
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()
