"""Extract original assignment text and organise it for the historical wizard.

This is a read-only, extractive adapter: it copies source passages, never writes
answers, invents declarations, generates grades, or makes an AI service call.
"""
import io
import re
import zipfile
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import PurePosixPath
from xml.etree import ElementTree

from . import evidence_storage

MAX_DOWNLOAD = 30 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024
MAX_MEMBER = 10 * 1024 * 1024
MAX_FILES = 32
MAX_TEXT = 160_000


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, value):
        self.parts.append(value)

    def handle_endtag(self, tag):
        if tag in {'p', 'div', 'li', 'br', 'h1', 'h2', 'h3'}:
            self.parts.append('\n')


def plain_html(value):
    parser = PlainHTML()
    parser.feed(str(value or ''))
    return ''.join(parser.parts).strip()


def _zip_members(archive):
    members = archive.infolist()
    if len(members) > 2000 or sum(item.file_size for item in members) > MAX_EXPANDED:
        raise ValueError('Archive is too large to preview safely.')
    return members


def _read_member(archive, member):
    if member.file_size > MAX_MEMBER or member.file_size > max(member.compress_size, 1) * 250:
        raise ValueError('An archive entry exceeds the preview limit.')
    if member.flag_bits & 1:
        raise ValueError('Encrypted entries cannot be previewed.')
    return archive.read(member)


def extract_documents(data, name):
    """Read bounded PDF/Office/text content, including files inside ZIPs, in memory."""
    documents, notices, files = [], [], []
    budget = {'text': MAX_TEXT, 'expanded': MAX_EXPANDED, 'entries': 0}

    def add(name, text, kind='assignment'):
        text = text.replace('\x00', '').strip()
        if len(text) > budget['text']:
            text = text[:budget['text']]
            notices.append('Text preview is limited; the original file contains the complete submission.')
        budget['text'] -= len(text)
        if text:
            documents.append({'name': name, 'text': text, 'kind': kind})
        else:
            notices.append(f'{name}: no readable text was found. Open the original to see images or scanned pages.')

    def parse(data, name, depth=0):
        if budget['entries'] >= MAX_FILES or budget['text'] <= 0:
            notices.append('Only the first part of this submission is previewed. Open the original for all files.')
            return
        budget['entries'] += 1
        files.append(name)
        suffix = PurePosixPath(name.lower()).suffix
        try:
            if suffix == '.pdf':
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                if len(reader.pages) > 100:
                    notices.append(f'{name}: only the first 100 pages are previewed.')
                parts, size = [], 0
                for page in reader.pages[:100]:
                    part = page.extract_text() or ''
                    parts.append(part[:budget['text'] - size])
                    size += len(parts[-1])
                    if size >= budget['text']:
                        notices.append(f'{name}: text preview was shortened; open the original for the complete document.')
                        break
                add(name, '\n\n'.join(parts))
            elif suffix in {'.docx', '.pptx', '.xlsx'}:
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    members = _zip_members(archive)
                    if suffix == '.docx':
                        xml = _read_member(archive, archive.getinfo('word/document.xml'))
                        root = ElementTree.fromstring(xml)
                        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                        add(name, '\n\n'.join(''.join(p.itertext()) for p in root.findall('.//w:p', ns)))
                    elif suffix == '.pptx':
                        slides = sorted((m for m in members if re.fullmatch(r'ppt/slides/slide\d+\.xml', m.filename)),
                                        key=lambda m: int(re.search(r'slide(\d+)', m.filename)[1]))
                        parts = []
                        for member in slides[:100]:
                            root = ElementTree.fromstring(_read_member(archive, member))
                            parts.append('\n'.join(node.text or '' for node in root.iter() if node.tag.endswith('}t')))
                        add(name, '\n\n'.join(parts), 'presentation')
                    else:
                        from openpyxl import load_workbook
                        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
                        try:
                            parts, sheets = [], []
                            for sheet in workbook.worksheets[:10]:
                                parts.append(sheet.title)
                                rows = []
                                for values in sheet.iter_rows(max_row=500, max_col=40, values_only=True):
                                    if any(v is not None for v in values):
                                        cells = [str(v) if v is not None else '' for v in values]
                                        rows.append(cells)
                                        parts.append(' | '.join(cells))
                                sheets.append({'name': sheet.title, 'rows': rows})
                            add(name, '\n'.join(parts))
                            if documents and documents[-1]['name'] == name:
                                documents[-1]['sheets'] = sheets
                            if len(workbook.worksheets) > 10 or any((s.max_row or 0) > 500 or (s.max_column or 0) > 40 for s in workbook.worksheets[:10]):
                                notices.append(f'{name}: this spreadsheet preview is shortened. Open the original for all cells and charts.')
                        finally:
                            workbook.close()
            elif suffix in {'.txt', '.md', '.csv'}:
                add(name, data.decode('utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig', errors='replace'))
            elif suffix == '.zip':
                if depth >= 2:
                    raise ValueError('Nested archive preview limit reached.')
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    for member in _zip_members(archive):
                        if member.is_dir() or member.filename.startswith('__MACOSX/'):
                            continue
                        if budget['entries'] >= MAX_FILES:
                            notices.append('Archive file limit reached; open the original for remaining files.')
                            break
                        label = f'{name} / {member.filename}'
                        if member.file_size > budget['expanded']:
                            notices.append(f'{label}: archive preview size limit reached.')
                            break
                        try:
                            content = _read_member(archive, member)
                            budget['expanded'] -= len(content)
                            parse(content, label, depth + 1)
                        except (ValueError, RuntimeError, zipfile.BadZipFile):
                            notices.append(f'{label}: could not preview this archive entry.')
            else:
                notices.append(f'{name}: this file is preserved in the original; a text preview is unavailable.')
        except Exception:
            # Do not expose parser internals or storage paths to the client.
            notices.append(f'{name}: text could not be extracted. Open the original file.')

    if len(data) > MAX_DOWNLOAD:
        return {'documents': [], 'files': [name], 'notices': ['This file exceeds the text-preview size limit. Open the original file.']}
    parse(data, name)
    return {'documents': documents, 'files': files, 'notices': list(dict.fromkeys(notices))}


@lru_cache(maxsize=32)
def _cached_extract(blob, name, revision):
    data = evidence_storage.download_blob_bytes('fetch-aptem-evidences', blob, max_bytes=MAX_DOWNLOAD)
    return extract_documents(data, name)


PATTERNS = {
    2: re.compile(r'\b(?:[KSB]\d{1,3}|KSBs?|knowledge,? skills|off.the.job|hours? spent|time spent)\b', re.I),
    3: re.compile(r'\b(?:reflect\w*|lessons? learnt|lessons? learned|I (?:learned|learnt|understood)|learning outcomes?|personal development|gained skills|new knowledge)\b', re.I),
    4: re.compile(r'\b(?:business impact|employer benefit|business benefit|cost savings?|efficien\w*|measurable|productivity|return on investment|improved|improvement in|benefits?|organisational impact)\b', re.I),
    5: re.compile(r'\b(?:action plan|next steps?|next month|future|I will|I plan|recommend\w*|areas? for improvement|EPA|end.point assessment|further develop)\b', re.I),
    7: re.compile(r'\b(?:coaching (?:meeting|session)|meeting minutes|presentation|slide deck)\b', re.I),
}


def _passages(text):
    """Keep source paragraphs, grouping PDF lines rather than inventing a summary."""
    paragraphs = re.split(r'\n\s*\n', text)
    if len(paragraphs) < 3:
        paragraphs = text.splitlines()
    return [part.strip() for part in paragraphs if part.strip()]


def organise_content(row, extraction, feedbacks):
    titles = ['Assignment answer', 'Evidence & cross-referencing', 'KSBs & hours claimed',
              'Full-month reflection', 'Impact & employer benefit', 'Action plan & EPA',
              'Quality checks', 'Coaching & presentation']
    cards = [{'title': title, 'sections': []} for title in titles]

    def section(index, label, text, source, kind='original'):
        if text:
            cards[index]['sections'].append({'label': label, 'text': text, 'source': source, 'kind': kind})

    for document in extraction['documents']:
        name, text = document['name'], document['text']
        if document['kind'] == 'presentation':
            section(7, 'Original presentation content', text, name)
            continue
        section(0, 'Original assignment content', text, name)
        for index, pattern in PATTERNS.items():
            matches = [p for p in _passages(text) if pattern.search(p)]
            section(index, 'Extracts from the assignment', '\n\n'.join(matches), name)

    section(1, 'Submitted files', '\n'.join(extraction['files']), row.get('evidence_name') or 'Original submission', 'record')
    section(6, 'Original assessment status', row.get('evidence_status') or 'Not recorded', 'Source assessment record', 'record')
    for feedback in feedbacks:
        if not isinstance(feedback, dict):
            continue
        text = plain_html(feedback.get('message'))
        source = 'Tutor feedback' + (f" — {feedback['author']}" if feedback.get('author') else '')
        section(6, 'Original tutor feedback', text, source, 'feedback')
        for index in (2, 3, 4, 5):
            matches = [p for p in _passages(text) if PATTERNS[index].search(p)]
            section(index, 'Tutor observations (not a learner declaration)', '\n\n'.join(matches), source, 'feedback')
    empty = [
        'No readable assignment text was found. Open the original file.',
        'No file references were recorded.',
        'No explicit KSB or time passages were found in the readable source text.',
        'No reflection passage was identified in the readable source text.',
        'No employer-impact passage was identified in the readable source text.',
        'No action-plan or EPA passage was identified in the readable source text.',
        'No original assessment information was recorded.',
        'No coaching or presentation content was found in the readable files. No booking or presentation export has been inferred.',
    ]
    for index, card in enumerate(cards):
        card['emptyMessage'] = empty[index]
    return {'cards': cards, 'notices': extraction['notices'], 'method': 'source-extracts'}


def load_assignment_content(row, feedbacks):
    name = row.get('evidence_name') or 'Assignment'
    try:
        if not row.get('file_blob'):
            raise ValueError('Missing file')
        extraction = _cached_extract(row['file_blob'], name,
                                     (str(row.get('source_hash') or ''), str(row.get('source_updated_at') or '')))
    except Exception:
        extraction = {'documents': [], 'files': [name],
                      'notices': ['The original file could not be read for this preview. Reload to retry, or open the original file.']}
    return organise_content(row, extraction, feedbacks)
