"""Extract reference styling and create branded presentations."""
import io
import json
import logging
import re
import textwrap
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from login.permissions import learner_self_only

DEFAULT = {'accent': '51258B', 'font': 'Aptos', 'ratio': 16 / 9}


def extract_design(content):
    if len(content) > 15 * 1024 * 1024:
        raise ValueError('The design reference must be no larger than 15 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 60 * 1024 * 1024:
                raise ValueError('The reference presentation is too large to process.')
            if any(e.flag_bits & 1 or e.file_size > 20 * 1024 * 1024 for e in entries):
                raise ValueError('Encrypted or oversized presentation entries are not supported.')
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
            presentation = ET.fromstring(archive.read('ppt/presentation.xml'))
            size = presentation.find('p:sldSz', ns)
            design = dict(DEFAULT)
            if size is not None:
                design['ratio'] = max(1.0, min(2.4, int(size.attrib['cx']) / int(size.attrib['cy'])))
            theme = ET.fromstring(archive.read('ppt/theme/theme1.xml'))
            accent = theme.find('.//a:clrScheme/a:accent1/a:srgbClr', ns)
            font = theme.find('.//a:fontScheme/a:minorFont/a:latin', ns)
            if accent is not None and re.fullmatch('[0-9a-fA-F]{6}', accent.get('val', '')):
                design['accent'] = accent.get('val').upper()
            if font is not None and font.get('typeface'):
                design['font'] = font.get('typeface')[:80]
            return design
    except (zipfile.BadZipFile, KeyError, ET.ParseError, ZeroDivisionError) as exc:
        raise ValueError('Upload a valid, unencrypted .pptx presentation.') from exc


@csrf_exempt
@require_POST
@learner_self_only(query_param='learnerId')
def upload_design(request):
    try:
        from .presentation_template import load_owned_template
        from pptx import Presentation
        payload = json.loads(request.body or '{}')
        if not isinstance(payload, dict) or payload.get('learnerKind') not in ('commercial', 'apprenticeship') or not payload.get('evidenceId'):
            raise ValueError('Upload and select an approved PowerPoint reference first.')
        name, content = load_owned_template(payload['learnerKind'], request.GET.get('learnerId'), payload['evidenceId'])
        design = extract_design(content)
        deck = Presentation(io.BytesIO(content))
        if not deck.slides:
            raise ValueError('Choose a reference with at least one slide.')
        return JsonResponse({'design': {**design, 'name': name, 'evidenceId': payload['evidenceId'],
                                      'slideCount': len(deck.slides), 'coverSlide': 1, 'contentSlide': min(2, len(deck.slides))}})
    except (ValueError, TypeError) as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    except Exception:
        logging.getLogger(__name__).exception('Could not load presentation reference')
        return JsonResponse({'error': 'Could not read the stored reference. Try uploading it again.'}, status=503)


def build_deck(slides, design=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    design = design if isinstance(design, dict) else {}
    accent = str(design.get('accent', DEFAULT['accent']))
    if not re.fullmatch('[0-9a-fA-F]{6}', accent):
        accent = DEFAULT['accent']
    font = str(design.get('font') or DEFAULT['font'])[:80]
    ratio = max(1.0, min(2.4, float(design.get('ratio') or DEFAULT['ratio'])))
    deck = Presentation()
    deck.core_properties.title = "MCM Meeting"
    deck.slide_width, deck.slide_height = Inches(12), Inches(12 / ratio)
    logo = Path(__file__).parent / 'assets' / 'kbc-logo.png'
    height = 12 / ratio
    line_budget = max(6, int((height - 2.4) / 0.32))
    page_number = 0
    for content in slides:
        lines = []
        for paragraph in content['body'].splitlines():
            lines.extend(textwrap.wrap(paragraph, width=82, break_long_words=True) or [''])
        for start in range(0, len(lines), line_budget):
            page_number += 1
            if page_number > 500:
                raise ValueError('This presentation exceeds 500 pages. Split it into smaller presentations.')
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor.from_string(accent)
            bar.line.fill.background()
            slide.shapes.add_picture(str(logo), Inches(10.4), Inches(0.3), width=Inches(1.05))
            title_shape = slide.shapes.title
            title_shape.left, title_shape.top = Inches(0.55), Inches(0.35)
            title_shape.width, title_shape.height = Inches(9.5), Inches(1.05)
            title = title_shape.text_frame
            title.word_wrap = True
            title.text = ('MCM Meeting' if page_number == 1 else content['title'] + (' (continued)' if start else ''))
            for p in title.paragraphs:
                p.font.name = font
                p.font.size = Pt(25)
                p.font.bold = True
                brightness = sum(int(accent[i:i + 2], 16) * weight for i, weight in zip((0, 2, 4), (0.2126, 0.7152, 0.0722)))
                p.font.color.rgb = RGBColor.from_string(accent if brightness < 165 else '172033')
            frame = slide.shapes.add_textbox(Inches(0.65), Inches(1.55), Inches(10.7), Inches(height - 2.25)).text_frame
            frame.word_wrap = True
            frame.text = '\n'.join(lines[start:start + line_budget])
            for p in frame.paragraphs:
                p.font.name = font
                p.font.size = Pt(18)
                p.font.color.rgb = RGBColor.from_string('172033')
                p.space_after = Pt(0)
            footer = slide.shapes.add_textbox(Inches(0.65), Inches(height - 0.5), Inches(10.7), Inches(0.3)).text_frame
            footer.text = f'KBC | Monthly learning presentation                                      {page_number}'
            footer.paragraphs[0].font.size = Pt(10)
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()
