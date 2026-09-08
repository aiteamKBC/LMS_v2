"""Pure/fully mocked tests: never create a test database or change stored data."""
import io
import json
from copy import deepcopy
from inspect import unwrap
from types import SimpleNamespace
from unittest.mock import patch, MagicMock

from django.core import signing
from django.test import RequestFactory, SimpleTestCase, override_settings
from pptx import Presentation

from .monthly_assignment import assignment_checks, booked_coaching, export_presentation, month_bounds, presentation_fingerprint, valid_presentation


@override_settings(SECRET_KEY="monthly-assignment-test-key")
class MonthlyAssignmentTests(SimpleTestCase):
    def payload(self):
        sentence = " ".join(["learning"] * 25)
        monthly = {
            "version": 2, "month": "2026-09", "understood": sentence, "gainedSkills": sentence,
            "evidence": [{"id": "file-1", "name": "report.pdf", "points": "1"}],
            "claims": [{"code": "K1", "explanation": sentence, "evidenceIds": ["file-1"]}],
            "plannedReviewed": True, "newKnowledge": True, "newSkills": True, "sharingConsent": True,
            "lmsReflection": sentence, "integratedReflection": sentence, "careerImpact": sentence,
            "jobImpact": sentence, "employerImpact": sentence, "employerBenefit": True,
            "actionPlan": sentence, "epaPreparedness": sentence, "meetingKey": "meeting-1",
            "slides": [{"title": "My work", "body": sentence}], "presentationReviewed": True,
        }
        payload = {"learnerKind": "commercial", "learnerId": "1", "activityId": "COMP-1", "activityType": "assignment",
                   "assignmentAnswer": " ".join(["work"] * 120), "whatYouLearned": sentence,
                   "businessImpact": sentence, "actualTimeHours": "8", "monthlyAssignment": monthly}
        monthly["presentationToken"] = signing.dumps(presentation_fingerprint(payload), salt="monthly-assignment-pptx")
        return payload

    def checks(self, payload, **kwargs):
        return {c["key"]: c["passed"] for c in assignment_checks(payload, evidence_ids=kwargs.get("evidence_ids", {"file-1"}), meeting_booked=kwargs.get("meeting_booked", True), allowed_ksbs={"K1"})}

    def test_complete_submission_has_thirteen_passing_checks_and_no_six_hour_cap(self):
        checks = self.checks(self.payload())
        self.assertEqual(len(checks), 13)
        self.assertTrue(all(checks.values()))

    def test_empty_draft_is_safe_to_check_but_not_ready(self):
        self.assertFalse(any(self.checks({}, meeting_booked=False).values()))

    def test_import_flag_from_a_client_does_not_bypass_checks(self):
        for payload in [{"submissionOrigin": "imported_legacy"}, {"monthlyAssignment": {"version": 1, "imported": True}}]:
            self.assertFalse(self.checks(payload)["answer"])
            self.assertFalse(self.checks(payload)["presentation"])

    def test_evidence_must_belong_to_learner_and_be_approved(self):
        checks = self.checks(self.payload(), evidence_ids=set())
        self.assertFalse(checks["evidence"])
        self.assertFalse(checks["ksbs"])

    def test_evidence_must_reference_existing_answer_points(self):
        for points in ["", "0", "2", "1,abc"]:
            payload = self.payload()
            payload["monthlyAssignment"]["evidence"][0]["points"] = points
            self.assertFalse(self.checks(payload)["evidence"], points)

    def test_http_links_can_be_used_but_javascript_cannot(self):
        payload = self.payload()
        entry = payload["monthlyAssignment"]["evidence"][0]
        entry.update(id="link:1", url="https://example.com/report")
        self.assertTrue(self.checks(payload)["evidence"])
        entry["url"] = "javascript:alert(1)"
        self.assertFalse(self.checks(payload)["evidence"])

    def test_unknown_and_duplicate_ksbs_are_rejected(self):
        payload = self.payload()
        payload["monthlyAssignment"]["claims"][0]["code"] = "UNKNOWN"
        self.assertFalse(self.checks(payload)["ksbs"])
        payload = self.payload()
        payload["monthlyAssignment"]["claims"] *= 2
        self.assertFalse(self.checks(payload)["ksbs"])

    def test_outside_hours_confirmation_is_required(self):
        payload = self.payload()
        payload["outsideWorkingHours"] = True
        self.assertFalse(self.checks(payload)["hours"])
        payload["outsideWorkingHoursConfirmed"] = True
        self.assertTrue(self.checks(payload)["hours"])

    def test_hours_reject_zero_negative_nonfinite_or_invalid(self):
        for value in ["0", "-1", "NaN", "Infinity", "bad"]:
            payload = self.payload()
            payload["actualTimeHours"] = value
            self.assertFalse(self.checks(payload)["hours"], value)

    def test_declared_booking_is_not_enough(self):
        self.assertFalse(self.checks(self.payload(), meeting_booked=False)["meeting"])

    def test_presentation_edit_or_identity_change_invalidates_export(self):
        payload = self.payload()
        self.assertTrue(valid_presentation(payload))
        for field, value in [("learnerId", "2"), ("activityId", "OTHER")]:
            modified = {**payload, field: value}
            self.assertFalse(valid_presentation(modified))
        payload["monthlyAssignment"]["slides"][0]["body"] += " edit"
        self.assertFalse(valid_presentation(payload))

    def test_review_checkbox_without_export_is_not_enough(self):
        payload = self.payload()
        payload["monthlyAssignment"]["presentationToken"] = "forged"
        self.assertFalse(valid_presentation(payload))

    def test_answer_edits_invalidate_presentation_but_timer_and_booking_do_not(self):
        payload = self.payload()
        payload["actualTimeHours"] = "9"
        payload["monthlyAssignment"]["step"] = 7
        payload["monthlyAssignment"]["meetingKey"] = "another-booking"
        self.assertTrue(valid_presentation(payload))
        payload["assignmentAnswer"] += " New workplace outcome."
        self.assertFalse(valid_presentation(payload))

    def test_month_boundaries_include_leap_year(self):
        self.assertEqual(month_bounds("2028-02")[1].day, 29)
        self.assertEqual(month_bounds("2026-09")[1].day, 30)
        self.assertEqual(month_bounds("2026-99"), (None, None))

    @patch("learner_api.calendar._learner_calendar_record")
    def test_coaching_uses_real_owned_meeting_in_last_ten_days(self, get_record):
        from datetime import date
        record = SimpleNamespace(event_type="mcr", scheduled_date=date(2026, 9, 21), status="scheduled")
        get_record.return_value = record
        self.assertTrue(booked_coaching(self.payload()))
        get_record.assert_called_with("commercial", 1, "meeting-1")
        record.scheduled_date = date(2026, 9, 20)
        self.assertFalse(booked_coaching(self.payload()))
        record.scheduled_date = date(2026, 9, 30)
        record.status = "cancelled"
        self.assertFalse(booked_coaching(self.payload()))

    def test_export_is_a_real_powerpoint_and_signs_its_contents(self):
        payload = self.payload()
        request = RequestFactory().post("/", data=json.dumps(payload), content_type="application/json")
        response = unwrap(export_presentation)(request)
        self.assertEqual(response.status_code, 200)
        deck = Presentation(io.BytesIO(response.content))
        self.assertEqual(deck.slides[0].shapes.title.text, "My work")
        payload["monthlyAssignment"]["presentationToken"] = response["X-Presentation-Token"]
        self.assertTrue(valid_presentation(payload))

    def test_export_rejects_empty_or_malformed_slides(self):
        for slides in [[], [None], [{"title": "Title", "body": ""}]]:
            payload = deepcopy(self.payload())
            payload["monthlyAssignment"]["slides"] = slides
            response = unwrap(export_presentation)(RequestFactory().post("/", data=json.dumps(payload), content_type="application/json"))
            self.assertEqual(response.status_code, 400)


class MonthlyDraftPersistenceTests(SimpleTestCase):
    def post(self, payload, existing=None):
        from .reflection_submissions import _submit_reflection
        cursor = MagicMock()
        cursor.fetchone.side_effect = [existing, ("submission-1",)]
        with patch("learner_api.reflection_submissions.connections") as connections_mock, patch("learner_api.reflection_submissions.transaction.atomic"), patch("learner_api.reflection_submissions._reflection_lineage", return_value={}):
            connections_mock.__getitem__.return_value.cursor.return_value.__enter__.return_value = cursor
            response = unwrap(_submit_reflection)(RequestFactory().post("/", data=json.dumps(payload), content_type="application/json"))
        return response, cursor

    def test_partial_drafts_are_saved_without_new_submission_checks(self):
        payload = {"learnerKind": "commercial", "learnerId": "1", "activityId": "C1", "activityType": "assignment",
                   "submissionMode": "draft", "monthlyAssignment": {"version": 2, "actionPlan": "Unfinished"}}
        response, cursor = self.post(payload)
        self.assertEqual(response.status_code, 201)
        parameters = cursor.execute.call_args.args[1]
        self.assertEqual(parameters[11], "draft")
        stored = json.loads(parameters[30])
        self.assertEqual(stored["monthlyAssignment"]["actionPlan"], "Unfinished")
        self.assertEqual(stored["submissionOrigin"], "learner")

    def test_client_cannot_set_import_provenance_even_on_a_draft(self):
        response, cursor = self.post({"learnerKind": "commercial", "learnerId": "1", "activityId": "C1", "activityType": "assignment", "submissionMode": "draft", "submissionOrigin": "imported_legacy"})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(json.loads(cursor.execute.call_args.args[1][30])["submissionOrigin"], "learner")

    def test_late_draft_cannot_overwrite_a_submitted_assignment(self):
        response, cursor = self.post({"learnerKind": "commercial", "learnerId": "1", "activityId": "C1", "activityType": "assignment", "submissionMode": "draft"}, ("submitted_for_tutor_review", {}))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(cursor.execute.call_count, 1)

    def test_imported_history_is_not_overwritten_by_learner_form(self):
        response, cursor = self.post({"learnerKind": "commercial", "learnerId": "1", "activityId": "C1", "activityType": "assignment", "submissionMode": "draft"}, ("draft", {"submissionOrigin": "imported_legacy"}))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(cursor.execute.call_count, 1)

    @patch("learner_api.monthly_assignment.assignment_checks", return_value=[{"passed": True}])
    @patch("learner_api.reflection_submissions._reflection_lineage", return_value={"progress_entry_id": "progress-1"})
    @patch("learner_api.monthly_assignment.transaction.atomic")
    @patch("learner_api.monthly_assignment.connections")
    def test_progress_and_final_submission_share_transaction(self, connection_mock, atomic, lineage, checks):
        from .monthly_assignment import complete_saved_assignment
        cursor = connection_mock.__getitem__.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = ("submission-1", "draft", {"monthlyAssignment": {"claims": [{"code": "K1"}]}})
        save = MagicMock()
        record = {}
        complete_saved_assignment("commercial", "1", "C1", record, save)
        atomic.assert_called_once_with(using="enrolment")
        self.assertIn("FOR UPDATE", cursor.execute.call_args_list[0].args[0])
        self.assertEqual(record["ksbs"], ["K1"])
        save.assert_called_once()
        self.assertIn("UPDATE", cursor.execute.call_args_list[-1].args[0])

    @patch("learner_api.monthly_assignment.assignment_checks", return_value=[{"passed": True}])
    @patch("learner_api.monthly_assignment.transaction.atomic")
    @patch("learner_api.monthly_assignment.connections")
    def test_failed_progress_never_marks_assignment_submitted(self, connection_mock, atomic, checks):
        from .monthly_assignment import complete_saved_assignment
        cursor = connection_mock.__getitem__.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchone.return_value = ("submission-1", "draft", {"monthlyAssignment": {"claims": []}})
        with self.assertRaisesRegex(ValueError, "Progress failed"):
            complete_saved_assignment("commercial", "1", "C1", {}, MagicMock(side_effect=ValueError("Progress failed")))
        self.assertEqual(cursor.execute.call_count, 1)
