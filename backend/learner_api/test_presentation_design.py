import io
from unittest import TestCase
from pptx import Presentation
from .presentation_design import extract_design, build_deck

class PresentationDesignTests(TestCase):
    def test_reference_and_branded_export_preserve_text(self):
        reference = Presentation(); stream = io.BytesIO(); reference.save(stream)
        design = extract_design(stream.getvalue())
        self.assertAlmostEqual(design['ratio'], reference.slide_width / reference.slide_height)
        body = ' '.join(f'word{i}' for i in range(700))
        output = build_deck([{'title': 'Monthly learning', 'body': body}], design)
        result = Presentation(io.BytesIO(output))
        self.assertGreater(len(result.slides), 1)
        contents = []
        for slide in result.slides:
            self.assertTrue(any(shape.shape_type == 13 for shape in slide.shapes))
            contents.extend(shape.text for shape in slide.shapes if shape.has_text_frame)
        combined = ' '.join(contents)
        for i in range(700): self.assertIn(f'word{i}', combined)
        self.assertIn('KBC', combined)

    def test_invalid_reference_is_rejected(self):
        with self.assertRaises(ValueError): extract_design(b'not a powerpoint')
