"""Authoritative monthly assignment checks and real PPTX export.

Versioned answers live in learning_reflection_submissions.full_submission.
No schema creation and no import/write-back of historical assignments here.
"""
import calendar
import hashlib
import io
import json
import math
from datetime import date
from urllib.parse import urlparse

from django.core import signing
from django.db import connections, DatabaseError, transaction
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from login.permissions import learner_self_only


def text(value):
    return value.strip() if isinstance(value, str) else ""


def words(value):
    return len(text(value).split())


def mapping(value):
    return value if isinstance(value, dict) else {}


def items(value):
    return value if isinstance(value, list) else []


def month_bounds(month):
    try:
        start = date.fromisoformat(text(month) + "-01")
        end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
        return start, end
    except ValueError:
        return None, None


def presentation_fingerprint(payload):
    monthly = mapping(payload.get("monthlyAssignment"))
    contents = {key: payload.get(key) for key in ("learnerKind", "learnerId", "activityId")}
    # Timer ticks and booking selection must not invalidate a deck; changes to
    # the learner's written/evidence content do require a new reviewed export.
    contents.update({key: payload.get(key) for key in ("assignmentAnswer", "whatYouLearned", "businessImpact")})
    contents["monthly"] = {key: value for key, value in monthly.items()
                           if key not in ("step", "meetingKey", "presentationReviewed", "presentationToken")}
    return hashlib.sha256(json.dumps(contents, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def valid_presentation(payload):
    monthly = mapping(payload.get("monthlyAssignment"))
    slides = items(monthly.get("slides"))
    if not (1 <= len(slides) <= 30 and monthly.get("presentationReviewed") is True):
        return False
    try:
        digest = signing.loads(text(monthly.get("presentationToken")), salt="monthly-assignment-pptx")
        return digest == presentation_fingerprint(payload)
    except signing.BadSignature:
        return False


def approved_evidence_ids(payload):
    with connections["enrolment"].cursor() as cur:
        cur.execute(
            'SELECT id::text FROM "Learner"."evidence_files" '
            "WHERE learner_kind = %s AND learner_id = %s AND status = 'approved'",
            [payload.get("learnerKind"), str(payload.get("learnerId"))],
        )
        return {str(row[0]) for row in cur.fetchall()}


def booked_coaching(payload):
    from .calendar import _learner_calendar_record
    monthly = mapping(payload.get("monthlyAssignment"))
    start, end = month_bounds(monthly.get("month"))
    if not start or not text(monthly.get("meetingKey")):
        return False
    record = _learner_calendar_record(payload.get("learnerKind"), int(payload.get("learnerId")), monthly["meetingKey"])
    return bool(record and record.event_type == "mcr" and record.scheduled_date
                and end.replace(day=end.day - 9) <= record.scheduled_date <= end
                and record.status in ("scheduled", "in-progress", "completed", "awaiting-signature"))


def available_ksb_codes(payload):
    from .models import CommercialUser, EnrolmentUser
    from .identity import learner_profile_for_source
    from .components import component_ksb_codes
    model = CommercialUser if payload.get("learnerKind") == "commercial" else EnrolmentUser
    source = model.all_learners.filter(pk=payload.get("learnerId")).first()
    profile = learner_profile_for_source(source, payload.get("learnerId"), active_only=True) if source else None
    from .active_users import current_curriculum_ksb_items_for_learner
    programme_items = current_curriculum_ksb_items_for_learner(profile, source=source) if source else []
    return set(component_ksb_codes(payload.get("activityId"))) | {k["code"] for k in (programme_items or (profile.ksbs if profile else []))}


def assignment_checks(payload, *, evidence_ids=None, meeting_booked=None, allowed_ksbs=None):
    """Dependencies are injectable for schema-free tests; production verifies DB ownership."""
    monthly = mapping(payload.get("monthlyAssignment"))
    evidence = [mapping(e) for e in items(monthly.get("evidence"))]
    owned = approved_evidence_ids(payload) if evidence_ids is None and evidence else (evidence_ids or set())
    answer_lines = [line.strip() for line in text(payload.get("assignmentAnswer")).splitlines() if line.strip()]
    linked_ids = set()
    for entry in evidence:
        entry_id = text(entry.get("id"))
        url = text(entry.get("url"))
        is_link = entry_id.startswith("link:") and urlparse(url).scheme in ("https", "http") and bool(urlparse(url).hostname)
        point_text = text(entry.get("points")).strip()
        try:
            points = [int(p.strip()) for p in point_text.split(",") if p.strip()]
        except ValueError:
            points = []
        valid_points = not point_text or (bool(points) and all(1 <= p <= len(answer_lines) for p in points))
        if entry_id and (is_link or entry_id in owned) and valid_points:
            linked_ids.add(entry_id)
    claims = [mapping(c) for c in items(monthly.get("claims"))]
    allowed = available_ksb_codes(payload) if allowed_ksbs is None and claims else (allowed_ksbs or set())
    claimed_codes = [text(c.get("code")) for c in claims]
    try:
        hours = float(payload.get("actualTimeHours") or 0)
    except (TypeError, ValueError):
        hours = 0
    booked = booked_coaching(payload) if meeting_booked is None else meeting_booked
    checks = [
        ("answer", "Assignment answer: at least 120 words", words(payload.get("assignmentAnswer")) >= 120),
        ("learning", "Learned, understood and gained skills: at least 20 words each", all(words(v) >= 20 for v in [payload.get("whatYouLearned"), monthly.get("understood"), monthly.get("gainedSkills")])),
        ("evidence", "At least one available evidence item; answer point numbers are optional and must be valid if provided", bool(linked_ids)),
        ("ksbs", "Every claimed programme KSB has a 20-word explanation and linked evidence", bool(claims) and len(set(claimed_codes)) == len(claimed_codes) and set(claimed_codes) <= allowed and all(words(c.get("explanation")) >= 20 and bool(set(str(e) for e in items(c.get("evidenceIds"))) & linked_ids) for c in claims)),
        ("planned", "Planned hours and KSBs reviewed", monthly.get("plannedReviewed") is True),
        ("declarations", "New learning, skills and employer evidence-sharing declarations confirmed", all(monthly.get(k) is True for k in ["newKnowledge", "newSkills", "sharingConsent"])),
        ("hours", "Positive time recorded and any out-of-hours work confirmed (no six-hour cap)", math.isfinite(hours) and hours > 0 and (not payload.get("outsideWorkingHours") or payload.get("outsideWorkingHoursConfirmed") is True)),
        ("reflection", "Monthly LMS reflection and integrated understanding: at least 20 words each", all(words(monthly.get(k)) >= 20 for k in ["lmsReflection", "integratedReflection"])),
        ("benefit", "Employer benefit confirmed and measurable outcomes described (20 words)", monthly.get("employerBenefit") is True and words(payload.get("businessImpact")) >= 20),
        ("impact", "Career, job and employer impacts: at least 20 words each", all(words(monthly.get(k)) >= 20 for k in ["careerImpact", "jobImpact", "employerImpact"])),
        ("action", "Action plan and EPA preparedness: at least 20 words each", all(words(monthly.get(k)) >= 20 for k in ["actionPlan", "epaPreparedness"])),
        ("meeting", "Coaching meeting booked in the last ten days of the submission month", booked),
        ("presentation", "Presentation generated, exported and reviewed", valid_presentation(payload)),
    ]
    return [{"key": key, "label": label, "passed": bool(passed)} for key, label, passed in checks]


def parse_request(request):
    if request.method != "POST":
        raise ValueError("POST is required.")
    payload = json.loads(request.body or b"{}")
    if not isinstance(payload, dict) or payload.get("learnerKind") not in ("commercial", "apprenticeship") or not str(payload.get("learnerId", "")).isdigit() or not text(payload.get("activityId")):
        raise ValueError("A valid learner and assignment are required.")
    return payload


@csrf_exempt
@learner_self_only(body_field="learnerId")
def check_assignment(request):
    try:
        payload = parse_request(request)
        return JsonResponse({"checks": assignment_checks(payload)})
    except (ValueError, TypeError) as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    except DatabaseError:
        return JsonResponse({"error": "Could not verify evidence or coaching. Save your draft and retry."}, status=503)


@csrf_exempt
@learner_self_only(body_field="learnerId")
def export_presentation(request):
    try:
        payload = parse_request(request)
        slides = items(mapping(payload.get("monthlyAssignment")).get("slides"))
        if not 1 <= len(slides) <= 30:
            raise ValueError("Add between 1 and 30 slides before exporting.")
        if any(not text(mapping(s).get("title")) or not text(mapping(s).get("body")) or len(text(mapping(s).get("body"))) > 12000 for s in slides):
            raise ValueError("Every slide needs a title and content (up to 12,000 characters).")
        from pptx import Presentation
        from pptx.util import Pt
        deck = Presentation()
        for content in slides:
            # Long narratives are split so exported slides remain readable.
            body = text(content["body"])
            paragraphs = [body[i:i + 900] for i in range(0, len(body), 900)]
            for index, paragraph in enumerate(paragraphs):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = content["title"] + (" (continued)" if index else "")
                frame = slide.placeholders[1].text_frame
                frame.text = paragraph
                for p in frame.paragraphs:
                    p.font.size = Pt(18)
        stream = io.BytesIO()
        deck.save(stream)
        response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
        response["Content-Disposition"] = 'attachment; filename="monthly-assignment.pptx"'
        response["X-Presentation-Token"] = signing.dumps(presentation_fingerprint(payload), salt="monthly-assignment-pptx")
        response["Cache-Control"] = "no-store"
        return response
    except (ValueError, TypeError) as exc:
        return JsonResponse({"error": str(exc)}, status=400)


def complete_saved_assignment(kind, learner_id, component_id, record, save_progress):
    """Progress and final submission commit together; failed completion leaves a draft.

    The row lock also prevents a delayed autosave from overwriting submission.
    """
    from .reflection_submissions import _reflection_lineage
    from django.utils import timezone
    with transaction.atomic(using="enrolment"):
        with connections["enrolment"].cursor() as cur:
            cur.execute(
                'SELECT id, status, full_submission FROM "Learner"."learning_reflection_submissions" '
                "WHERE learner_kind = %s AND learner_id = %s AND activity_type = 'assignment' AND activity_id = %s FOR UPDATE",
                [kind, str(learner_id), component_id],
            )
            row = cur.fetchone()
            if not row or row[1] in ("accepted", "submitted_for_tutor_review"):
                raise ValueError("This assignment is missing or has already been submitted. Reload the page to see its current status.")
            payload = mapping(json.loads(row[2]) if isinstance(row[2], str) else row[2])
            if payload.get("submissionOrigin") == "imported_legacy":
                raise ValueError("Historical imported assignments cannot be completed again.")
            payload.update(learnerKind=kind, learnerId=str(learner_id), activityId=component_id)
            checks = assignment_checks(payload)
            if not all(c["passed"] for c in checks):
                raise ValueError("Some monthly submission requirements changed. Your draft is safe; recheck before submitting.")
            # Use the reviewed claims, not all planned KSBs automatically.
            record["ksbs"] = [c["code"] for c in payload["monthlyAssignment"]["claims"]]
            save_progress()
            lineage = _reflection_lineage(learner_id, component_id)
            payload.update(submissionMode="submit", qualityScore=100, qualityChecks=checks,
                           otjhConfirmed=True, signedDeclaration=True, dateCompleted=timezone.localdate().isoformat(),
                           outsideWorkingHours=record.get("outsideWorkingHours", False),
                           outsideWorkingHoursConfirmed=record.get("outsideWorkingHoursConfirmed", False),
                           outsideWorkingHoursConfirmedAt=record.get("outsideWorkingHoursConfirmedAt"))
            cur.execute(
                'UPDATE "Learner"."learning_reflection_submissions" '
                "SET status = 'submitted_for_tutor_review', full_submission = %s::jsonb, quality_score = 100, "
                "submitted_at = now(), date_completed = %s, otjh_confirmed = true, signed_declaration = true, "
                "progress_entry_id = %s WHERE id = %s",
                [json.dumps(payload), payload["dateCompleted"], lineage["progress_entry_id"], row[0]],
            )
