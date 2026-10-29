"""Manager-supplied MCM template, versioned separately from saved older decks.

Only existing text slots are populated. Artwork, masters, layouts and slide
dimensions come from the supplied file; no decorative artwork is redrawn.
"""
import io
import textwrap
from copy import deepcopy
from pathlib import Path

FORMAT = 'kbc-mcm-v2'
TEMPLATE_PATH = Path(__file__).parent / 'assets' / 'mcm-standard-v2' / 'learner-monthly-coaching.pptx'
# Stable shape ids in the approved, immutable seven-slide reference.
SLOTS = {
    1: (8, [(None, 11), (21, 20)]),
    2: (3, [(7, 8), (11, 12), (15, 16)]),
    3: (3, [(10, 11), (14, 15), (18, 19)]),
    4: (3, [(7, 8), (11, 12), (15, 16)]),
    5: (3, [(7, 8), (11, 12), (15, 16)]),
    6: (3, [(7, 8), (11, 12), (15, 16)]),
    7: (3, [(7, 8), (11, 12)]),
}


def _validate(slides):
    if not isinstance(slides, list) or not 1 <= len(slides) <= 200:
        raise ValueError('Add between 1 and 200 slides before exporting.')
    for number, section in enumerate(slides, 1):
        if not isinstance(section, dict) or any(not isinstance(section.get(k), str) or not section[k].strip() for k in ('title', 'body')):
            raise ValueError(f'Slide {number}: add a title and content.')
        if len(section['title']) > 100 or len(section['body']) > 1800:
            raise ValueError(f'Slide {number}: shorten the title/highlights and put full detail in speaker notes.')
        if not isinstance(section.get('notes', ''), str) or len(section.get('notes', '')) > 2000000:
            raise ValueError(f'Slide {number}: speaker notes must be text of no more than 2,000,000 characters.')
        if type(section.get('templateSlide', 2)) is not int or section.get('templateSlide', 2) not in SLOTS:
            raise ValueError(f'Slide {number}: choose a valid template layout.')
    if sum(len(s.get('notes', '')) for s in slides) > 2000000:
        raise ValueError('The speaker notes are too large. Split the material into separate presentations.')


def _fill(shape, text, size, number, *, check_fit=False):
    from pptx.util import Pt
    from pptx.enum.text import MSO_AUTO_SIZE
    frame = shape.text_frame
    if check_fit:
        width = shape.width - frame.margin_left - frame.margin_right
        height = shape.height - frame.margin_top - frame.margin_bottom
        cols = max(1, int(width / 12700 / (size * .60)))
        lines = sum(max(1, len(textwrap.wrap(p, cols, break_long_words=True))) for p in text.splitlines())
        if lines * size * 1.15 > height / 12700:
            raise ValueError(f'Slide {number}: a column is too long for this template. Shorten its highlights, move detail to notes, or add a slide.')
    first = frame.paragraphs[0]
    paragraph_style = deepcopy(first._p.pPr) if first._p.pPr is not None else None
    run_style = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, line in enumerate(text.splitlines() or ['Not recorded.']):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if paragraph._p.pPr is not None:
            paragraph._p.remove(paragraph._p.pPr)
        if paragraph_style is not None:
            paragraph._p.insert(0, deepcopy(paragraph_style))
        run = paragraph.add_run()
        if run_style is not None:
            run._r.insert(0, deepcopy(run_style))
        run.text = line
        run.font.size = Pt(size)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.15


def _columns(section, count):
    if 'templateSlide' not in section:
        chunks = textwrap.wrap(section['body'], 220, break_long_words=False, break_on_hyphens=False)
        if len(chunks) > count:
            raise ValueError('This added slide is too long. Shorten the highlights or add another slide.')
        return [(('Details' if i == 0 else 'More detail'), chunks[i] if i < len(chunks) else 'Not recorded.') for i in range(count)]
    blocks = section['body'].split('\n\n')
    if len(blocks) > count:
        raise ValueError(f'This layout has {count} columns. Keep one block per column, or add a slide.')
    result = []
    for block in blocks:
        heading, _, content = block.partition('\n')
        if len(heading) > 45:
            raise ValueError('Keep each column heading within 45 characters.')
        result.append((heading, content.strip() or 'Not recorded.'))
    return result + [('Details', 'Not recorded.')] * (count - len(result))


def build_standard_deck(slides, reference=None, design=None, *, fetch_image=None):
    from pptx import Presentation
    from .presentation_template import _clone_slide
    from .compact_presentation import _notes
    _validate(slides)
    if reference is not None:
        # An explicitly selected personal reference still wins over the default.
        from .compact_presentation import _from_template
        adapted = [{**s, 'body': '\n'.join(f'{h}: {b.replace(chr(10), " ")}' for h, b in _columns(s, len(SLOTS[s.get('templateSlide', 2)][1])))} for s in slides]
        return _from_template(adapted, reference, design or {})
    deck = Presentation(str(TEMPLATE_PATH))
    originals = list(deck.slides)
    if len(originals) != 7:
        raise ValueError('The approved MCM template is unavailable or invalid.')
    for number, section in enumerate(slides, 1):
        layout = section.get('templateSlide', 2)
        slide = _clone_slide(deck, originals[layout - 1])
        shapes = {s.shape_id: s for s in slide.shapes}
        title_id, cells = SLOTS[layout]
        _fill(shapes[title_id], section['title'], 28 if layout == 1 else 26, number, check_fit=True)
        for (heading_id, body_id), (heading, body) in zip(cells, _columns(section, len(cells))):
            if heading_id is not None:
                heading_shape = shapes[heading_id]
                size = 18 if layout == 1 else 17
                columns = max(1, int((heading_shape.width - heading_shape.text_frame.margin_left - heading_shape.text_frame.margin_right) / 12700 / (size * .60)))
                room = (shapes[body_id].top - heading_shape.top) / 12700
                if len(textwrap.wrap(heading, columns)) * size * 1.15 > room:
                    raise ValueError(f'Slide {number}: shorten the column heading to fit the template.')
                _fill(shapes[heading_id], heading, 18 if layout == 1 else 17, number)
            _fill(shapes[body_id], body, (17 if body_id == 11 else 16) if layout == 1 else 13, number, check_fit=True)
        if fetch_image:
            from pptx.util import Inches
            for (_, body_id), ref in zip(cells, section.get('evidenceImages') or []):
                if not ref or ref == 'Not available':
                    continue
                target = shapes[body_id]
                body = target.text_frame.text
                raw = fetch_image(ref)
                if raw:
                    from PIL import Image
                    try:
                        with Image.open(io.BytesIO(raw)) as picture:
                            picture.verify()
                            ratio = picture.width / picture.height
                    except (OSError, ValueError):
                        _fill(target, body + '\nImage unavailable.', 13, number)
                        continue
                    width = min(target.width, int((target.height - Inches(.65)) * ratio))
                    height = int(width / ratio)
                    slide.shapes.add_picture(io.BytesIO(raw), target.left, target.top + target.height - height,
                                             width=width, height=height)
                    # Keep the evidence title and full account in notes; the
                    # caption remains above the image inside the same card.
                    _fill(target, body[:70], 13, number)
                else:
                    _fill(shapes[body_id], body + '\nImage unavailable.', 13, number)
        _notes(slide, section)
    for _ in originals:
        slide_id = deck.slides._sldIdLst[0]
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    deck.core_properties.title = 'Learner Monthly Coaching Meeting'
    deck.core_properties.author = 'Kent Business College'
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()
