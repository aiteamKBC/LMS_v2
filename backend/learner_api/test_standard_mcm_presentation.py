import io
from copy import deepcopy
from unittest.mock import patch

from django.test import SimpleTestCase
from pptx import Presentation

from .standard_mcm_presentation import TEMPLATE_PATH, FORMAT, build_standard_deck, SLOTS


def sections(count=8):
    layouts = [1, 2, 3, 4, 5, 6, 2] + [2] * (count - 8) + [7]
    return [{'title': 'Learner Monthly Coaching Meeting' if i == 0 else f'{i}. My learning', 'templateSlide': layout,
             'body': '\n\n'.join(f'Column {j + 1}\nMy original learning and workplace example.' for j in range(len(SLOTS[layout][1]))),
             'notes': f'Full original source {i}: ' + 'My original account. ' * 500} for i, layout in enumerate(layouts)]


class StandardMcmTests(SimpleTestCase):
    def setUp(self):
        for target in ('socket.socket.connect', 'socket.create_connection'):
            self.enterContext(patch(target, side_effect=AssertionError('No external access permitted')))

    def test_original_template_geometry_artwork_and_notes_are_retained_for_eight_to_twelve_slides(self):
        reference = Presentation(str(TEMPLATE_PATH))
        for count in (8, 9, 10, 12):
            source = sections(count); before = deepcopy(source)
            result = Presentation(io.BytesIO(build_standard_deck(source)))
            self.assertEqual(len(result.slides), count)
            self.assertEqual((result.slide_width, result.slide_height), (reference.slide_width, reference.slide_height))
            for data, slide in zip(source, result.slides):
                original = reference.slides[data['templateSlide'] - 1]
                self.assertEqual([(s.shape_id, s.left, s.top, s.width, s.height, s.shape_type) for s in slide.shapes],
                                 [(s.shape_id, s.left, s.top, s.width, s.height, s.shape_type) for s in original.shapes])
                self.assertIn(data['notes'], slide.notes_slide.notes_text_frame.text)
                self.assertIn(data['title'], '\n'.join(s.text for s in slide.shapes if s.has_text_frame))
            self.assertEqual(source, before)

    def test_manual_slides_can_exceed_automatic_ceiling_without_changing_saved_sources(self):
        source = sections(12) + [{'title': 'My extra example', 'body': 'My own additional explanation.', 'notes': 'Full additional account'}]
        result = Presentation(io.BytesIO(build_standard_deck(source)))
        self.assertEqual(len(result.slides), 13)
        self.assertIn('My own additional explanation.', result.slides[-1].notes_slide.notes_text_frame.text)

    def test_column_overflow_is_rejected_instead_of_creating_fifty_pages_or_hiding_text(self):
        source = sections()
        source[1]['body'] = 'Column\n' + 'Long content ' * 110
        with self.assertRaisesRegex(ValueError, 'column is too long'):
            build_standard_deck(source)

    def test_new_format_reaches_learner_personal_and_coach_export_paths(self):
        from inspect import unwrap
        import json
        from django.test import RequestFactory
        from .monthly_assignment import export_presentation
        from .personal_learning_activities import presentation
        from coach_api.presentations import presentation_file
        for kind in ('commercial', 'apprenticeship'):
            payload = {'learnerKind': kind, 'learnerId': '1', 'activityId': 'C1', 'monthlyAssignment': {'presentationFormat': FORMAT, 'month': '2026-10', 'slides': sections()}}
            request = RequestFactory().post('/', data=json.dumps(payload), content_type='application/json')
            response = unwrap(export_presentation)(request)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(len(Presentation(io.BytesIO(response.content)).slides), 8)
            self.assertEqual(len(Presentation(io.BytesIO(presentation({}, payload).content)).slides), 8)
            self.assertEqual(len(Presentation(io.BytesIO(presentation_file(payload, lambda *a: {})[0])).slides), 8)
