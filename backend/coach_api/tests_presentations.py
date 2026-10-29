"""Owned read SQL in private SQLite, real PPTX output, no live blob/DB/network."""
import io
import json
import re
import sqlite3
from copy import deepcopy
from inspect import unwrap
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from uuid import UUID

from django.db import DatabaseError
from django.test import RequestFactory, SimpleTestCase
from django.urls import resolve
from pptx import Presentation

from . import presentations as decks

IDENTIFIER = '00000000-0000-4000-8000-000000000001'
SUBMITTED = '2026-10-28T09:00:00+00:00'


def visible_text(content):
    deck = Presentation(io.BytesIO(content))
    return '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def reference_file():
    deck = Presentation()
    for index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f'Original template {index}'
        slide.placeholders[1].text = 'Original template body'
    output = io.BytesIO(); deck.save(output)
    return output.getvalue()


class CoachPresentationTests(SimpleTestCase):
    def setUp(self):
        for target in ('socket.socket.connect', 'socket.create_connection'):
            self.enterContext(patch(target, side_effect=AssertionError('No network permitted')))
        self.db = sqlite3.connect(':memory:'); self.db.row_factory = sqlite3.Row
        self.addCleanup(self.db.close)
        self.db.create_function('btrim', 1, lambda value: str(value or '').strip())
        self.db.executescript('''ATTACH DATABASE ':memory:' AS Learner;
            ATTACH DATABASE ':memory:' AS enrolment;
            CREATE TABLE enrolment.Created_users(id,Email);
            INSERT INTO enrolment.Created_users VALUES(7,'learner@example.test'),(8,'commercial@example.test');
            CREATE TABLE Learner.learners(id,enrolment_id,learner_type,coach_email,email);
            INSERT INTO Learner.learners VALUES(70,7,'apprenticeship','coach@example.test','learner@example.test'),
                (80,8,'commercial','coach@example.test','commercial@example.test');
            CREATE TABLE Learner.learning_reflection_submissions(id,learner_kind,learner_id,activity_id,activity_type,status,full_submission,submitted_at);
            CREATE TABLE Learner.evidence_files(id,learner_kind,learner_id,section_ref,original_filename,container,blob_name,size_bytes,status);
        ''')
        self.document = {'activityId': 'COMP-1', 'signedDeclaration': True, 'actualTimeHours': '2',
            'monthlyAssignment': {'month': '2026-10', 'slides': [
                {'title': 'Saved learner title', 'body': 'Original saved learning, including the learner’s own edits.'},
                {'title': 'Saved application', 'body': 'Original workplace application. No new content or AI generation.'}]},
            'assignmentAttemptHistory': [{'status': 'rejected', 'content': {'monthlyAssignment': {'slides': [{'title': 'History', 'body': 'Prior attempt unchanged.'}]}}}]}
        self.db.execute('INSERT INTO Learner.learning_reflection_submissions VALUES(?,?,?,?,?,?,?,?)',
            [IDENTIFIER, 'apprenticeship', '7', 'COMP-1', 'assignment', 'submitted_for_tutor_review', json.dumps(self.document), SUBMITTED])
        self.statements = []
        self.enterContext(patch.object(decks, 'query', side_effect=self.query))
        self.storage = self.enterContext(patch('learner_api.evidence_storage.download_blob_bytes', return_value=reference_file()))
        self.factory = RequestFactory()

    def query(self, sql, params=()):
        self.statements.append(sql)
        sql = re.sub(r'(\w+(?:\.\w+)?)::text', r'CAST(\1 AS TEXT)', sql)
        return [dict(row) for row in self.db.execute(sql.replace('%s', '?'), params).fetchall()]

    def update(self, **changes):
        for key, value in changes.items():
            if key == 'full_submission': value = json.dumps(value)
            self.db.execute(f'UPDATE Learner.learning_reflection_submissions SET {key}=?', [value])

    def save_document(self):
        self.update(full_submission=self.document)

    def file(self, identifier='file-1', **changes):
        row = dict(id=identifier, learner_kind='apprenticeship', learner_id='7', section_ref='COMP-1',
                   original_filename='original-upload.pptx', container='approved-test', blob_name='synthetic-file',
                   size_bytes=len(self.storage.return_value), status='approved')
        row.update(changes)
        self.db.execute(f'INSERT INTO Learner.evidence_files({",".join(row)}) VALUES({",".join("?" for _ in row)})', list(row.values()))

    def get(self, query=None, email='coach@example.test', method='GET'):
        request = self.factory.generic(method, '/presentation') if method != 'GET' else self.factory.get('/presentation', query or {})
        request.coach_email = email
        return unwrap(decks.download)(request, UUID(IDENTIFIER))

    def test_real_powerpoint_uses_saved_content_without_writes(self):
        before = [tuple(row) for row in self.db.execute('SELECT * FROM Learner.learning_reflection_submissions')]
        response = self.get({'submittedAt': SUBMITTED, 'learnerId': '999', 'slides': 'Client replacement must be ignored'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], decks.PPTX)
        self.assertIn('presentation-2026-10.pptx', response['Content-Disposition'])
        self.assertEqual(response['X-Presentation-Source'], 'generated')
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
        contents = visible_text(response.content)
        for slide in self.document['monthlyAssignment']['slides']:
            self.assertIn(slide['body'], contents)
        self.assertNotIn('Client replacement', contents)
        self.assertNotIn('Prior attempt unchanged.', contents)
        self.assertEqual(before, [tuple(row) for row in self.db.execute('SELECT * FROM Learner.learning_reflection_submissions')])
        self.assertTrue(all(sql.lstrip().upper().startswith('SELECT') for sql in self.statements))
        self.storage.assert_not_called()

    def test_commercial_learner_uses_the_same_owned_download(self):
        self.update(learner_kind='commercial', learner_id='8')
        self.assertEqual(self.get().status_code, 200)

    def test_foreign_coach_wrong_kind_unlinked_account_and_ambiguous_profiles_are_refused(self):
        self.assertEqual(self.get(email='other@example.test').status_code, 404)
        self.update(learner_kind='commercial')
        self.assertEqual(self.get().status_code, 404)
        self.update(learner_kind='apprenticeship')
        self.db.execute("UPDATE enrolment.Created_users SET Email='other@example.test' WHERE id=7")
        self.assertEqual(self.get().status_code, 404)
        self.db.execute("UPDATE enrolment.Created_users SET Email='learner@example.test' WHERE id=7")
        self.db.execute("INSERT INTO Learner.learners VALUES(71,7,'apprenticeship','coach@example.test','learner@example.test')")
        self.assertEqual(self.get().status_code, 404)
        self.db.execute("UPDATE Learner.learners SET coach_email='other@example.test' WHERE id=71")
        self.assertEqual(self.get().status_code, 404)
        self.storage.assert_not_called()

    def test_draft_reflection_extra_and_missing_submissions_are_not_exported(self):
        self.update(status='draft'); self.assertEqual(self.get().status_code, 404)
        self.update(status='accepted', activity_type='reflection'); self.assertEqual(self.get().status_code, 404)
        self.update(activity_type='extra_activity'); self.assertEqual(self.get().status_code, 404)
        self.db.execute('DELETE FROM Learner.learning_reflection_submissions')
        self.assertEqual(self.get().status_code, 404)

    def test_saved_upload_is_returned_exactly_and_has_priority_over_slides(self):
        self.file(); self.document['monthlyAssignment']['uploadedPresentation'] = {'id': 'file-1', 'name': 'untrusted-name.exe'}
        self.save_document()
        with patch('learner_api.presentation_design.build_deck', side_effect=AssertionError('Do not rebuild an uploaded file')):
            response = self.get()
        self.assertEqual(response.content, self.storage.return_value)
        self.assertIn('original-upload.pptx', response['Content-Disposition'])
        self.assertEqual(response['X-Presentation-Source'], 'uploaded')
        self.storage.assert_called_once_with('approved-test', 'synthetic-file', max_bytes=decks.MAX_UPLOAD)

    def test_original_legacy_ppt_and_unicode_filename_are_preserved_safely(self):
        self.storage.return_value = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1Synthetic old PowerPoint'
        self.file(original_filename='C:\\folder\\تدريب\r\n.ppt')
        self.document['monthlyAssignment']['uploadedPresentation'] = {'id': 'file-1'}; self.save_document()
        response = self.get()
        self.assertEqual(response.content, self.storage.return_value)
        self.assertEqual(response['Content-Type'], decks.PPT)
        self.assertIn("filename*=utf-8''", response['Content-Disposition'].lower())
        self.assertNotIn('\r', response['Content-Disposition']); self.assertNotIn('\n', response['Content-Disposition'])
        self.assertNotIn('folder', response['Content-Disposition'])

    def test_upload_scope_status_size_and_extension_are_enforced_without_fallback(self):
        self.document['monthlyAssignment']['uploadedPresentation'] = {'id': 'file-1'}; self.save_document()
        variants = [{'learner_id': '8'}, {'learner_kind': 'commercial'}, {'section_ref': 'OTHER'},
                    {'status': 'pending'}, {'status': 'rejected'}, {'size_bytes': 0},
                    {'size_bytes': decks.MAX_UPLOAD + 1}, {'original_filename': 'file.exe'}]
        for changes in variants:
            with self.subTest(changes=changes):
                self.db.execute('DELETE FROM Learner.evidence_files'); self.file(**changes)
                self.assertEqual(self.get().status_code, 404)
        self.storage.assert_not_called()

    def test_missing_upload_and_storage_errors_do_not_create_a_different_presentation(self):
        self.document['monthlyAssignment']['uploadedPresentation'] = {'id': 'file-1'}; self.save_document()
        self.assertEqual(self.get().status_code, 404)
        self.file(); self.storage.side_effect = RuntimeError('Private storage details must not be exposed')
        response = self.get(); self.assertEqual(response.status_code, 503)
        self.assertNotIn('Private storage', response.content.decode())
        self.storage.side_effect = None; self.storage.return_value = b''
        self.assertEqual(self.get().status_code, 503)
        self.storage.return_value = b'partial data'
        self.assertEqual(self.get().status_code, 503)

    def test_no_presentation_and_incomplete_saved_slides_have_clear_errors(self):
        self.document['monthlyAssignment']['slides'] = []; self.save_document()
        self.assertEqual(self.get().status_code, 404)
        for slides in ([{'title': 'Title', 'body': ''}], ['invalid'], [{'title': 'Title', 'body': 'x' * 12001}], 'invalid'):
            self.document['monthlyAssignment']['slides'] = slides; self.save_document()
            self.assertEqual(self.get().status_code, 422)

    def test_saved_template_is_reused_without_altering_its_content_or_ownership(self):
        self.file('reference', original_filename='reference.pptx', section_ref='reused-owned-template')
        self.document['monthlyAssignment']['presentationDesign'] = {'evidenceId': 'reference', 'coverSlide': 1, 'contentSlide': 2}
        self.save_document(); response = self.get()
        self.assertEqual(response.status_code, 200)
        self.assertIn('Original saved learning', visible_text(response.content))
        self.storage.assert_called_once_with('approved-test', 'synthetic-file', max_bytes=decks.MAX_REFERENCE)
        self.db.execute("UPDATE Learner.evidence_files SET learner_id='8'")
        self.assertEqual(self.get().status_code, 404)

    def test_unavailable_or_invalid_reference_never_silently_changes_the_design(self):
        self.document['monthlyAssignment']['presentationDesign'] = {'accent': '000000'}; self.save_document()
        self.assertEqual(self.get().status_code, 422)
        self.document['monthlyAssignment']['presentationDesign'] = {'evidenceId': 'reference'}; self.save_document()
        self.assertEqual(self.get().status_code, 404)
        self.file('reference', original_filename='reference.pptx'); self.storage.return_value = b'invalid pptx'
        self.assertEqual(self.get().status_code, 503)

    def test_submission_timestamp_guards_races_and_preserves_utc_instants(self):
        self.assertEqual(self.get({'submittedAt': '2026-10-27T09:00:00Z'}).status_code, 409)
        for timestamp in ('invalid', '2026-02-31T12:00:00Z', '2026-10-28T09:00:00'):
            self.assertEqual(self.get({'submittedAt': timestamp}).status_code, 400)
        self.assertEqual(self.get({'submittedAt': '2026-10-28T10:00:00+01:00'}).status_code, 200)

    def test_database_and_record_failures_are_reported_without_writes(self):
        with patch.object(decks, 'query', side_effect=DatabaseError('private connection detail')):
            response = self.get(); self.assertEqual(response.status_code, 503)
            self.assertNotIn('private connection', response.content.decode())
        self.db.execute("UPDATE Learner.learning_reflection_submissions SET full_submission='invalid json'")
        self.assertEqual(self.get().status_code, 503)
        self.assertEqual(self.get(method='POST').status_code, 405)


class PersonalPresentationTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.document = {'id': IDENTIFIER, 'activityType': 'assignment', 'activityId': 'C1', 'status': 'accepted',
            'version': 7, 'submittedAt': SUBMITTED, 'monthlyAssignment': {'month': '2026-10',
                'slides': [{'title': 'Personal learning', 'body': 'Saved personal course content.'}]}}
        self.state = {'submissions': {'assignment:C1': self.document}, 'progress': [], 'evidence': {}}
        self.course = {'accountId': 22, 'moduleId': 'MODULE-A'}
        self.selected = self.enterContext(patch('coach_api.personal_learning.selected_course', return_value=(self.course, self.document)))
        self.load = self.enterContext(patch('learner_api.personal_learning_store.load', return_value=self.state))
        self.enterContext(patch.object(decks, 'query', side_effect=AssertionError('Do not read official learner tables')))
        self.storage = self.enterContext(patch('learner_api.evidence_storage.download_blob_bytes', return_value=reference_file()))

    def get(self, query=None):
        return unwrap(decks.download)(self.factory.get('/presentation', query or {}), UUID(IDENTIFIER), personal=True)

    def test_personal_content_is_downloaded_from_its_own_store_without_mutation(self):
        before = deepcopy(self.state)
        response = self.get({'version': '7', 'submittedAt': SUBMITTED})
        self.assertEqual(response.status_code, 200)
        self.assertIn('Saved personal course content.', visible_text(response.content))
        self.load.assert_called_once_with(22, 'MODULE-A')
        self.assertEqual(self.state, before)

    def test_personal_missing_foreign_and_stale_work_are_refused(self):
        self.assertEqual(self.get({'version': '6'}).status_code, 409)
        self.assertEqual(self.get({'version': '0'}).status_code, 400)
        self.document['status'] = 'draft'; self.assertEqual(self.get().status_code, 404)
        self.document['status'] = 'accepted'; self.document['activityType'] = 'reading'; self.assertEqual(self.get().status_code, 404)
        self.selected.side_effect = LookupError('Other course')
        self.assertEqual(self.get().status_code, 404)

    def test_personal_upload_is_owned_by_course_and_activity_and_must_be_approved(self):
        self.document['monthlyAssignment']['uploadedPresentation'] = {'id': 'personal-file'}
        file = dict(id='personal-file', filename='personal.pptx', sizeBytes=len(self.storage.return_value), status='approved',
                    sectionRef='C1', container='approved-test', blobName='personal-private-file')
        self.state['evidence']['personal-file'] = file
        self.assertEqual(self.get().content, self.storage.return_value)
        for changes in ({'sectionRef': 'OTHER'}, {'status': 'pending'}, {'deletedAt': SUBMITTED}):
            self.state['evidence']['personal-file'] = {**file, **changes}
            self.assertEqual(self.get().status_code, 404)
        self.state['evidence'] = {}
        self.assertEqual(self.get().status_code, 404)


class PresentationAccessTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.account = SimpleNamespace(subject_type='staff', subject_id=3, role='staff')
        self.staff = SimpleNamespace(id=3, username='Actual Coach', email='COACH@example.test', access='coach')
        self.enterContext(patch('login.permissions.authenticate_request', side_effect=self.authenticate))
        self.access = self.enterContext(patch('login.permissions._accesses_of', return_value={'coach'}))
        manager = self.enterContext(patch('coach_api.auth.StaffUser.objects', new=MagicMock()))
        manager.filter.return_value.only.return_value.first.return_value = self.staff
        self.enterContext(patch('coach_api.auth._staff_access', side_effect=lambda staff: staff.access))
        self.source = self.enterContext(patch.object(decks, '_official', return_value=({}, None)))
        self.enterContext(patch.object(decks, 'presentation_file', return_value=(b'powerpoint', 'saved.pptx', decks.PPTX, 'generated')))

    def authenticate(self, request):
        request.login_account = self.account
        return self.account

    def test_no_session_learner_and_unrelated_staff_are_refused_before_reading(self):
        with patch('login.permissions.authenticate_request', return_value=None):
            self.assertEqual(decks.download(self.factory.get('/presentation'), IDENTIFIER).status_code, 401)
        self.account.role = 'learner'
        self.assertEqual(decks.download(self.factory.get('/presentation'), IDENTIFIER).status_code, 403)
        self.account.role = 'staff'; self.access.return_value = {'tutor'}
        self.assertEqual(decks.download(self.factory.get('/presentation'), IDENTIFIER).status_code, 403)
        self.source.assert_not_called()

    def test_authenticated_coach_is_canonical_and_cannot_assert_another_owner(self):
        response = decks.download(self.factory.get('/presentation'), IDENTIFIER)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.source.call_args.args[0].coach_email, 'coach@example.test')
        request = self.factory.get('/presentation', {'ownerEmail': 'other@example.test'})
        self.assertEqual(decks.download(request, IDENTIFIER).status_code, 403)

    def test_admin_view_as_is_a_read_with_selected_coach_scope(self):
        self.staff.access = 'super-admin'; self.access.return_value = {'super-admin'}
        self.assertEqual(decks.download(self.factory.get('/presentation'), IDENTIFIER).status_code, 403)
        selected = SimpleNamespace(id=4, email='selected@example.test')
        with patch('coach_api.auth._find_coach_staff', return_value=selected):
            request = self.factory.get('/presentation', {'viewAsCoach': selected.email})
            self.assertEqual(decks.download(request, IDENTIFIER).status_code, 200)
            self.assertEqual(self.source.call_args.args[0].coach_email, selected.email)
            self.assertEqual(self.source.call_args.args[0].coach_view_as_admin.id, 3)

    def test_official_and_personal_urls_resolve_to_the_guarded_download(self):
        official = resolve(f'/coach_api/coach/marking-queue/{IDENTIFIER}/presentation')
        personal = resolve(f'/coach_api/coach/personal-marking/{IDENTIFIER}/presentation')
        self.assertIs(official.func, decks.download); self.assertIs(personal.func, decks.download)
        self.assertTrue(personal.kwargs['personal'])
