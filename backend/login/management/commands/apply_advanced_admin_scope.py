"""Replace the restricted workspace scope with approved Active PCP/ME learners.

The workbook is read in place. Names and IDs are never copied into the repo or
printed by this command. The scope table follows this project's out-of-band
schema convention for the login/enrolment database.
"""
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import connections, transaction

from learner_api.models import LearnerProfile


def selected_rows(workbook_path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise CommandError('openpyxl is required to read the supplied workbook.') from exc
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headings = {str(value or '').strip(): index for index, value in enumerate(next(rows))}
        required = {'Id', 'Current programme', 'Programme status'}
        if not required.issubset(headings):
            raise CommandError('The workbook is missing an expected learner column.')
        selected = {}
        for row in rows:
            if str(row[headings['Programme status']] or '').strip().casefold() != 'active':
                continue
            programme = str(row[headings['Current programme']] or '').strip().casefold()
            code = ('PCP' if 'project controls professional' in programme else
                    'ME' if 'marketing executive' in programme else None)
            if code is None:
                continue
            aptem_id = str(row[headings['Id']] or '').strip()
            if not aptem_id.isdigit() or aptem_id in selected:
                raise CommandError('The selected learner IDs must be unique, nonblank Aptem numbers.')
            selected[aptem_id] = code
        if len(selected) != 211:
            raise CommandError(f'Expected exactly 211 Active PCP/ME learners; found {len(selected)}.')
        return selected
    finally:
        workbook.close()


def select_named_rows(workbook_rows, profiles, names):
    """Resolve requested names uniquely within the approved workbook IDs."""
    def name_key(value):
        return ''.join(char for char in value.casefold() if char.isalnum())

    by_name = {}
    for aptem_id, full_name in profiles:
        aptem_id = str(aptem_id)
        if aptem_id in workbook_rows:
            by_name.setdefault(name_key(full_name), []).append(aptem_id)
    selected = {}
    requested = set()
    for name in names:
        key = name_key(name.strip())
        if not key or key in requested or len(by_name.get(key, [])) != 1:
            raise CommandError('Every requested learner name must match one unique approved profile.')
        requested.add(key)
        aptem_id = by_name[key][0]
        selected[aptem_id] = workbook_rows[aptem_id]
    return selected


class Command(BaseCommand):
    help = 'Replace the Advanced Admin allowlist with learners from the approved workbook.'

    def add_arguments(self, parser):
        parser.add_argument('--workbook', required=True)
        selection = parser.add_mutually_exclusive_group(required=True)
        selection.add_argument('--learner-name', action='append', dest='learner_names')
        selection.add_argument('--all-approved', action='store_true')
        parser.add_argument('--expect-existing-count', type=int, required=True)
        parser.add_argument('--apply', action='store_true')

    def handle(self, *args, **options):
        path = Path(options['workbook'])
        if not path.is_file():
            raise CommandError('Workbook not found.')
        workbook_rows = selected_rows(path)
        profiles = list(LearnerProfile.objects.using('enrolment').filter(
            aptem_id__in=[int(value) for value in workbook_rows]
        ).values_list('aptem_id', 'full_name'))
        if len(profiles) != 211:
            raise CommandError('The 211 approved IDs do not resolve uniquely in Learner.learners.')
        selected = (workbook_rows if options['all_approved'] else
                    select_named_rows(workbook_rows, profiles, options['learner_names']))
        target = connections['enrolment'].settings_dict
        self.stdout.write(f"Target: {target['HOST']} / {target['NAME']}; requested learners: {len(selected)}")
        if not options['apply']:
            self.stdout.write('Dry run; no changes made.')
            return
        with transaction.atomic(using='enrolment'):
            with connections['enrolment'].cursor() as cursor:
                cursor.execute('''create table if not exists login."Advanced_admin_learner_scope" (
                    aptem_id text primary key,
                    programme_code varchar(3) not null check (programme_code in ('PCP','ME')),
                    created_at timestamptz not null default now()
                )''')
                cursor.execute('select aptem_id from login."Advanced_admin_learner_scope"')
                existing = {row[0] for row in cursor.fetchall()}
                if len(existing) != options['expect_existing_count'] or existing - workbook_rows.keys():
                    raise CommandError('Existing Advanced Admin scope differs from the expected workbook scope; no changes made.')
                for aptem_id, code in selected.items():
                    cursor.execute('''insert into login."Advanced_admin_learner_scope" (aptem_id,programme_code)
                        values (%s,%s) on conflict (aptem_id) do update
                        set programme_code = excluded.programme_code''', [aptem_id, code])
                cursor.execute('''delete from login."Advanced_admin_learner_scope"
                    where aptem_id <> all(%s::text[])''', [list(selected)])
                cursor.execute('select aptem_id from login."Advanced_admin_learner_scope"')
                if {row[0] for row in cursor.fetchall()} != selected.keys():
                    raise CommandError('Scope membership did not verify; no changes made.')
        self.stdout.write(self.style.SUCCESS(
            f'Advanced Admin learner scope saved and verified: {len(selected)}.'))
