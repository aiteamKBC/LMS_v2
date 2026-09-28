"""Spreadsheet attendance and token-scoped post-event feedback access."""
from __future__ import annotations

import csv
from datetime import timedelta
from html import escape
from io import BytesIO, StringIO
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from login import email_azure
from login.invitations import frontend_base_url
from login.models import LoginAccount
from login.security import generate_token, hash_token

from .models import Event, FeedbackEventCampaign, FeedbackEventRecipient, FeedbackForm

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
TOKEN_TTL = timedelta(days=30)
PRESENT_VALUES = {'present', 'attended', 'yes', 'y', '1', 'حاضر', 'حاضرة'}
HEADER_ALIASES = {
    'name': {'name', 'full name', 'attendee name', 'student name', 'learner name', 'الاسم'},
    'email': {'email', 'email address', 'attendee email', 'البريد', 'البريد الإلكتروني'},
    'status': {'attendance status', 'status', 'attendance', 'حالة الحضور'},
}


def _normalise_header(value):
    return ' '.join(str(value or '').strip().casefold().replace('_', ' ').split())


def _header_map(values):
    normalised = [_normalise_header(value) for value in values]
    result = {}
    for key, aliases in HEADER_ALIASES.items():
        positions = [index for index, value in enumerate(normalised) if value in aliases]
        if len(positions) != 1:
            raise ValueError(f'The attendance file needs one {key.title()} column.')
        result[key] = positions[0]
    return result


def _xlsx_rows(upload):
    workbook = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _csv_rows(upload):
    raw = upload.read()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ValueError('CSV files must use UTF-8 encoding.') from exc
    return [list(row) for row in csv.reader(StringIO(text))]


def parse_attendance(upload):
    """Return a validation preview; no database writes occur here."""
    if upload is None:
        raise ValueError('Choose an attendance .xlsx or .csv file.')
    size = getattr(upload, 'size', 0)
    if size <= 0 or size > MAX_UPLOAD_BYTES:
        raise ValueError('Attendance files must be between 1 byte and 5 MB.')
    suffix = Path(getattr(upload, 'name', '')).suffix.casefold()
    if suffix == '.xlsx':
        try:
            rows = _xlsx_rows(upload)
        except Exception as exc:  # openpyxl has several malformed-workbook exceptions
            raise ValueError('The Excel workbook could not be read.') from exc
    elif suffix == '.csv':
        rows = _csv_rows(upload)
    else:
        raise ValueError('Upload an Excel .xlsx or CSV file.')
    if not rows:
        raise ValueError('The attendance file is empty.')
    columns = _header_map(rows[0])
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f'Attendance files can contain at most {MAX_ROWS} data rows.')

    accepted, ignored, errors, seen = [], 0, [], set()
    for row_number, row in enumerate(rows[1:], start=2):
        def cell(key):
            index = columns[key]
            return str(row[index] if index < len(row) and row[index] is not None else '').strip()
        name, email, status = cell('name'), cell('email').casefold(), cell('status').casefold()
        if not any((name, email, status)):
            continue
        if status not in PRESENT_VALUES:
            ignored += 1
            continue
        row_errors = []
        if not name:
            row_errors.append('name is required')
        try:
            validate_email(email)
        except ValidationError:
            row_errors.append('email is invalid')
        if email in seen:
            row_errors.append('email is duplicated')
        if row_errors:
            errors.append({'row': row_number, 'error': ', '.join(row_errors)})
            continue
        seen.add(email)
        accepted.append({'row': row_number, 'name': name[:255], 'email': email})
    return {'attendees': accepted, 'presentCount': len(accepted), 'ignoredCount': ignored, 'errors': errors}


def _learner_ids_by_email(emails):
    rows = LoginAccount.objects.filter(
        subject_type='learner', email__in=emails,
    ).values_list('email', 'subject_id')
    matches = {}
    for email, subject_id in rows:
        key = str(email).strip().casefold()
        matches.setdefault(key, set()).add(str(subject_id))
    return {email: next(iter(ids)) for email, ids in matches.items() if len(ids) == 1}


@transaction.atomic
def import_event_attendance(*, event_id, form_ids, preview, created_by):
    """Create/update one event roster and attach one or more post-event forms."""
    if preview['errors']:
        raise ValueError('Fix the attendance file errors before importing it.')
    event = Event.objects.filter(pk=event_id).first()
    if event is None:
        raise ValueError('Event not found.')
    form_ids = list(dict.fromkeys(int(value) for value in form_ids))
    forms = list(FeedbackForm.objects.filter(
        id__in=form_ids, form_type='post_event', status='published', is_current=True,
    ))
    if not form_ids or len(forms) != len(form_ids):
        raise ValueError('Choose published Post-event feedback forms only.')
    if FeedbackForm.objects.filter(
        id__in=form_ids, sections__questions__question_type='photo_upload',
    ).exists():
        raise ValueError('Post-event forms sent to guests cannot contain photo upload questions.')
    for form in forms:
        FeedbackEventCampaign.objects.update_or_create(
            event=event, form=form,
            defaults={'status': 'open', 'created_by': created_by},
        )
    learners = _learner_ids_by_email([item['email'] for item in preview['attendees']])
    active_emails = [item['email'] for item in preview['attendees']]
    FeedbackEventRecipient.objects.filter(
        event=event, revoked_at__isnull=True,
    ).exclude(attendee_email__in=active_emails).update(
        revoked_at=timezone.now(), token_hash=None, token_expires_at=None,
    )
    imported = []
    for item in preview['attendees']:
        recipient, _ = FeedbackEventRecipient.objects.update_or_create(
            event=event, attendee_email=item['email'],
            defaults={
                'attendee_name': item['name'],
                'learner_id': learners.get(item['email'], ''),
                'attendance_source': 'spreadsheet',
                'attendance_reference': f'row:{item["row"]}',
                'revoked_at': None,
            },
        )
        if not recipient.token_hash and recipient.invite_status == 'sent':
            recipient.invite_status = 'pending'
            recipient.invitation_sent_at = None
            recipient.invitation_error = ''
            recipient.save(update_fields=['invite_status', 'invitation_sent_at', 'invitation_error', 'updated_at'])
        imported.append(recipient)
    return event, forms, imported


def event_feedback_link(token):
    # Keep the bearer value in the URL fragment so browsers do not send it in
    # the initial page request, referrer headers, or reverse-proxy access logs.
    return f'{frontend_base_url()}/login#feedback={token}'


def send_event_invitations(event, recipients):
    """Mint a fresh hash-only token per recipient, then send outside a transaction."""
    results = []
    form_count = FeedbackEventCampaign.objects.filter(event=event, status='open').count()
    for recipient in recipients:
        token = generate_token()
        recipient.token_hash = hash_token(token)
        recipient.token_expires_at = timezone.now() + TOKEN_TTL
        recipient.invite_status = 'pending'
        recipient.invitation_error = ''
        recipient.save(update_fields=['token_hash', 'token_expires_at', 'invite_status', 'invitation_error', 'updated_at'])
        link = event_feedback_link(token)
        event_title = escape(event.title)
        name = escape(recipient.attendee_name)
        subject = f'Feedback for {event.title}'
        text = (
            f'Hello {recipient.attendee_name},\n\nThank you for attending {event.title}. '
            f'Please complete the {form_count} feedback form(s) available at:\n{link}\n\n'
            f'This personal link expires in {TOKEN_TTL.days} days and must not be shared.'
        )
        html = (
            f'<p>Hello {name},</p><p>Thank you for attending <strong>{event_title}</strong>.</p>'
            f'<p><a href="{escape(link)}">Open event feedback</a></p>'
            f'<p>This personal link expires in {TOKEN_TTL.days} days and must not be shared.</p>'
        )
        sent, detail = email_azure.send_mail(
            to=recipient.attendee_email, subject=subject, html_body=html,
            text_body=text, save_to_sent=True,
        )
        recipient.invite_status = 'sent' if sent else 'failed'
        recipient.invitation_sent_at = timezone.now() if sent else None
        recipient.invitation_error = '' if sent else str(detail or 'Email delivery failed')[:1000]
        recipient.save(update_fields=['invite_status', 'invitation_sent_at', 'invitation_error', 'updated_at'])
        results.append({'id': recipient.id, 'email': recipient.attendee_email, 'sent': sent})
    return results


def recipient_for_token(token):
    if not token:
        return None
    return FeedbackEventRecipient.objects.select_related('event').filter(
        token_hash=hash_token(token), token_expires_at__gt=timezone.now(),
        revoked_at__isnull=True,
    ).first()
