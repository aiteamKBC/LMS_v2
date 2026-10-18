"""Independent spreadsheet campaigns for public non-LMS feedback recipients."""
from __future__ import annotations

import csv
from datetime import timedelta
from html import escape
from io import BytesIO, StringIO
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db.models import F
from django.utils import timezone
from openpyxl import load_workbook

from login import email_azure
from login.invitations import frontend_base_url
from login.security import generate_token, hash_token

from .event_emails import event_logo_attachment
from .models import FeedbackExternalRecipient

CAMPAIGN_TYPES = {'information_session', 'commitment_meeting'}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
TOKEN_TTL = timedelta(days=30)
PUBLIC_PATH = '/session-feedback'


def _normalise_header(value):
    return ' '.join(str(value or '').strip().casefold().replace('_', ' ').split())


def _rows(upload):
    if upload is None:
        raise ValueError('Choose an Excel .xlsx or CSV file.')
    size = getattr(upload, 'size', 0)
    if size <= 0 or size > MAX_UPLOAD_BYTES:
        raise ValueError('Recipient files must be between 1 byte and 5 MB.')
    suffix = Path(getattr(upload, 'name', '')).suffix.casefold()
    if suffix == '.xlsx':
        try:
            workbook = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True)
            try:
                return [list(row) for row in workbook.active.iter_rows(values_only=True)]
            finally:
                workbook.close()
        except Exception as exc:
            raise ValueError('The Excel workbook could not be read.') from exc
    if suffix == '.csv':
        try:
            text = upload.read().decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('CSV files must use UTF-8 encoding.') from exc
        return [list(row) for row in csv.reader(StringIO(text))]
    raise ValueError('Upload an Excel .xlsx or CSV file.')


def parse_recipients(upload):
    """Return disjoint preview groups without writing or sending anything."""
    rows = _rows(upload)
    if not rows:
        raise ValueError('The recipient file is empty.')
    headers = [_normalise_header(value) for value in rows[0]]
    positions = {}
    for key in ('name', 'email'):
        matches = [index for index, value in enumerate(headers) if value == key]
        if len(matches) != 1:
            raise ValueError(f'The recipient file needs exactly one {key.title()} column.')
        positions[key] = matches[0]
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f'Recipient files can contain at most {MAX_ROWS} data rows.')

    valid, invalid_emails, missing_rows, duplicates, seen = [], [], [], [], set()
    for row_number, row in enumerate(rows[1:], start=2):
        value = lambda key: str(row[positions[key]] if positions[key] < len(row) and row[positions[key]] is not None else '').strip()
        name, raw_email = value('name'), value('email')
        if not name and not raw_email:
            continue
        email = raw_email.casefold()
        item = {'row': row_number, 'name': name, 'email': email}
        if not name or not email:
            missing_rows.append({**item, 'missing': 'Name' if not name else 'Email'})
            continue
        try:
            validate_email(email)
        except ValidationError:
            invalid_emails.append(item)
            continue
        if email in seen:
            duplicates.append(item)
            continue
        seen.add(email)
        valid.append({**item, 'name': name[:255]})
    return {
        'valid': valid, 'invalidEmails': invalid_emails,
        'missingRows': missing_rows, 'duplicates': duplicates,
    }


def parse_manual_recipients(text):
    """Parse one ``Name, Email`` recipient per line without writing or sending."""
    if not isinstance(text, str) or not text.strip():
        raise ValueError('Enter at least one recipient as Name, Email.')
    rows = list(csv.reader(StringIO(text)))
    if len(rows) > MAX_ROWS:
        raise ValueError(f'You can enter at most {MAX_ROWS} recipients.')
    if any(len(row) > 2 for row in rows):
        raise ValueError('Enter one recipient per line using Name, Email.')

    valid, invalid_emails, missing_rows, duplicates, seen = [], [], [], [], set()
    for row_number, row in enumerate(rows, start=1):
        name = str(row[0] if row else '').strip()
        raw_email = str(row[1] if len(row) > 1 else '').strip()
        if not name and not raw_email:
            continue
        email = raw_email.casefold()
        item = {'row': row_number, 'name': name, 'email': email}
        if not name or not email:
            missing_rows.append({**item, 'missing': 'Name' if not name else 'Email'})
            continue
        try:
            validate_email(email)
        except ValidationError:
            invalid_emails.append(item)
            continue
        if email in seen:
            duplicates.append(item)
            continue
        seen.add(email)
        valid.append({**item, 'name': name[:255]})
    return {
        'valid': valid, 'invalidEmails': invalid_emails,
        'missingRows': missing_rows, 'duplicates': duplicates,
    }


def public_link(token):
    return f'{frontend_base_url()}{PUBLIC_PATH}#token={token}'


def _email_copy(campaign, recipient, link):
    label = 'Information Session' if campaign.campaign_type == 'information_session' else 'Commitment Meeting'
    subject = f'{label} feedback: {campaign.name}'
    greeting = f'Hello {recipient.recipient_name},'
    intro = f'Please share your feedback for {campaign.name}, held on {campaign.session_date.isoformat()}.'
    notice = f'This personal link expires in {TOKEN_TTL.days} days, can be used once, and must not be shared.'
    safe_link = escape(link, quote=True)
    html = f'''<!doctype html><html><body style="font-family:Segoe UI,Arial,sans-serif;background:#f4f1f8;padding:24px">
<div style="max-width:600px;margin:auto;background:#fff;border-radius:14px;padding:30px">
<p>{escape(greeting)}</p><p>{escape(intro)}</p>
<p><a href="{safe_link}" style="display:inline-block;background:#5b21b6;color:#fff;padding:12px 22px;border-radius:8px;text-decoration:none;font-weight:700">Open feedback form</a></p>
<p style="font-size:12px;color:#6f6478">{escape(notice)}</p></div></body></html>'''
    text = f'{greeting}\n\n{intro}\n\nOpen feedback form: {link}\n\n{notice}'
    return subject, text, html


def send_invitations(recipients):
    """Send fresh token links; callers select pending/failed recipients only."""
    results = []
    for recipient in recipients:
        token = generate_token()
        recipient.token_hash = hash_token(token)
        recipient.token_expires_at = timezone.now() + TOKEN_TTL
        recipient.delivery_status = 'pending'
        recipient.invitation_error = ''
        recipient.save(update_fields=['token_hash', 'token_expires_at', 'delivery_status', 'invitation_error', 'updated_at'])
        subject, text, html = _email_copy(recipient.campaign, recipient, public_link(token))
        sent, detail = email_azure.send_mail(
            to=recipient.recipient_email, subject=subject, html_body=html,
            text_body=text, save_to_sent=True, attachments=[event_logo_attachment()],
        )
        recipient.delivery_status = 'sent' if sent else 'failed'
        recipient.invitation_sent_at = timezone.now() if sent else None
        recipient.invitation_error = '' if sent else str(detail or 'Email delivery failed')[:1000]
        recipient.save(update_fields=['delivery_status', 'invitation_sent_at', 'invitation_error', 'updated_at'])
        results.append({'id': recipient.id, 'sent': bool(sent)})
    return results


def recipient_for_token(token, *, lock=False):
    if not token:
        return None
    queryset = FeedbackExternalRecipient.objects.select_related('campaign__form')
    if lock:
        queryset = queryset.select_for_update()
    return queryset.filter(
        token_hash=hash_token(token), token_expires_at__gt=timezone.now(),
        campaign__status='open', campaign__form__status='published',
        campaign__form__form_type=F('campaign__campaign_type'),
    ).first()
